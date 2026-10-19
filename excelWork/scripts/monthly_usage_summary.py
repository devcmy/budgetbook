"""Replace September–December's monthly footer with a six-line summary.

Patch the original package, not an openpyxl/LibreOffice save of the XLSM.
Only summary rows 72–86, print areas, and appended summary styles may change.
Prepayments here are gross payments excluded from the consumption calendar;
they are not personal consumption, card-bill repayments, or wallet top-ups.
"""
import argparse
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl

from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from update_cash_details_20261019 import import_changed_caches
from validate_excel_structure import validate_package
from verify_reserve_redesign import assert_no_errors, recalculate

MONTHS = (9, 10, 11, 12)
ACTUAL_STATUSES = ('확정', '결제', 'RAW 확인', '기존 입력')
ADVANCE_CATEGORIES = ('대납', '대리결제', '미리결제', '선결제')
LABELS = ('월급 예정', '고정비', '가계부 달력합계', '테트리스 달력합계',
          '실 사용금액 합계', '최종 사용금액')
RAW_END = 10018  # Existing workbook's provisioned future-entry range.


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def fixed_total(rows, basis='actual'):
    if basis not in ('actual', 'actual_and_planned'):
        raise ValueError(basis)
    return sum(amount or 0 for status, amount in rows
               if basis == 'actual_and_planned' or status in ACTUAL_STATUSES)


def advance_total(rows, start):
    """Sum original payment G, not personal-burden K (zero for an advance).

    Recovery status is deliberately not filtered: this is this month's gross
    payment volume, not the outstanding receivable at the end of the month.
    """
    return sum(row[6] or 0 for row in rows
               if row[1] == start and row[7] == '지출'
               and row[9] in ADVANCE_CATEGORIES)


def totals(fixed, calendar, advances, additional):
    base = fixed + calendar + advances
    return base, base + additional


def advance_formula():
    r = lambda c: f'RAW!${c}$2:${c}${RAW_END}'
    return '+'.join(f'SUMIFS({r("G")},{r("B")},$T$1,{r("H")},"지출",'
                    f'{r("J")},"{category}")' for category in ADVANCE_CATEGORIES)


def fixed_formula(basis):
    if basis == 'actual_and_planned':
        return '=SUM(Q4:Q24)'
    return '=' + '+'.join(f'SUMIF($U$4:$U$24,"{status}",$Q$4:$Q$24)'
                           for status in ACTUAL_STATUSES)


def prepare(baseline, candidate, basis):
    p = ExtendedPackage(baseline)
    # Keep every existing font/fill/style ID; append a dark bold body font and
    # three summary cell formats (the title uses a white font on a dark fill).
    head = p.cell('10월', 'A72').get('s')
    body = p.cell('10월', 'A73').get('s')
    money = p.cell('10월', 'E73').get('s')
    style_tree = p.tree('xl/styles.xml')
    xfs = style_tree.find(Q('cellXfs'))
    fonts = style_tree.find(Q('fonts'))
    bold_font = deepcopy(fonts[int(xfs[int(body)].get('fontId'))])
    if bold_font.find(Q('b')) is None:
        ET.SubElement(bold_font, Q('b'), val='1')
    bold_font_id = str(len(fonts))
    fonts.append(bold_font)
    fonts.set('count', str(len(fonts)))
    styles = {'header': head, 'body': body, 'money': money}
    for key, original in [('bold', body), ('bold_money', money), ('right_money', money)]:
        node = deepcopy(xfs[int(original)])
        if key != 'right_money':
            node.set('fontId', bold_font_id)
        if key.endswith('money'):
            align = node.find(Q('alignment'))
            if align is None:
                align = ET.SubElement(node, Q('alignment'))
            align.set('horizontal', 'right')
            align.set('vertical', 'center')
        styles[key] = str(len(xfs))
        xfs.append(node)
    xfs.set('count', str(len(xfs)))
    w = openpyxl.load_workbook(baseline, data_only=True)
    for month in MONTHS:
        name = f'{month}월'
        s = p.sheet(name)
        p.cell(name, 'A72')  # Initialize cell/row indices.
        data = s.find(Q('sheetData'))
        # Drop the old explanatory footer, not any financial source cells.
        for row in list(data):
            if 72 <= int(row.get('r')) <= 86:
                for c in row:
                    p.cell_maps[name].pop(c.get('r'), None)
                p.row_maps[name].pop(int(row.get('r')), None)
                data.remove(row)
        merges = s.find(Q('mergeCells'))
        if merges is None:
            merges = ET.Element(Q('mergeCells'))
            # Schema order: after sheetData, before conditionalFormatting etc.
            s.insert(list(s).index(data) + 1, merges)
        for merge in list(merges):
            if int(''.join(filter(str.isdigit, merge.get('ref').split(':')[0]))) >= 72:
                merges.remove(merge)
        def merge(ref):
            ET.SubElement(merges, Q('mergeCell'), ref=ref)
        merge('A72:U72')
        for r in range(73, 79):
            for ref in [f'A{r}:D{r}', f'E{r}:H{r}', f'J{r}:U{r}']:
                merge(ref)
        merges.set('count', str(len(merges)))
        p.write(name, 'A72', '=TEXT($T$1,"m월")&" 사용금액 요약"', True)
        p.cell(name, 'A72').set('s', styles['header'])
        salary = w[name]['C37'].value or w['10월']['E73'].value
        formulas = {'E74': fixed_formula(basis), 'E75': '=I33',
                    'E76': '=SUM(I4,I9,I14,I19,I24,I29)',
                    'E77': '=E74+E75+' + advance_formula(),
                    'E78': '=E77+H39'}
        for r, label in enumerate(LABELS, 73):
            p.write(name, f'A{r}', label)
            p.cell(name, f'A{r}').set('s', styles['bold' if r >= 77 else 'body'])
            p.cell(name, f'E{r}').set('s', styles['bold_money' if r >= 77 else 'right_money'])
            p.cell(name, f'J{r}').set('s', styles['body'])
        # Each month keeps an explicit reference estimate, never booked income.
        p.write(name, 'E73', salary)
        for addr, formula in formulas.items():
            p.write(name, addr, formula, True)
        p.write(name, 'J73', '참고 예정액 · 실제 급여는 위 수입표에서 확인합니다.')
        if basis == 'actual':
            p.write(name, 'J74', '="확정·기존 입력 기준. 미결제 예정 "&TEXT(MAX(0,Q25-E74),"#,##0")&"원은 합계에서 제외."', True)
        else:
            p.write(name, 'J74', '실적 + 미결제 예정 기준 · 미결제 비용도 아래 합계에 포함합니다.')
        p.write(name, 'J75', '달력 실지출 합계 · 대리·미리결제는 별도로 더합니다.')
        p.write(name, 'J76', '거래별 천원 올림 합계 · 아래 사용금액 합계에는 다시 더하지 않습니다.')
        p.write(name, 'J77', '="고정비 + 가계부 달력 + 대리·미리결제 "&TEXT(E77-E74-E75,"#,##0")&"원. 대납 포함 결제 총액이며 본인 소비와 구분."', True)
        p.write(name, 'J78', '="실 사용금액 + 추가 실지출 "&TEXT(H39,"#,##0")&"원 (이미 집계된 금액 제외)."', True)
        for r in range(72, 79):
            p.row_maps[name][r].set('ht', '27' if r == 72 else '32')
            p.row_maps[name][r].set('customHeight', '1')
        s.find(Q('dimension')).set('ref', 'A1:AA78')
    # Only these four print areas change; preserve workbook properties and IDs.
    book = p.tree('xl/workbook.xml')
    sheetnames = [n.get('name') for n in book.find(Q('sheets'))]
    for defined in book.find(Q('definedNames')):
        if defined.get('name') == '_xlnm.Print_Area':
            name = sheetnames[int(defined.get('localSheetId'))]
            if name in {f'{m}월' for m in MONTHS}:
                defined.text = f"'{name}'!$A$1:$V$78"
    p.save(candidate)
    w.close()


def verify(baseline, final, basis):
    b = openpyxl.load_workbook(baseline, data_only=True)
    w = openpyxl.load_workbook(final, data_only=True)
    formulas = openpyxl.load_workbook(final)
    assert w.sheetnames == b.sheetnames and '검증' not in w.sheetnames
    rows = list(w['RAW'].values)[1:]
    report = {}
    for month in MONTHS:
        name = f'{month}월'; s = w[name]; original = b[name]
        fixed = fixed_total([(s[f'U{r}'].value, s[f'Q{r}'].value) for r in range(4, 25)], basis)
        advances = advance_total(rows, s['T1'].value)
        calendar = s['I33'].value
        tetris = sum(s[f'I{r}'].value or 0 for r in [4, 9, 14, 19, 24, 29])
        additional = sum(s[f'H{r}'].value or 0 for r in range(52, 62))
        base, final_total = totals(fixed, calendar, advances, additional)
        expected = [s['E73'].value, fixed, calendar, tetris, base, final_total]
        assert [s[f'A{r}'].value for r in range(73, 79)] == list(LABELS)
        assert [s[f'E{r}'].value for r in range(73, 79)] == expected, (name, expected)
        for addr in ['A77', 'A78', 'E77', 'E78']:
            assert s[addr].font.bold
            assert s[addr].font.color == s['A73'].font.color, (name, addr, 'legible total text')
        assert s['H39'].value == additional
        assert not any(c.value is not None for row in s.iter_rows(min_row=79, max_row=86) for c in row)
        assert str(formulas[name].print_area) == f"'{name}'!$A$1:$V$78"
        # The complete original month, all calendar days and inputs, unchanged.
        for row in original.iter_rows(max_row=71):
            for c in row:
                assert c.value == s[c.coordinate].value, (name, c.coordinate, c.value, s[c.coordinate].value)
        report[name] = dict(salary_reference=expected[0], fixed=fixed, calendar_actual=calendar,
                            calendar_tetris=tetris, gross_advances=advances,
                            additional_actual=additional, base_actual=base, final_actual=final_total)
    assert report['10월']['gross_advances'] == 103600
    assert report['10월']['additional_actual'] == 110075
    for month in [11, 12]:
        assert report[f'{month}월']['gross_advances'] == 0
        assert report[f'{month}월']['calendar_actual'] == 0
        if basis == 'actual':
            assert report[f'{month}월']['final_actual'] == 0
    assert w['잔액']['G16'].value == b['잔액']['G16'].value == 6559691
    assert w['잔액']['B248'].value == '출금'
    assert w['RAW']['G2376'].value == -1300
    assert_no_errors(w)
    structure = validate_package(final)
    assert not structure['errors'], structure
    before, after = ExtendedPackage(baseline), ExtendedPackage(final)
    allowed = {before.sheets[f'{m}월'] for m in MONTHS} | {'xl/styles.xml', 'xl/workbook.xml'}
    with ZipFile(baseline) as bz, ZipFile(final) as az:
        assert bz.namelist() == az.namelist() and az.testzip() is None
        changed = [p for p in bz.namelist() if bz.read(p) != az.read(p)]
        assert set(changed) <= allowed, changed
        for month in MONTHS:
            name = f'{month}월'
            for c in before.sheet(name).iter(Q('c')):
                r = int(''.join(filter(str.isdigit, c.get('r'))))
                if r < 72:
                    assert ET.tostring(c) == ET.tostring(after.cell(name, c.get('r'))), (name, c.get('r'))
        oldstyle, newstyle = before.tree('xl/styles.xml'), after.tree('xl/styles.xml')
        for component in oldstyle:
            target = newstyle.find(component.tag)
            if component.tag in {Q('cellXfs'), Q('fonts')}:
                assert [ET.tostring(x) for x in component] == [ET.tostring(x) for x in list(target)[:len(component)]]
                assert len(target) == len(component) + (3 if component.tag == Q('cellXfs') else 1)
            else:
                assert ET.tostring(component) == ET.tostring(target)
        oldbook, newbook = before.tree('xl/workbook.xml'), after.tree('xl/workbook.xml')
        # Ignore only the four authorized print-area changes.
        for book in [oldbook, newbook]:
            for item in book.find(Q('definedNames')):
                if item.get('name') == '_xlnm.Print_Area' and any(f"'{m}월'!" in (item.text or '') for m in MONTHS):
                    item.text = 'authorized monthly print area'
        assert ET.tostring(oldbook) == ET.tostring(newbook)
    for x in [b, w, formulas]:
        x.close()
    return dict(date='2026-10-19', baseline_sha256=digest(baseline), workbook_sha256=digest(final),
                months=report, fixed_basis=basis, summary_rows=[72, 78],
                raw_ledger_reserve_other_months_calculators_byte_identical=True,
                vba_drawings_controls_relationships_byte_identical=True,
                existing_styles_preserved=True, changed_parts=changed,
                formula_errors=0, structure_validation=structure,
                native_excel_opened=False, sharepoint_file_updated=False)


def mutation_regression(final, folder, basis):
    """Verify future entries and shared extra-expense formulas in a throwaway copy."""
    base = openpyxl.load_workbook(final, data_only=True)
    original_tetris = base['10월']['E76'].value
    original_tetris_cell = base['10월']['I4'].value
    base10, final10 = [base['10월'][a].value for a in ['E77', 'E78']]
    original9 = [base['9월'][a].value for a in ['I33', 'E77', 'E78']]
    base.close()
    def edit(p):
        p.write('11월', 'O4', 400000)
        p.write('11월', 'U4', '확정')
        p.write('11월', 'D52', 2000)
        p.write('11월', 'D53', 17000)
        p.write('11월', 'G53', 17000)  # Already counted: additional zero.
        for r, month, category, kind, amount in [
            (2494, 11, '대납', '지출', 11200),
            (2495, 11, '선결제', '지출', 8880),
            (2496, 11, '내부이체', '이체지출', 20080),
            (2497, 11, '대납', '수입', 1000),
            (2498, 12, '미리결제', '지출', 9900),
        ]:
            for c, value in {'A': datetime(2026, month, 15), 'B': datetime(2026, month, 1),
                             'C': '카드', 'F': '회귀 검증용',
                             'G': amount, 'H': kind, 'J': category, 'K': 0,
                             'L': '정산완료'}.items():
                p.write('RAW', f'{c}{r}', value)
        p.sheet('RAW').find(Q('dimension')).set('ref', 'A1:S2498')
        p.write('10월', 'I4', original_tetris_cell + 1000)
        p.write('9월', 'I33', original9[0] + 777)
    output = recalculate(final, folder, edit)
    w = openpyxl.load_workbook(output, data_only=True)
    for month in [11, 12]:
        s = w[f'{month}월']
        fixed = fixed_total([(s[f'U{r}'].value, s[f'Q{r}'].value) for r in range(4, 25)], basis)
        advance = 20080 if month == 11 else 9900
        extra = 2000 if month == 11 else 0
        assert s['E74'].value == fixed
        assert s['E77'].value == fixed + advance
        assert s['E78'].value == fixed + advance + extra
    assert w['11월']['H53'].value == 0
    assert w['10월']['E76'].value == original_tetris + 1000
    assert w['10월']['E77'].value == base10 and w['10월']['E78'].value == final10
    assert w['9월']['E77'].value == original9[1] + 777
    assert w['9월']['E78'].value == original9[2] + 777
    assert_no_errors(w)
    w.close()
    return dict(future_fixed_status_change=True, future_raw_rows_beyond_current_end=True,
                month_filtering=True, original_zero_personal_burden_advance_included=True,
                recovered_status_remains_gross_payment=True, transfer_and_recovery_not_new_payment=True,
                already_counted_extra_not_duplicated=True, tetris_not_added_to_actual=True,
                calendar_actual_changes_flow_to_both_totals=True)


def update_audit(previous, report, output):
    a = json.loads(Path(previous).read_text())
    a['current_verification'] = 'monthly_usage_summary_20261019_validation.json'
    a['workbook_sha256'] = report['workbook_sha256']
    # The old E73:E86 addresses belong to the removed, superseded footer.
    a['october_summary'] = {k: v for k, v in a['october_summary'].items()
                            if k not in {'E73', 'E74', 'E75', 'E76', 'E77', 'E78', 'E80'}}
    a['monthly_usage_summary_20261019'] = {
        **report, 'details': 'monthly_usage_summary_20261019_validation.json',
        'guide': '2026-10-19_월간사용금액요약_검증.md',
        'previous_version': 'history/2026-10-19/05'}
    a['delta_20261019']['scope'] = 'RAW 갱신 당시 검증. 거래·달력·잔액 수치는 유지하며 요약표와 최신 파일 해시는 monthly_usage_summary_20261019_validation.json을 따릅니다.'
    Path(output).write_text(json.dumps(a, ensure_ascii=False, indent=2) + '\n')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for key in ['baseline', 'workdir', 'output', 'report']:
        parser.add_argument('--' + key, required=True)
    parser.add_argument('--basis', choices=['actual', 'actual_and_planned'], default='actual')
    parser.add_argument('--audit-input')
    parser.add_argument('--audit-output')
    args = parser.parse_args()
    folder = Path(args.workdir)
    candidate = folder / 'candidate.xlsm'
    prepare(args.baseline, candidate, args.basis)
    print('Four six-line summaries prepared. Recalculating a disposable macro-free copy.', flush=True)
    calculated = recalculate(candidate, folder / 'default')
    caches = import_changed_caches(candidate, calculated, args.output)
    result = verify(args.baseline, args.output, args.basis)
    result['changed_formula_caches'] = caches
    result['mutation_regression'] = mutation_regression(args.output, folder / 'mutations', args.basis)
    Path(args.report).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    if args.audit_input and args.audit_output:
        update_audit(args.audit_input, result, args.audit_output)
    print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
