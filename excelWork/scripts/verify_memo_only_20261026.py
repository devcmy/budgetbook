"""Regression tests for independent memos and their live cross-tab displays."""
from datetime import datetime
from hashlib import sha256
import argparse
import json
import math
from pathlib import Path
import re
import subprocess
import sys

import openpyxl
from openpyxl.utils.datetime import to_excel

from memo_only_20261026 import TABLE
from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from validate_excel_structure import validate_package
from verify_reserve_redesign import recalculate
from verify_raw_auto_link_20261026 import errors,append_row


def financial_snapshot(path):
    w=openpyxl.load_workbook(path,data_only=True);assert not errors(w),errors(w)[:10]
    result={}
    for s in w:
        for c in s._cells.values():
            if s.title=='RAW' and c.column in [9,13,40,41]:continue
            if re.fullmatch(r'\d+월',s.title) and (c.coordinate=='AC40' or (35<=c.row<=39 and 22<=c.column<=27)):continue
            if re.fullmatch(r'\d+월계산기',s.title) and 158<=c.row<=162 and c.column<=21:continue
            value=to_excel(c.value) if isinstance(c.value,datetime) else c.value
            if isinstance(value,(int,float,bool)):
                result[s.title,c.coordinate]=value
    for m in range(1,13):
        s=w[f'{m}월'];t=w[f'{m}월계산기']
        for address in ['I33','H37','H38','H44','C41','E78','E79']:
            result[s.title,address]=s[address].value
        for address in ['W12','W13']+[f'C{r}' for r in range(3,11)]:result[t.title,address]=t[address].value
    w.close();return result


def same_finances(reference,actual):
    a=financial_snapshot(reference);b=financial_snapshot(actual)
    def equal(x,y):
        if isinstance(x,(int,float)) and isinstance(y,(int,float)):return math.isclose(x,y,rel_tol=0,abs_tol=1e-6)
        return x==y
    differences=[(key,x,b.get(key)) for key,x in a.items() if not equal(x,b.get(key))]
    assert not differences,differences[:20]
    return {'numeric_cells_checked':len(a),'financial_differences':0}


def clear_memos(p):
    for r in range(2,4002):p.write('RAW',f'I{r}',None)


def note_rows(w):
    # Independent ledger-only historical notes are intentionally not RAW notes.
    linked_ledger={r for r in range(64,501) if any(isinstance(w['잔액'][f'{c}{r}'].value,(int,float)) and w['잔액'][f'{c}{r}'].value>=2 for c in ['W','X'])}
    return linked_ledger


def check_cleared(source,result,metadata):
    check=same_finances(source,result);w=openpyxl.load_workbook(result,data_only=True)
    ledger=note_rows(w);checked=0
    for name,addr in metadata['memo_destinations']:
        if name=='예비비' and addr.startswith('I') and w[name][f'C{int(addr[1:])}'].value not in ledger:continue
        assert w[name][addr].value in [None,''],(name,addr,w[name][addr].value)
        checked+=1
    for m in range(1,13):assert w[f'{m}월']['AC40'].value==0
    w.close();return {**check,'raw_linked_memo_cells_empty':checked,'all_12_month_browsers_empty':True}


def edit_memos(p):
    clear_memos(p)
    for r,text in [(2444,'이자 메모 수정 시험'),(2493,'보험 메모 수정 시험'),(356,'예비비 메모 수정 시험')]:p.write('RAW',f'I{r}',text)


def check_edited(source,result):
    check=same_finances(source,result);w=openpyxl.load_workbook(result,data_only=True)
    for name,addr,expected in [('10월','V13','이자 메모 수정 시험'),('10월','I54','보험 메모 수정 시험'),
                               ('잔액','G249','보험 메모 수정 시험'),('예비비','H102','예비비 메모 수정 시험'),
                               ('10월','W35','이자 메모 수정 시험'),('10월','W36','보험 메모 수정 시험'),
                               ('10월계산기','B158','이자 메모 수정 시험'),('10월계산기','B159','보험 메모 수정 시험'),
                               ('2월계산기','B158','예비비 메모 수정 시험')]:
        assert w[name][addr].value==expected,(name,addr,w[name][addr].value,expected)
    w.close();return {**check,'monthly_fixed_extra_and_browsers_follow_raw':True,'calculator_balance_reserve_follow_raw':True}


def business_fixture(p):
    # Empty prepared registers also follow a RAW/ledger link added in the future.
    p.write('RAW','I2493','미래 연결 메모')
    p.write('잔액','W250',2493)
    p.write('예비비','F399',2493)
    p.write('예비비','C729',250)
    # Identical business fields but different note text are duplicates. A real
    # separate purchase needs the explicit M discriminator, not a note keyword.
    for row,memo in [(2494,'대납 · 대리결제'),(2495,'분류와 관계없는 다른 메모'),(2496,'')]:
        append_row(p,row,10,26,4200,title='메모 분리 결제 시험',memo=memo)
    p.write('RAW','M2496','실제별도거래2')
    append_row(p,2497,10,26,51800,title='명시 대납 시험',memo='평범한 메모')
    p.write('RAW','J2497','대납')
    # Page 2 is mirrored by the calculator; a note containing a wildcard stays text.
    for k in range(6):
        append_row(p,2498+k,11,26,1000+k,title=f'월 메모 시험{k}',memo=f'11월 메모 {k} ? *')
    p.write('11월','W40',2)
    for m in range(1,13):append_row(p,2510+m,m,26,1500+m,title=f'12개월 회귀 {m}')


def check_business(source,result):
    a=openpyxl.load_workbook(source,data_only=True);w=openpyxl.load_workbook(result,data_only=True)
    assert not errors(w),errors(w)[:10]
    s=w['RAW']
    assert s['J2494'].value=='변동비' and s['Y2494'].value==4200
    assert s['Q2495'].value==2 and s['Y2495'].value==0
    assert s['Q2496'].value==1 and s['Y2496'].value==4200
    assert s['T2497'].value=='대납' and s['W2497'].value==0 and s['AG2497'].value==51800
    for m in range(1,13):
        expected=1500+m+(8400 if m==10 else 6015 if m==11 else 0)
        rounded=2000+(10000 if m==10 else 11000 if m==11 else 0)
        assert w[f'{m}월']['I33'].value-a[f'{m}월']['I33'].value==expected,(m,'calendar')
        assert w[f'{m}월계산기']['W12'].value-a[f'{m}월계산기']['W12'].value==rounded,(m,'calculator')
    assert w['10월']['E79'].value-a['10월']['E79'].value==8400+1510+51800
    assert w['11월']['W35'].value==w['11월계산기']['B158'].value=='11월 메모 5 ? *'
    assert w['11월']['W36'].value in [None,'']
    assert w['잔액']['G250'].value==w['예비비']['H399'].value==w['예비비']['I729'].value=='미래 연결 메모'
    assert w['잔액']['G16'].value==a['잔액']['G16'].value
    assert w['예비비']['I13'].value==a['예비비']['I13'].value
    a.close();w.close()
    return {'note_keywords_do_not_classify':True,'memo_changes_do_not_distinguish_duplicates':True,
            'explicit_discriminator_retains_real_separate_purchase':True,'explicit_advance_keeps_zero_personal_cost':True,
            'memo_paging_and_literal_wildcards':True,'future_linked_registers_show_memos':True,
            'all_12_month_raw_calendar_calculator_regression':True}


def structure_checks(source,final):
    structure=validate_package(final);assert not structure['errors'],structure
    w=openpyxl.load_workbook(final)
    assert len(w.sheetnames)==46 and '검증' not in w.sheetnames
    raw=w['RAW'];table=raw.tables[TABLE]
    assert table.ref=='A1:AO4001' and table.tableColumns[8].name=='메모' and table.tableColumns[12].name=='동일거래 구분'
    assert raw.protection.sheet and not raw.protection.autoFilter and raw.protection.sort
    assert not raw['I2494'].protection.locked and not raw['M2494'].protection.locked
    refs=[]
    for r in range(2,4002):
        for column in ['P','Q','AJ']:
            formula=raw[f'{column}{r}'].value
            if re.search(r'\bI\$?'+str(r)+r'\b|\$I\$2|\[메모\]',formula):refs.append((r,column))
    assert not refs,refs[:10]
    w.close()
    a,b=ExtendedPackage(source),ExtendedPackage(final)
    for n in a.parts:
        if n=='xl/vbaProject.bin' or n.startswith(('xl/media/','xl/activeX/','xl/ctrlProps/','xl/drawings/')):
            assert a.parts[n]==b.parts[n],n
    return {'structure':structure,'vba_controls_drawings_media_preserved':True,
            'memo_is_single_source_column':'RAW!I','classification_and_dedupe_do_not_read_memo':True,
            'protected_input_and_month_filter_preserved':True}


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for name in ['source','workbook','metadata','workdir','report']:parser.add_argument('--'+name,required=True)
    args=parser.parse_args();folder=Path(args.workdir);folder.mkdir(parents=True,exist_ok=True)
    metadata=json.loads(Path(args.metadata).read_text())
    report=structure_checks(args.source,args.workbook)
    report['baseline_financial_preservation']=same_finances(args.source,args.workbook)
    # Original memo texts are merged losslessly into I, not parsed as instructions.
    a=openpyxl.load_workbook(args.source,data_only=True);b=openpyxl.load_workbook(args.workbook,data_only=True)
    for r in range(2,4002):
        old=[str(a['RAW'][f'{c}{r}'].value) for c in ['I','M'] if a['RAW'][f'{c}{r}'].value not in [None,'']]
        assert (b['RAW'][f'I{r}'].value or '')=='\n'.join(dict.fromkeys(old)),r
    a.close();b.close();report['legacy_note_texts_preserved_in_single_column']=True
    for name,edit,check in [('cleared',clear_memos,lambda p:check_cleared(args.workbook,p,metadata)),
                            ('edited',edit_memos,lambda p:check_edited(args.workbook,p)),
                            ('business',business_fixture,lambda p:check_business(args.workbook,p))]:
        result=recalculate(args.workbook,folder/name,edit);report[name]=check(result)
        print(json.dumps({'passed':name,**report[name]},ensure_ascii=False),flush=True)
    tests=subprocess.run([sys.executable,'-m','unittest','discover','-s',str(Path(__file__).parent),'-p','test_*.py'],capture_output=True,text=True,check=True)
    report['unit_tests_passed']=int(re.search(r'Ran (\d+) tests',tests.stderr)[1])
    report['workbook_sha256']=sha256(Path(args.workbook).read_bytes()).hexdigest()
    report['source_sha256']=sha256(Path(args.source).read_bytes()).hexdigest()
    report['native_excel_ui_executed']=False
    report['date']='2026-10-26';report['data_as_of']='2026-10-19'
    report['previous_version']='history/2026-10-26/02'
    report['known_limitations']=['RAW full reset still uses historical anchors; not a reset workflow.',
                               'Replacing old rows can retain J/K/L overrides; append new records to empty rows.',
                               'Independent handwritten non-RAW notes and hidden historical archives are retained.']
    Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'report':args.report,'tests':report['unit_tests_passed']},ensure_ascii=False))
