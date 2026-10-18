"""Append the confirmed October 16–18 delta to the XML-repaired XLSM.

Keep stable RAW/ledger identities, original controls and the selected-month UI.
Only disposable macro-free copies are recalculated; never deliver a Calc export.
"""
import argparse
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
import re

from lxml import etree as ET
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter

from redesign_reserve import ExtendedPackage
from update_20261015 import Q, CATEGORIES
from validate_excel_structure import validate_package


ASOF = datetime(2026, 10, 18)
FIRST = 2323


def prepare(baseline, source, previous_source, output, summary):
    w = openpyxl.load_workbook(baseline)
    d = openpyxl.load_workbook(baseline, data_only=True)
    incoming = openpyxl.load_workbook(source, data_only=True).active
    previous = openpyxl.load_workbook(previous_source, data_only=True).active
    assert list(next(incoming.values)) == list(next(previous.values)) == [
        '날짜', '자산', '분류', '소분류', '내용', '금액(원)', '수입/지출', '메모', '금액', '화폐']
    normalize = lambda v: tuple(None if x == '' else x for x in v)
    before = Counter(normalize(v) for v in list(previous.values)[1:])
    after = Counter(normalize(v) for v in list(incoming.values)[1:])
    assert not before - after, 'Prior source transactions changed or disappeared.'
    assert sum(before.values()) == 144 and sum(after.values()) == 164
    raw = d['RAW']
    assert raw.max_row == 2466
    old = defaultdict(list)
    for r in range(FIRST, raw.max_row + 1):
        old[(raw[f'A{r}'].value, raw[f'G{r}'].value, raw[f'H{r}'].value)].append(r)
    additions, mapping, matched = [], [], set()
    for sr in range(2, incoming.max_row + 1):
        v = list(normalize(incoming.cell(sr, c).value for c in range(1, 11)))
        date, asset, _, _, _, amount, kind, _, amount2, currency = v
        assert isinstance(date, datetime) and date.year == 2026 and date.month == 10
        assert currency == 'KRW' and amount == amount2 and isinstance(amount, int)
        matches = [r for r in old[(date, amount, kind)] if r not in matched]
        if matches:
            assert len(matches) == 1
            rr = matches[0]
            assert raw[f'C{rr}'].value == ('모임통장' if asset == '20260606' else asset)
            matched.add(rr)
            mapping.append({'source_row': sr, 'workbook_row': rr, 'status': 'preserved'})
        else:
            additions.append((sr, v))
    assert len(matched) == 144 and len(additions) == 20
    additions.sort(key=lambda item: (item[1][0], item[0]))
    p = ExtendedPackage(baseline)
    end = raw.max_row + len(additions)
    # Extend only exact RAW range endpoints, including fixed October checks.
    for sheet in w:
        if sheet.title.startswith(('기존_', '보관_')):
            continue
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    formula = re.sub(r'(RAW!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?)2466\b',
                                     lambda m: m[1] + str(end), cell.value)
                    if formula != cell.value:
                        p.write(sheet.title, cell.coordinate, formula, True)
    new_rows, personal, colors = [], [], Counter()
    for rr, (sr, v) in enumerate(additions, raw.max_row + 1):
        date, asset, category, sub, title, amount, kind, note, _, _ = v
        payment, direction, burden = '본인', None, amount
        memo = f'10/18 첨부 RAW {sr}행 · 신규 거래. 기존 RAW 행번호와 이전 분류는 보존.'
        if asset == '미래에셋_WRAP형CMA':
            assert kind == '이체지출' and category == '카카오페이머니' and amount == 15700
            classification, scope, direction = '내부이체', 'WRAP', '이체출금'
            memo += ' 카카오페이 충전은 WRAP에서 1회 차감; 소비는 후속 결제에서만 반영.'
            assert d['잔액']['A248'].value is None
            ledger = {'A248': date, 'B248': '이체', 'C248': '생활비', 'D248': '외부',
                      'E248': amount, 'F248': '카카오페이머니 충전',
                      'G248': f'RAW {rr}행. WRAP→카카오페이머니 충전 15,700원. 실제 구매 17,000원은 RAW에서 소비로 별도 집계하며 WRAP에 중복 차감하지 않음.',
                      'J248': asset, 'K248': category, 'L248': 'WRAP계좌이체', 'W248': rr}
            for addr, value in ledger.items():
                p.write('잔액', addr, value)
        elif asset == '20260606':
            assert kind == '지출' and category == '여행/약속/공사'
            asset, classification, scope, payment, burden = '모임통장', '약속·만남', '타계좌', '모임통장', amount / 2
            memo += ' 기존 모임통장 분류·본인분 1/2 기준 유지. 월급 납입과 중복 소비 제외.'
        else:
            assert kind == '지출' and (asset == '카카오페이머니' or '카드' in asset or asset == 'KB쿠팡와우')
            scope = '타계좌' if asset == '카카오페이머니' else '카드사용'
            classification = '변동비'
            palette = {'식비': 3, '기타': 4, '문화/생활': 6, '패션/미용': 6,
                       '건강/보험/의료': 7, '잡화/소품/가구/생필품': 9}[category]
            units = (amount + 999) // 1000
            personal.append({'row': rr, 'date': date.isoformat(), 'amount': amount,
                             'units': units, 'calculator_row': palette})
            colors[palette] += units
            memo += ' 본인 생활비·달력·계산기 반영. WRAP 원장에는 결제를 재기록하지 않음.'
        p.clone_row('RAW', 2464, rr)
        values = [date, datetime(2026, 10, 1), asset, category, sub, title, amount, kind,
                  note, classification, burden, payment, memo, None, None, None, None, scope, direction]
        for col, value in enumerate(values, 1):
            p.write('RAW', f'{get_column_letter(col)}{rr}', value)
        link = '(COUNTIF(잔액!$W$64:$W$500,ROW())+COUNTIF(잔액!$X$64:$X$500,ROW()))'
        p.write('RAW', f'N{rr}', f'=IF(R{rr}="WRAP",IF({link}=1,"연결",IF({link}=0,"누락","중복")),IF(R{rr}="카드사용","월간·계산기","타계좌·잔액제외"))', True)
        mapping.append({'source_row': sr, 'workbook_row': rr, 'status': 'added'})
        new_rows.append({'source_row': sr, 'row': rr, 'date': date.isoformat(),
                         'amount': amount, 'classification': classification, 'scope': scope})
    assert len(personal) == 17 and sum(r['amount'] for r in personal) == 177160
    assert sum(colors.values()) == 180
    for tag in ['dimension', 'autoFilter']:
        p.sheet('RAW').find(Q(tag)).set('ref', f'A1:S{end}')
    for cf in p.sheet('RAW').findall(Q('conditionalFormatting')):
        cf.set('sqref', cf.get('sqref').replace('N2323:N2466', f'N2323:N{end}'))
    p.write('RAW', 'Q1', '10월 RAW 갱신 · 2026-10-18 기준')
    p.write('RAW', 'Q2', '10/1~18 RAW 164건: 기존 144건 유지 + 신규 20건. WRAP 15건 연결·타계좌 48건·카드사용 101건.')
    p.write('10월', 'U1', ASOF)
    for date_row in [3, 8, 13, 18, 23, 28]:
        for col in 'BCDEFGH':
            date = d['10월'][f'{col}{date_row}'].value
            if not isinstance(date, datetime) or not 16 <= date.day <= 18:
                continue
            assert date.month == 10
            actual_row, tetris_row = date_row + 2, date_row + 1
            p.write('10월', f'{col}{actual_row}', Translator(w['10월']['E15'].value, origin='E15').translate_formula(f'{col}{actual_row}'), True)
            rr = f'RAW!$A${FIRST}:$A${end}'
            cl = f'RAW!$J${FIRST}:$J${end}'
            terms = '+'.join(f'({cl}="{c}")' for c in sorted(CATEGORIES))
            p.write('10월', f'{col}{tetris_row}',
                    f'=SUMPRODUCT(({rr}>={col}{date_row})*({rr}<{col}{date_row}+1)*(RAW!$H${FIRST}:$H${end}="지출")*(RAW!$L${FIRST}:$L${end}="본인")*({terms})*ROUNDUP(RAW!$K${FIRST}:$K${end}/1000,0)*1000)', True)
    p.write('10월', 'E75', f'=SUMIFS(RAW!$G${FIRST}:$G${end},RAW!$C${FIRST}:$C${end},"모임통장",RAW!$H${FIRST}:$H${end},"지출")', True)
    for addr, formula in {'E83': f'=COUNT(RAW!G{FIRST}:G{end})-164', 'E84': '=I33-1476900',
                          'E85': '=SUM(I4,I9,I14,I19,I24,I29)-1501000'}.items():
        p.write('10월', addr, formula, True)
    p.write('10월', 'A70', '10/1~18 RAW 164건. 기존 144건 유지 + 신규 20건. 본인 생활비 177,160원 추가; 공금 사용·카카오페이 충전은 중복 소비 제외.')
    p.write('10월', 'A72', '10/18 실제 / 예정 / 공동자금 · 달력 합계와 연결')
    p.write('10월', 'J83', '0이면 최신 10/18 첨부 RAW 164건 모두 유지. 기존 분류·행번호 보존.')
    p.write('10월', 'J86', '급여 미입금: 10/18 중간 집계이며 월말 결과가 아님')
    for r in range(4, 25):
        note = w['10월'][f'V{r}'].value
        if note:
            p.write('10월', f'V{r}', note.replace('10/15 기준', '10/18 기준').replace('10/15 RAW 집계', '10/18 RAW 집계'))
    # Paint only the new tail, retaining every pre-existing color and border.
    calc = w['10월계산기']
    order = [f'{get_column_letter(c)}{r}' for r0 in range(15, 145, 5)
             for c0 in range(2, 22, 2) for c in range(c0, c0 + 2) for r in range(r0, r0 + 5)]
    assert [calc[a].fill.patternType == 'solid' for a in order] == [True] * 1321 + [False] * (len(order) - 1321)
    xfs = p.tree('xl/styles.xml').find(Q('cellXfs'))
    cache = {ET.tostring(x): i for i, x in enumerate(xfs)}
    index = 1321
    for row in personal:
        for _ in range(row['units']):
            cell = p.cell('10월계산기', order[index])
            xf = deepcopy(xfs[int(cell.get('s', '0'))])
            legend = p.cell('10월계산기', f'B{row["calculator_row"]}')
            xf.set('fillId', xfs[int(legend.get('s', '0'))].get('fillId'))
            key = ET.tostring(xf)
            if key not in cache:
                cache[key] = len(xfs)
                xfs.append(xf)
            cell.set('s', str(cache[key]))
            index += 1
    xfs.set('count', str(len(xfs)))
    for r in range(3, 11):
        p.write('10월계산기', f'AA{r}', d['10월계산기'][f'AA{r}'].value + colors[r])
    p.write('10월계산기', 'A154', '10/18까지 본인 생활비 반영. 이전 색상·블록 유지 + 180칸 추가. 공금 사용·미회수 대납 103,600원·생신 예비비 239,000원은 중복 합산하지 않음.')
    p.write('잔액', 'D3', '2026-10-18 기준 · WRAP RAW 15건 연결. 카카오페이 충전 15,700원 1회 차감. 카드 사용·타계좌 결제는 WRAP 중복 제외.')
    p.write('검증', 'I44', '=COUNTIF(잔액!$I$231:$I$248,"입력 확인")+COUNTIF(잔액!$I$231:$I$248,"WRAP 범위 확인")', True)
    p.write('검증', 'I70', '=COUNT(잔액!W231:X248)', True)
    for row, expected in [(35, 164), (63, 164), (64, 15), (65, 48), (66, 15), (70, 15),
                          (171, 2553539), (174, 6669766)]:
        p.write('검증', f'M{row}', f'=I{row}={expected}', True)
    p.write('검증', 'A72', '10/18 WRAP 15건만 연결. 타계좌 48건·카드사용 101건은 RAW에 보존. 생신 사용 239,000원·철도 미회수 대납 103,600원 별도 표시.')
    p.write('검증', 'Q14', '10/18까지 · 대납 2건 103,600원·가족 공동자금 분리')
    p.write('사용안내', 'A22', '10월은 10/18까지 실제 자료입니다. 미수령 급여·미결제 고정비·공동자금·대납은 월간 확인표에 별도 표시합니다.')
    p.write('사용안내', 'A23', '10월 RAW 164건 중 WRAP 15건을 잔액 원장과 대조합니다. 타계좌 48건·카드사용 101건을 보존하며 충전·실제 결제를 중복 소비로 잡지 않습니다.')
    cp = p.tree('xl/workbook.xml').find(Q('calcPr'))
    for key, value in [('calcMode', 'auto'), ('fullCalcOnLoad', '1'), ('forceFullCalc', '1')]:
        cp.set(key, value)
    p.save(output)
    structure = validate_package(output)
    assert not structure['errors'], structure
    result = {'as_of': ASOF.date().isoformat(), 'source_sha256': sha256(Path(source).read_bytes()).hexdigest(),
              'baseline_sha256': sha256(Path(baseline).read_bytes()).hexdigest(),
              'preserved_october_rows': 144, 'october_count': 164, 'raw_total': 2485,
              'last_raw_row': end, 'new_rows': new_rows, 'new_personal': personal,
              'source_mapping': sorted(mapping, key=lambda x: x['source_row']),
              'new_personal_actual': 177160, 'new_personal_tetris': 180000,
              'personal_actual': 1476900, 'personal_tetris': 1501000,
              'wrap_total': 6669766, 'living_balance': 4519766,
              'new_common_spend': 36800, 'new_wrap_transfer': 15700,
              'selected_month_reserve_usage': 239000, 'outstanding_advances': 103600,
              'structure_validation': structure}
    Path(summary).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({k: v for k, v in result.items() if k not in ['new_rows', 'new_personal', 'source_mapping']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['baseline', 'source', 'previous-source', 'output', 'summary']:
        parser.add_argument('--' + name, required=True)
    a = parser.parse_args()
    prepare(a.baseline, a.source, a.previous_source, a.output, a.summary)
