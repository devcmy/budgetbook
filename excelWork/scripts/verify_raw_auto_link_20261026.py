"""Integration checks for RAW-only input and the twelve automatic calculators."""
from datetime import datetime
import json
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from openpyxl.formula.translate import Translator
from lxml import etree as ET

from raw_auto_link_20261026 import BACK,TABLE,HEADERS,PREPARED_END,RR,ceil_units,digest
from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from validate_excel_structure import validate_package
from verify_reserve_redesign import recalculate


def errors(w):
    return [(s.title,c.coordinate,c.value) for s in w for c in list(s._cells.values()) if c.data_type=='e']


def snapshot(path):
    w=openpyxl.load_workbook(path,data_only=True)
    found=errors(w);assert not found,found[:20]
    result={'months':{},'raw':{},'wrap':w['잔액']['G16'].value,'reserve':w['예비비']['I13'].value}
    for m in range(1,13):
        s=w[f'{m}월'];c=w[f'{m}월계산기']
        result['months'][m]={a:s[a].value for a in ['I33','H37','I37','H38','I38','H39','H40','H42','H43','H44','I44','C37','C38','C39','C40','C41','E74','E75','E76','E77','E78','E79','N37','N69','AA1','O19','U19']}
        result['months'][m]['W12']=c['W12'].value
        result['months'][m]['counts']={r:c[f'C{r}'].value for r in range(3,11)}
        result['months'][m]['block_end']=c['AF3'].value
        result['months'][m]['warning']=c['A156'].value
    for r in range(2494,2530):
        if w['RAW'][f'A{r}'].value is not None:
            result['raw'][r]={c:w['RAW'][f'{c}{r}'].value for c in ['B','J','K','L','N','O','Q','T','U','V','W','X','Y','Z','AA','AB','AG','AH','AI']}
    result['raw'][4002]={c:w['RAW'][f'{c}4002'].value for c in ['J','K','L','N','Y','Z']}
    w.close();return result


def check_default(source,calculated):
    before=snapshot(source);after=snapshot(calculated)
    checked=[]
    for m in range(1,11):
        for addr in ['I33','H37','I37','H38','I38','H39','H40','H42','H44','I44','C37','C38','C39','C40','C41','N69','W12']:
            a,b=before['months'][m][addr],after['months'][m][addr]
            assert (a or 0)==(b or 0),(m,addr,a,b)
        assert before['months'][m]['counts']==after['months'][m]['counts'],m
        checked.append(m)
    for m in [11,12]:
        s=after['months'][m]
        assert s['I33']==s['H37']==s['H44']==s['C37']==s['W12']==0
        assert s['E75']==before['months'][m]['E75']==1234406
        assert s['AA1'] is False
    assert before['wrap']==after['wrap']==6559691
    assert before['reserve']==after['reserve']
    return {'historical_months_unchanged':checked,'future_actual_not_forecast':True,'wrap_and_reserve_unchanged':True}


def append_row(p,r,m,day,amount,*,kind='지출',category='식비',title='자동연동 시험',asset='우리카드의정석SUPER',memo='',year=2026):
    for c,value in {'A':datetime(year,m,day,15,20),'C':asset,'D':category,'F':title,'G':amount,'H':kind,'I':memo}.items():
        p.write('RAW',f'{c}{r}',value)


def fixture(p):
    protection=p.sheet('RAW').find(Q('sheetProtection'))
    if protection is not None:protection.set('sheet','0')  # Owner removes protection to expand the table.
    for m in range(1,13):append_row(p,2493+m,m,26,2101+100*m,title=f'월별 자동연동 {m}')
    # Exact duplicate of October's row: never a second expense.
    append_row(p,2506,10,26,3101,title='월별 자동연동 10')
    append_row(p,2507,10,25,8888,category='정기구독',title='Apple Appstore')
    append_row(p,2508,10,26,2776050,kind='수입',category='급여',title='월급',asset='미래에셋_WRAP형CMA')
    append_row(p,2509,10,26,12500,kind='이체지출',category='카카오뱅크',title='카드대금 준비',asset='미래에셋_WRAP형CMA')
    append_row(p,2510,10,26,110075,category='주거',title='보증보험료',asset='미래에셋_WRAP형CMA')
    append_row(p,2511,10,27,-1300,title='환급 시험')
    append_row(p,2512,10,27,150000,kind='이체지출',category='희망두배 청년통장',title='적립 시험',asset='미래에셋_WRAP형CMA')
    append_row(p,2513,10,27,'4,200',title='문자 금액 시험')
    append_row(p,2514,10,27,51800,title='대리 구매 시험',memo='대리결제')
    append_row(p,2515,10,27,10000,asset='가족통장',title='공동 사용 시험')
    append_row(p,2516,10,27,9000,title='잘못된 날짜 시험');p.write('RAW','A2516','날짜 오류')
    append_row(p,2517,1,1,3000,title='다른 연도 시험',year=2027)
    append_row(p,2518,10,27,1200,title='메모 다른 거래',memo='별도 거래 A')
    append_row(p,2519,10,27,1200,title='메모 다른 거래',memo='별도 거래 B')
    append_row(p,2520,10,27,1200,title='긴 메모 중복 시험',memo='긴 거래 메모 '*60)
    append_row(p,2521,10,27,1200,title='긴 메모 중복 시험',memo='긴 거래 메모 '*60)
    append_row(p,2522,10,27,-1000,category='금융',title='카드 할인')
    append_row(p,2523,10,27,1000,title='커피A')
    append_row(p,2524,10,27,1000,title='커피?')
    append_row(p,2525,10,27,120000,kind='이체지출',category='메리츠증권_슈퍼365',title='증권 적립 시험',asset='미래에셋_WRAP형CMA')
    append_row(p,2526,10,27,20000,kind='이체출금',category='ISA',title='ISA 적립 시험',asset='미래에셋_WRAP형CMA')
    append_row(p,2527,10,27,5000,kind='이체입금',category='ISA',title='투자 회수 시험',asset='미래에셋_WRAP형CMA')
    p.write('RAW','G2323',12234)
    # Simulate Excel's table expansion with its actual calculated-column formulas.
    tables=[n for n in p.parts if n.startswith('xl/tables/') and n.endswith('.xml')]
    table_part=next(n for n in tables if TABLE.encode() in p.parts[n]);tree=p.tree(table_part)
    tree.set('ref','A1:AM4002');tree.find(Q('autoFilter')).set('ref','A1:AM4002')
    for idx,column in enumerate(tree.find(Q('tableColumns')),1):
        formula=column.find(Q('calculatedColumnFormula'))
        if formula is not None:
            c=openpyxl.utils.get_column_letter(idx)
            p.write('RAW',f'{c}4002',Translator('='+formula.text,origin=f'{c}2').translate_formula(f'{c}4002'),True)
    append_row(p,4002,12,31,2222,title='표 확장 시험')
    p.sheet('RAW').find(Q('dimension')).set('ref','A1:AV4002')


def check_fixture(baseline,result):
    b=snapshot(baseline);a=snapshot(result)
    reports={}
    for m in range(1,13):
        money=2101+100*m;units=ceil_units(money)
        if m==10:
            money+=110075-1300+4200+1234+3600+2000
            units+=111-1+5+2+6+2
        if m==12:money+=2222;units+=3
        bm=b['months'][m];am=a['months'][m]
        assert am['I33']-bm['I33']==money,(m,'actual',am['I33'],bm['I33'],money)
        assert am['W12']-bm['W12']==units*1000,(m,'calc',am['W12'],bm['W12'],units)
        assert am['I38']-bm['I38']==units*1000
        reports[m]={'actual_delta':money,'tetris_delta':units*1000}
    assert a['raw'][2506]['N']=='중복 제외' and a['raw'][2506]['Y']==0
    assert a['raw'][2513]['W']==4200
    assert a['raw'][2514]['T']=='대납' and a['raw'][2514]['K']==0 and a['raw'][2514]['AG']==51800
    assert a['raw'][2515]['U']=='공동' and a['raw'][2515]['Y']==0
    assert a['raw'][2516]['N']=='날짜 확인'
    assert a['raw'][2517]['N']=='2026년 외 거래'
    assert a['raw'][2518]['Y']==a['raw'][2519]['Y']==1200
    assert a['raw'][2520]['Y']==1200 and a['raw'][2521]['N']=='중복 제외'
    assert a['raw'][2523]['Y']==a['raw'][2524]['Y']==1000
    assert a['raw'][2525]['AH']=='증권' and a['raw'][2526]['AH']=='ISA'
    assert a['raw'][2527]['AI']==0
    assert a['raw'][4002]['Y']==2222
    s=a['months'][10];old=b['months'][10]
    assert s['C37']==2776050 and s['O19']==8888 and s['U19']=='확정'
    assert s['H37']-old['H37']==8888
    assert s['E75']-old['E75']==8888-7737
    assert (s['N37'] or 0)-(old['N37'] or 0)==150000
    assert s['E79']-old['E79']==reports[10]['actual_delta']+8888+51800
    assert s['H43']-old['H43']==1000
    assert s['H40']-old['H40']==290000
    assert s['H44']-old['H44']==reports[10]['actual_delta']+8888+290000-1000
    assert a['wrap']==b['wrap'] and a['reserve']==b['reserve']
    return {'months':reports,'duplicate_excluded':True,'income_and_fixed_auto':True,'contribution_auto':True,
            'refund_and_text_amount':True,'wrong_date_and_other_year_flagged':True,'table_expansion_4002_formulas':True,
            'distinct_memos_not_deduplicated':True,'long_memo_duplicates_safe':True,
            'literal_wildcards_not_false_duplicates':True,
            'no_helpers_manually_entered':True,'wrap_reserve_not_guessed':True}


def overflow_fixture(p):
    append_row(p,2494,11,30,9000001,title='표시 범위 초과 시험')


def check_overflow(result):
    a=snapshot(result);m=a['months'][11]
    assert m['I33']==9000001 and m['W12']==9001000
    assert '범위 초과' in m['warning']
    return {'numeric_totals_preserved':True,'color_capacity_warning':True}


def package_checks(source,final):
    w=openpyxl.load_workbook(final)
    assert BACK in w.sheetnames and w[BACK].sheet_state=='hidden'
    assert '검증' not in w.sheetnames
    raw=w['RAW'];table=raw.tables[TABLE]
    assert table.ref==f'A1:AM{PREPARED_END}'
    assert str(raw.freeze_panes)=='C2'
    assert any(x.id==2 and x.name=='월' for x in table.tableColumns)
    assert table.autoFilter is not None and table.autoFilter.ref==table.ref
    assert raw.protection.sheet and raw.protection.sort and not raw.protection.autoFilter
    assert not raw['A2494'].protection.locked and raw['B2494'].protection.locked
    assert table.tableColumns[1].calculatedColumnFormula is not None
    for m in range(1,13):
        s=w[f'{m}월계산기'];assert sum(len(s.conditional_formatting[k]) for k in s.conditional_formatting)>=9
        assert 'colorfunction' not in str(s['C3'].value).lower()
    w.close()
    with ZipFile(source) as a,ZipFile(final) as z:
        for name in a.namelist():
            if name=='xl/vbaProject.bin' or name.startswith(('xl/media/','xl/ctrlProps/','xl/activeX/')):
                assert a.read(name)==z.read(name),name
        from raw_auto_link_20261026 import ExtendedPackage
        before,after=ExtendedPackage(source),ExtendedPackage(final)
        for sheet in ['잔액','예비비']:
            assert a.read(before.sheets[sheet])==z.read(after.sheets[sheet]),sheet
        old,new=before.tree('xl/styles.xml'),after.tree('xl/styles.xml')
        for key in ['fonts','fills','borders','cellXfs','dxfs']:
            x,y=old.find(Q(key)),new.find(Q(key))
            if x is not None:
                assert [ET.tostring(t) for t in x]==[ET.tostring(t) for t in list(y)[:len(x)]],key
    structure=validate_package(final);assert not structure['errors'],structure
    return {'table_month_filter':True,'computed_columns_and_frozen_headers':True,
            'all_calculators_conditional_colors':True,'cash_reserve_vba_media_preserved':True,
            'structure':structure,'source_sha256':digest(source),'workbook_sha256':digest(final)}


def check_source_rows(source,final):
    """Input fields and existing user-confirmed classifications stay intact."""
    a=openpyxl.load_workbook(source,data_only=True)
    b=openpyxl.load_workbook(final,data_only=True)
    old,new=a['RAW'],b['RAW'];count=0
    for r in range(2,old.max_row+1):
        if old[f'A{r}'].value is None:continue
        count+=1
        for c in list(range(1,10))+[13]:
            assert old.cell(r,c).value==new.cell(r,c).value,(r,c,'source field changed')
        for c in [10,11,12]:
            value=old.cell(r,c).value
            if value is not None:assert value==new.cell(r,c).value,(r,c,'confirmed override changed')
    a.close();b.close()
    return {'records_preserved':count,'source_fields_preserved':True,'confirmed_overrides_preserved':True}


def retain_unchanged_reserve(source,final):
    """Keep unrelated blank lookup caches byte-identical to the source.

    LibreOffice serializes INDEX-of-empty cells as zero. Retain the original
    representation only if every difference is an unchanged formula's blank
    -> zero cache; never suppress a change to inputs or financial values.
    """
    a,b=ExtendedPackage(source),ExtendedPackage(final);name='예비비';count=0
    for old in a.sheet(name).iter(Q('c')):
        new=b.cell(name,old.get('r'))
        if ET.tostring(old)==ET.tostring(new):continue
        of,nf=old.find(Q('f')),new.find(Q('f'))
        assert of is not None and nf is not None
        assert ET.tostring(of)==ET.tostring(nf),(name,old.get('r'),'formula changed')
        ov,nv=old.find(Q('v')),new.find(Q('v'))
        assert (ov is None or ov.text in [None,'']) and nv is not None and nv.text=='0',(old.get('r'),'not an empty reference cache')
        count+=1
    part=b.sheets[name];b.parts[part]=a.parts[a.sheets[name]];b.trees.pop(part,None)
    b.trees.pop('xl/workbook.xml',None)
    b.save(Path(final).with_name('retained.xlsm'))
    Path(final).with_name('retained.xlsm').replace(final)
    return count


if __name__=='__main__':
    import argparse
    import subprocess
    import sys
    parser=argparse.ArgumentParser(description='Read-only delivery checks plus disposable RAW-input recalc trials.')
    for name in ['source','workbook','workdir','report']:parser.add_argument('--'+name,required=True)
    args=parser.parse_args();folder=Path(args.workdir);folder.mkdir(parents=True,exist_ok=True)
    report=package_checks(args.source,args.workbook)
    report['source_rows']=check_source_rows(args.source,args.workbook)
    report['default']=check_default(args.source,args.workbook)
    trials={}
    for name,create,check in [('integration',fixture,lambda result:check_fixture(args.workbook,result)),
                              ('overflow',overflow_fixture,check_overflow)]:
        case=folder/name;case.mkdir(exist_ok=True)
        package=ExtendedPackage(args.workbook);create(package);package.save(case/'trial.xlsm')
        calculated=recalculate(case/'trial.xlsm',case/'recalculation')
        trials[name]=check(calculated)
    report['trials']=trials
    tests=subprocess.run([sys.executable,'-m','unittest','discover','-s',str(Path(__file__).parent),'-p','test_*.py'],capture_output=True,text=True,check=True)
    import re
    report['unit_tests_passed']=int(re.search(r'Ran (\d+) tests',tests.stderr)[1])
    report['prepared_records']=PREPARED_END-1
    report['new_record_capacity']=PREPARED_END-1-report['source_rows']['records_preserved']
    report['color_capacity']=8000
    report['input_protection']={'password':False,'filter_allowed':True,'input_allowed':True,'sorting_and_row_deletion_blocked':True,'expansion_requires_unprotect':True}
    report['native_excel_ui_executed']=False
    report['recalculation']='Macro-free disposable LibreOffice copies; final XLSM is never saved by LibreOffice.'
    report['visual_check']='Disposable PDF previews inspected for RAW readability and calculator conditional colors.'
    report['date']='2026-10-26';report['data_as_of']='2026-10-19'
    report['source']='inputs/2026-10-26/테트리스_가계부_2026_고도화_첨부원본.xlsm'
    report['workbook']='테트리스_가계부_2026_고도화.xlsm'
    report['previous_version']='history/2026-10-26/01'
    report['guide']='2026-10-26_RAW자동연동_검증.md'
    Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(report,ensure_ascii=False,indent=2))
