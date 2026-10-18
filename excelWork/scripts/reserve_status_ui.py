"""Patch the reserve UI without changing stable usage/settlement registers.

The savings plan remains at rows 19:29. A distinct, initially unfilled usage
plan is inserted at rows 32:40; display rows 32:69 move to 42:79. The caller
owns the financial engine and supplies its available/payment balance formulas.
"""
from copy import deepcopy
from datetime import datetime
import re

from lxml import etree as ET
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter as col, column_index_from_string

from redesign_reserve import FUNDS
from update_20261015 import Q


SHEET = '예비비'
FIRST, LAST, OFFSET = 32, 69, 10
CELL = re.compile(r'^(\$?[A-Z]{1,3})(\$?)(\d+)$')


def _move_reference(reference):
    """Relocate cell references, preserving absolute markers and range shape."""
    def move(match):
        prefix, absolute, number = match.groups()
        number = int(number)
        return f'{prefix}{absolute}{number + OFFSET if FIRST <= number <= LAST else number}'
    return ':'.join(CELL.sub(move, cell) for cell in reference.split(':'))


def _move_formula(formula, local=False):
    """Move only this sheet's references, never external ledger/source ranges."""
    original_equals = formula.startswith('=')
    tokens = Tokenizer(formula if original_equals else '=' + formula).items
    for token in tokens:
        if token.type != 'OPERAND' or token.subtype != 'RANGE':
            continue
        if '!' in token.value:
            sheet, address = token.value.rsplit('!', 1)
            if sheet.strip("'").replace("''", "'") == SHEET:
                token.value = sheet + '!' + _move_reference(address)
        elif local:
            token.value = _move_reference(token.value)
    result = ''.join(token.value for token in tokens)
    return '=' + result if original_equals else result


def apply_ui(p, available_formula, payment_formula):
    """Apply the display-only migration to an ExtendedPackage and return ranges."""
    s = p.sheet(SHEET)
    data = s.find(Q('sheetData'))
    assert not any(70 <= int(row.get('r')) <= 79 and len(row) for row in data), 'UI migration already applied or destination occupied'

    # Rewrite references before moving their owning rows. Absolute references
    # into stable registers (100:399, 430:729) and other sheets are unchanged.
    for name, part in p.sheets.items():
        # A read-only inspection must not enroll every untouched worksheet for
        # reserialization. Only retain another sheet's tree when a reference
        # actually changes, or when the caller already owns that edited tree.
        loaded = part in p.trees
        tree = p.trees[part] if loaded else ET.fromstring(p.parts[part])
        changed = False
        for cell in tree.iter(Q('c')):
            formula = cell.find(Q('f'))
            if formula is not None and formula.text and (name == SHEET or '예비비' in formula.text):
                updated = _move_formula(formula.text, local=name == SHEET)
                if updated != formula.text:
                    formula.text = updated
                    changed = True
                    value = cell.find(Q('v'))
                    if value is not None:
                        cell.remove(value)
        if changed and not loaded:
            p.trees[part] = tree
    for row in data:
        number = int(row.get('r'))
        if FIRST <= number <= LAST:
            row.set('r', str(number + OFFSET))
            for cell in row:
                cell.set('r', _move_reference(cell.get('r')))
    data[:] = sorted(data, key=lambda row: int(row.get('r')))
    p.cell_maps.pop(SHEET, None)
    p.row_maps.pop(SHEET, None)

    merges = s.find(Q('mergeCells'))
    for merged in list(merges):
        cells = merged.get('ref').split(':')
        start, finish = int(re.search(r'\d+', cells[0])[0]), int(re.search(r'\d+', cells[-1])[0])
        if 7 <= start and finish <= 13:
            merges.remove(merged)
        else:
            merged.set('ref', _move_reference(merged.get('ref')))
    validations = s.find(Q('dataValidations'))
    for validation in validations:
        validation.set('sqref', ' '.join(_move_reference(area) for area in validation.get('sqref').split()))
        for key in ['formula1', 'formula2']:
            formula = validation.find(Q(key))
            if formula is not None and formula.text:
                formula.text = _move_formula(formula.text, local=True)

    def write(addr, value, template=None):
        if template:
            p.cell(SHEET, addr).set('s', p.cell(SHEET, template).get('s', '0'))
        p.write(SHEET, addr, value, isinstance(value, str) and value.startswith('='))

    def span(area, value, template):
        for addr in _area_cells(area):
            write(addr, None, template)
        write(area.split(':')[0], value, template)
        ET.SubElement(merges, Q('mergeCell'), ref=area)

    def height(number, points):
        p.cell(SHEET, f'A{number}')
        row = p.row_maps[SHEET][number]
        row.set('ht', str(points))
        row.set('customHeight', '1')

    actual = {r: deepcopy(p.cell(SHEET, f'L{r}').find(Q('f'))).text for r in range(8, 13)}
    write('A2', '노란색만 입력합니다. 계획은 참고용이며 실제 사용만 반영합니다. 사용가능액은 잔액 탭의 같은 목적 계정과 일치합니다.')
    write('A6', '1  현황 | 계획은 참고 · 실제 사용은 한 번 · 사용가능액은 잔액 탭과 일치')
    for first, last, label in [('A', 'B', '목적'), ('C', 'E', '선택월 사용계획'), ('F', 'H', '선택월 실제 사용액'), ('I', 'O', '사용가능액')]:
        span(f'{first}7:{last}7', label, {'A': 'A7', 'C': 'C7', 'F': 'F7', 'I': 'I7'}[first])
    plan = '$D$34:$O$38'
    for r, fund in enumerate(FUNDS, 8):
        span(f'A{r}:B{r}', fund, f'A{r}')
        for first, last in [('C', 'E'), ('F', 'H'), ('I', 'O')]:
            span(f'{first}{r}:{last}{r}', None, f'{first}{r}')
        value = f'INDEX({plan},MATCH($A{r},$A$34:$A$38,0),MONTH($B$3))'
        write(f'C{r}', f'=IF(COUNT({value})=0,"미입력",{value})')
        write(f'F{r}', '=' + actual[r].lstrip('='))
        write(f'I{r}', '=' + available_formula(r).lstrip('='))
    for first, last in [('A', 'B'), ('C', 'E'), ('F', 'H'), ('I', 'O')]:
        span(f'{first}13:{last}13', '합계' if first == 'A' else None, f'{first}13')
    month_plan = f'INDEX({plan},0,MONTH($B$3))'
    write('C13', f'=IF(COUNT({month_plan})=0,"미입력",SUM({month_plan}))')
    write('F13', '=SUM(F8:F12)')
    write('I13', '=IF(COUNT(I8:I12)=0,"원장 없음",SUM(I8:I12))')
    write('A14', '10월 생신 사용 239,000원은 명절비 200,000원 + 임시보관 39,000원으로 이미 충당했습니다. 경조사비에서 중복 차감하지 않습니다.')
    write('A15', '=IF(COUNTIF($R$100:$R$399,"확인*")+COUNTIF($N$430:$N$729,"확인*")>0,"확인할 사용등록·이체 연결이 있습니다. 아래 등록 자료를 확인하세요.",IF(MIN($I$8:$I$12)<0,"사용가능액이 부족한 항목이 있습니다. 부담 항목 또는 실제 추가 적립을 확인하세요.","계획액은 잔액에서 차감하지 않습니다. 사용가능액은 같은 조회월의 잔액 탭 계정 잔액과 연결됩니다."))')

    # Preserve the existing savings plan and introduce a genuinely separate
    # spending plan. No amounts are invented from past actuals or savings.
    write('A18', '2-1  적립계획 | 목적별 월 적립액 · 노란 칸 입력')
    span('A32:O32', '2-2  사용계획 | 이번 달 쓸 예정인 금액 · 비워두면 미입력', 'A18')
    span('A33:C33', '예비비 항목', 'A19')
    for month in range(1, 13):
        letter = col(month + 3)
        write(f'{letter}33', datetime(2026, month, 1), f'{letter}19')
    for r, fund in enumerate(FUNDS, 34):
        span(f'A{r}:C{r}', fund, f'A{r - 14}')
        for number in range(4, 16):
            letter = col(number)
            write(f'{letter}{r}', None, f'{letter}{r - 14}')
        height(r, 25)
    span('A39:C39', '사용계획 합계', 'A29')
    for number in range(4, 16):
        letter = col(number)
        write(f'{letter}39', f'=IF(COUNT({letter}34:{letter}38)=0,"",SUM({letter}34:{letter}38))', f'{letter}29')
    span('A40:O40', '사용계획은 참고용입니다. 비워두면 미입력, 0을 입력하면 사용계획 없음입니다. 실제 사용·잔액에는 영향을 주지 않습니다.', 'A30')
    height(32, 30)
    height(33, 30)
    height(39, 25)
    height(40, 36)
    validation = ET.SubElement(validations, Q('dataValidation'), type='whole', operator='between', sqref='D34:O38', allowBlank='1', showInputMessage='1', showErrorMessage='1', errorStyle='stop', promptTitle='선택월 사용계획', prompt='해당 월에 쓸 예정 금액을 입력하세요. 미정이면 비워두세요.', errorTitle='금액 확인', error='0 이상 1,000,000,000 이하 정수를 입력하세요.')
    ET.SubElement(validation, Q('formula1')).text = '0'
    ET.SubElement(validation, Q('formula2')).text = '1000000000'
    validations.set('count', str(len(validations)))

    write('A42', '3  이체 내역 | 사용과 실제 자금 이동을 함께 확인')
    expression = payment_formula.lstrip('=')
    write('D43', f'="확보된 결제대금비: "&IF(ISNUMBER({expression}),TEXT({expression},"#,##0")&"원","원장 없음")&" · 카드대금 출금은 새 사용이 아닙니다."')
    write('H45', '아직 옮길 돈')
    write('A58', '당월·누적 정산은 실제 재원 이동을 연결한 금액입니다. 이미 반영한 사용액은 나중에 이체할 때 다시 지출로 차감하지 않습니다.')
    height(43, 36)
    height(45, 36)
    height(58, 36)
    merges.set('count', str(len(merges)))

    # Remove stale row spans where a wide merged amount now replaces old cells.
    for r in range(7, 14):
        p.row_maps[SHEET][r].attrib.pop('spans', None)
    workbook = p.tree('xl/workbook.xml')
    sheet_index = [node.get('name') for node in workbook.find(Q('sheets'))].index(SHEET)
    for name in workbook.find(Q('definedNames')):
        if name.get('name') == '_xlnm.Print_Area' and name.get('localSheetId') == str(sheet_index):
            name.text = "'예비비'!$A$1:$O$79"
    return {'status_rows': [7, 13], 'usage_plan_inputs': 'D34:O38', 'usage_plan_total': 'D39:O39', 'display_rows': [42, 79], 'display_shift': OFFSET, 'usage_register': 'A100:Z399', 'settlement_register': 'A430:O729', 'print_area': 'A1:O79'}


def _area_cells(area):
    first, last = area.split(':') if ':' in area else (area, area)
    start = re.match(r'([A-Z]+)(\d+)', first)
    end = re.match(r'([A-Z]+)(\d+)', last)
    for row in range(int(start[2]), int(end[2]) + 1):
        for column in range(column_index_from_string(start[1]), column_index_from_string(end[1]) + 1):
            yield f'{col(column)}{row}'
