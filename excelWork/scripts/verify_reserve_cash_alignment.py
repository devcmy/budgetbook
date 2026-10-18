"""Independent reserve-allocation checks on disposable recalculation copies.

The model below sums physical ledger inputs and independently validates use and
settlement registers. It deliberately does not read the allocation helpers or
reuse their Excel formulas. Only cached formula results are imported into a
delivery XLSM; controls, VBA and original formula XML remain in the package.
"""
import argparse
from collections import Counter, defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from datetime import datetime
import gc
from hashlib import sha256
import json
import math
from pathlib import Path
import posixpath
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel, to_excel

from redesign_reserve import ExtendedPackage, FUNDS, BACK, U0, U1, S0, S1, VARIABLE
from update_20261015 import Q, RNS
from validate_excel_structure import validate_package
from verify_reserve_redesign import recalculate, assert_no_errors


BILL = '결제대금비'
WRAP = '미래에셋_WRAP형CMA'
ACCOUNTS = ['생활비', *FUNDS, BILL, '금융', '임시보관']
CORE_MONTHLY = ['C37', 'C38', 'C39', 'C40', 'C41', 'H37', 'I37',
                'H38', 'I38', 'H40', 'I40', 'H41', 'I41', 'H44',
                'I44', 'H45', 'I45', 'N53']


def read(path, formulas=False):
    return openpyxl.load_workbook(path, data_only=not formulas)


def number(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def blank(value):
    return value is None or value == ''


def date(value):
    if isinstance(value, datetime):
        return value
    if number(value) and value > 0:
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted
    return None


def same_date(a, b):
    a, b = date(a), date(b)
    return a is not None and b is not None and abs((a-b).total_seconds()) < .01


def close(actual, expected, context):
    assert number(actual) and math.isclose(actual, expected, abs_tol=1e-6), (context, actual, expected)


def boundaries(month):
    return datetime(2026, month, 1), datetime(2027, 1, 1) if month == 12 else datetime(2026, month+1, 1)


def valid_registers(w):
    """Reconstruct the source contracts without trusting OK/cache helpers."""
    reserve, ledger, raw = w['예비비'], w['잔액'], w['RAW']
    uses = {r: {c: reserve[f'{c}{r}'].value for c in 'ABCDEFGHR'}
            for r in range(U0, U1+1)}
    ids = Counter(v['A'] for v in uses.values() if not blank(v['A']))
    raw_ids = Counter(v['F'] for v in uses.values() if not blank(v['F']))
    valid_uses, invalid_uses = {}, []
    for row, v in uses.items():
        uid, fund, when, amount = v['A'], v['C'], date(v['D']), v['E']
        if blank(uid):
            continue
        if fund == '생활비':
            assert v['R'] == '생활비', (row, v['R'])
            continue
        valid = fund in FUNDS and when is not None and number(amount) and amount > 0 and ids[uid] == 1
        raw_id = v['F']
        if blank(raw_id):
            valid = valid and str(uid).startswith('U') and v['G'] == '월기록'
        elif number(raw_id) and raw_id == int(raw_id) and 2 <= raw_id <= raw.max_row:
            rr = int(raw_id)
            rv = {c: raw[f'{c}{rr}'].value for c in 'AHJKL'}
            inherited = v['G'] == '기존반영' and uid in ['U003', 'U007']
            eligible = inherited or (rv['L'] == '본인' and rv['J'] in VARIABLE)
            fixed = (uid != 'U003' or rr == 356) and (uid != 'U007' or rr == 2203)
            fixed = fixed and (v['G'] != '기존반영' or uid in ['U003', 'U007'])
            valid = valid and raw_ids[raw_id] == 1 and rv['H'] == '지출' and same_date(v['D'], rv['A'])
            valid = valid and number(rv['K']) and amount <= rv['K'] and eligible and fixed
        else:
            valid = False
        assert (v['R'] == 'OK') == bool(valid), ('independent use validity', row, v['R'], valid)
        if valid:
            valid_uses[uid] = {'row': row, 'fund': fund, 'date': when, 'amount': amount}
        else:
            invalid_uses.append(row)

    physical = []
    for row in range(64, 501):
        v = {c: ledger[f'{c}{row}'].value for c in 'ABCDEIJKLY'}
        if all(blank(v[c]) for c in 'ABCDE'):
            continue
        when, amount = date(v['A']), v['E']
        in_scope = ((v['L'] == '원본WRAP' and when is not None and when < datetime(2026, 10, 1))
                    or v['J'] == WRAP or v['K'] == WRAP
                    or (v['C'] in ACCOUNTS and v['D'] in ACCOUNTS and v['C'] != v['D']))
        valid = (in_scope and when is not None and not blank(v['B']) and not blank(v['C'])
                 and not blank(v['D']) and number(amount) and amount > 0 and v['C'] != v['D'])
        assert (v['I'] == 'OK') == bool(valid), ('physical ledger validity', row, v['I'], valid)
        close(v['Y'] or 0, amount if valid else 0, ('ledger normalized amount', row))
        if valid:
            physical.append({'row': row, 'date': when, 'source': v['C'], 'destination': v['D'], 'amount': amount})
    physical_by_row = {v['row']: v for v in physical}

    settlements = {r: {c: reserve[f'{c}{r}'].value for c in 'BCDEFHJN'}
                   for r in range(S0, S1+1) if not blank(reserve[f'B{r}'].value)}
    by_use, by_ledger = defaultdict(float), defaultdict(float)
    for v in settlements.values():
        if number(v['D']):
            by_use[v['B']] += v['D']
            if not blank(v['C']):
                by_ledger[v['C']] += v['D']
    valid_settlements, invalid_settlements = [], []
    for row, v in settlements.items():
        uid, amount, when = v['B'], v['D'], date(v['E'])
        use = valid_uses.get(uid)
        valid = use is not None and number(amount) and amount > 0 and when is not None
        if valid:
            valid = by_use[uid] <= use['amount']
            opening = (v['H'] == '기초반영' and uid in ['U001', 'U002', 'U003']
                       and blank(v['C']) and when == datetime(2026, 3, 1))
            entry = physical_by_row.get(v['C'])
            entry_ok = entry is not None and v['H'] == '원장'
            if entry_ok:
                inherited = uid == 'U008' and v['C'] == 223 and use['fund'] == '경조사비' and v['F'] == '명절비'
                family = (uid == 'U008' and v['C'] == 208 and use['fund'] == '경조사비'
                          and v['F'] == '임시보관' and by_ledger[v['C']] <= 39000)
                source_ok = v['F'] == use['fund'] or v['F'] == BILL or inherited or family
                source_category = v['F'] in FUNDS or v['F'] == BILL or family
                entry_ok = (same_date(v['E'], entry['date']) and v['F'] == entry['source']
                            and source_ok and source_category and entry['destination'] not in FUNDS
                            and by_ledger[v['C']] <= entry['amount'])
            valid = valid and (opening or entry_ok)
        assert (v['N'] == 'OK') == bool(valid), ('independent settlement validity', row, v['N'], valid)
        if valid:
            valid_settlements.append({'row': row, 'uid': uid, 'date': when, 'amount': amount, 'source': v['F']})
        else:
            invalid_settlements.append(row)
    return valid_uses, valid_settlements, physical, invalid_uses, invalid_settlements


def independent_balances(w, cutoff, registers=None):
    uses, settlements, physical, _, _ = registers or valid_registers(w)
    balances = {fund: 0 for fund in ACCOUNTS}
    for entry in physical:
        if entry['date'] >= cutoff:
            continue
        if entry['source'] in balances:
            balances[entry['source']] -= entry['amount']
        if entry['destination'] in balances:
            balances[entry['destination']] += entry['amount']
    bank_total = sum(balances.values())
    allocated = {fund: 0 for fund in FUNDS}
    for uid, use in uses.items():
        related = [s for s in settlements if s['uid'] == uid and s['date'] < cutoff]
        settled = sum(s['amount'] for s in related)
        bill_paid = sum(s['amount'] for s in related if s['source'] == BILL)
        recognized = use['amount'] if use['date'] < cutoff else 0
        allocated[use['fund']] += bill_paid + max(0, recognized-settled)
    for fund, amount in allocated.items():
        balances[fund] -= amount
        balances[BILL] += amount
    close(sum(balances.values()), bank_total, ('independent allocation preserves bank', cutoff.isoformat()))
    return balances, bank_total, allocated


def verify_month(w, month):
    reserve, ledger = w['예비비'], w['잔액']
    start, end = boundaries(month)
    registers = valid_registers(w)
    uses = registers[0]
    current, bank, allocated = independent_balances(w, end, registers)
    prior, _, _ = independent_balances(w, start, registers)
    actual = {fund: sum(u['amount'] for u in uses.values() if u['fund'] == fund and start <= u['date'] < end)
              for fund in FUNDS}
    for r, fund in enumerate(FUNDS, 8):
        close(reserve[f'F{r}'].value, actual[fund], (month, fund, 'actual use'))
        close(w[f'{month}월'][f'N{r+56}'].value, actual[fund], (month, fund, 'monthly use link'))
        planned = reserve.cell(34+FUNDS.index(fund), month+3).value
        if blank(planned):
            assert reserve[f'C{r}'].value == '미입력', (month, fund, 'blank plan is not invented')
        else:
            close(reserve[f'C{r}'].value, planned, (month, fund, 'usage plan'))
        if month < 3:
            assert reserve[f'I{r}'].value == '원장 없음', (month, fund, 'no pre-March ledger')
        else:
            close(reserve[f'I{r}'].value, current[fund], (month, fund, 'available'))
            close(reserve[f'I{r}'].value, ledger[f'G{r}'].value, (month, fund, 'balance tab alignment'))
    close(reserve['F13'].value, sum(actual.values()), (month, 'actual sum'))
    selected_plans = [reserve.cell(r, month+3).value for r in range(34, 39)]
    if all(blank(x) for x in selected_plans):
        assert reserve['C13'].value == '미입력'
    else:
        close(reserve['C13'].value, sum(x for x in selected_plans if number(x)), (month, 'plan sum'))
    if month >= 3:
        close(reserve['I13'].value, sum(current[f] for f in FUNDS), (month, 'available sum'))
        for r, fund in enumerate(ACCOUNTS, 7):
            close(ledger[f'B{r}'].value, prior[fund], (month, fund, 'opening carry'))
            close(ledger[f'G{r}'].value, current[fund], (month, fund, 'closing balance'))
            arithmetic = sum((ledger[f'{c}{r}'].value or 0) * sign
                             for c, sign in [('B', 1), ('C', 1), ('D', -1), ('E', -1), ('F', 1)])
            close(ledger[f'G{r}'].value, arithmetic, (month, fund, 'row arithmetic'))
        close(ledger['G16'].value, bank, (month, 'WRAP total'))
        close(sum(ledger[f'G{r}'].value for r in range(7, 16)), bank, (month, 'account sum'))

    latest_physical = max(v['date'] for v in registers[2])
    for m in range(3, 13):
        col = get_column_letter(m)
        ms, me = boundaries(m)
        if ms > latest_physical:
            assert all(blank(ledger[f'{col}{r}'].value) for r in range(22, 31)), (m, 'future chart blank')
            assert blank(ledger[f'{col}32'].value), (m, 'future chart total blank')
            continue
        expected, total, _ = independent_balances(w, me, registers)
        for r, fund in enumerate(ACCOUNTS, 22):
            close(ledger[f'{col}{r}'].value, expected[fund], (month, m, fund, 'monthly chart'))
        close(ledger[f'{col}32'].value, total, (month, m, 'monthly chart total'))
    return {'month': month, 'physical_wrap_total': bank if month >= 3 else None,
            'purpose_available': sum(current[f] for f in FUNDS) if month >= 3 else None,
            'actual_use': sum(actual.values()), 'allocated_to_bill': sum(allocated.values()) if month >= 3 else None,
            'invalid_use_rows': registers[3], 'invalid_settlement_rows': registers[4]}


def cache_formulas(candidate, recalculated, output, baseline):
    """Copy only the default scenario caches into the original macro package."""
    package, values, before = ExtendedPackage(candidate), read(recalculated), read(baseline)
    bools = {(s.title, c.coordinate) for s in before for row in s for c in row if isinstance(c.value, bool)}
    count = 0
    for sheet in package.sheets:
        for cell in package.sheet(sheet).iter(Q('c')):
            if cell.find(Q('f')) is None:
                continue
            addr = cell.get('r')
            v = values[sheet][addr].value
            assert values[sheet][addr].data_type != 'e', (sheet, addr, v)
            if (sheet, addr) in bools and v in [0, 1]:
                v = bool(v)
            for child in list(cell):
                if child.tag in [Q('v'), Q('is')]:
                    cell.remove(child)
            cell.attrib.pop('t', None)
            if v is not None:
                if isinstance(v, datetime):
                    v = to_excel(v)
                if isinstance(v, bool):
                    cell.set('t', 'b'); text = '1' if v else '0'
                elif number(v):
                    text = str(v)
                else:
                    cell.set('t', 'str'); text = str(v)
                ET.SubElement(cell, Q('v')).text = text
            count += 1
    package.save(output)
    values.close(); before.close()
    return count


def verify_artifact(baseline, candidate):
    before, after, bf, af = read(baseline), read(candidate), read(baseline, True), read(candidate, True)
    expected = [n for n in before.sheetnames if n != '검증']
    assert after.sheetnames == expected
    assert len(after.sheetnames) == 45
    assert len([s for s in after if s.sheet_state == 'visible']) == 28
    assert [(s.title, s.sheet_state) for s in after] == [(s.title, s.sheet_state) for s in before if s.title != '검증']
    for row in before['RAW']:
        for c in row:
            target = after['RAW'][c.coordinate]
            assert c.value == target.value, ('RAW values', c.coordinate, c.value, target.value)
            assert bf['RAW'][c.coordinate].value == af['RAW'][c.coordinate].value, ('RAW formulas', c.coordinate)
            for attr in ['font', 'fill', 'border', 'alignment', 'number_format', 'protection']:
                assert copy(getattr(bf['RAW'][c.coordinate], attr)) == copy(getattr(af['RAW'][c.coordinate], attr)), ('RAW style', c.coordinate, attr)
    for row in before['잔액'].iter_rows(min_row=64, max_row=500, max_col=25):
        for c in row:
            assert c.value == after['잔액'][c.coordinate].value, ('physical ledger values', c.coordinate)
            assert bf['잔액'][c.coordinate].value == af['잔액'][c.coordinate].value, ('physical ledger formulas', c.coordinate)
    for month in range(1, 13):
        for addr in CORE_MONTHLY:
            assert before[f'{month}월'][addr].value == after[f'{month}월'][addr].value, (month, 'monthly core', addr)
        for addr in ['W12', 'W13', 'W14', 'O149', 'O150']:
            assert before[f'{month}월계산기'][addr].value == after[f'{month}월계산기'][addr].value, (month, 'calculator', addr)
    for r in range(20, 29):
        for col in range(1, 16):
            assert before['예비비'].cell(r, col).value == after['예비비'].cell(r, col).value, ('deposit plans', r, col)
    for start, stop, columns in [(U0, U1, 'ABCDEFGH'), (S0, S1, 'BCDHI')]:
        for r in range(start, stop+1):
            for col in columns:
                assert bf['예비비'][f'{col}{r}'].value == af['예비비'][f'{col}{r}'].value, ('stable register', col, r)
    for s in af:
        for row in s:
            for c in row:
                if c.data_type == 'f':
                    assert '검증!' not in c.value and "'검증'!" not in c.value, ('removed worksheet formula dependency', s.title, c.coordinate)
    for name in af.defined_names.values():
        assert '검증!' not in name.attr_text and "'검증'!" not in name.attr_text
    assert 'I8:O8' in set(map(str, af['예비비'].merged_cells.ranges))
    assert '정산대기' not in str(after['예비비']['F7'].value)
    assert after['예비비']['F13'].value == 239000
    assert after['예비비']['I10'].value == after['잔액']['G10'].value == 450000
    assert after['예비비']['I13'].value == 2150000
    assert after['잔액']['G16'].value == 6669766
    with ZipFile(baseline) as z0, ZipFile(candidate) as z1:
        assert z1.testzip() is None
        assert z0.read('xl/vbaProject.bin') == z1.read('xl/vbaProject.bin')
        for part in z0.namelist():
            if part.startswith(('xl/drawings/', 'xl/media/', 'xl/ctrlProps/', 'xl/activeX/')):
                assert part in z1.namelist() and z0.read(part) == z1.read(part), ('macro/control asset', part)
        p0, p1 = ExtendedPackage(baseline), ExtendedPackage(candidate)
        removed = p0.sheets['검증']
        assert removed not in z1.namelist(), ('removed worksheet XML', removed)
        assert removed.encode() not in z1.read('[Content_Types].xml')
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in z1.read('[Content_Types].xml')
        assert 'xl/calcChain.xml' not in z1.namelist()
        book = p1.tree('xl/workbook.xml')
        names = book.find(Q('definedNames'))
        for node in ([] if names is None else names):
            assert '검증!' not in (node.text or '') and "'검증'!" not in (node.text or '')
        rels = p1.tree('xl/_rels/workbook.xml.rels')
        for rel in rels:
            if rel.get('TargetMode') == 'External':
                continue
            target = rel.get('Target', '')
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
            assert target != removed
    assert_no_errors(after)
    structure = validate_package(candidate)
    assert not structure['errors'], structure['errors']
    verify_month(after, 10)
    for workbook in [before, after, bf, af]:
        workbook.close()
    return {'raw_values_formulas_and_styles_preserved': True, 'physical_ledger_preserved': True,
            'monthly_financial_and_calculator_core_values_preserved': True,
            'deposit_plans_and_stable_record_ids_preserved': True,
            'vba_drawings_controls_preserved': True, 'validation_sheet_removed_not_hidden': True,
            'removed_sheet_dependencies': 0, 'visible_sheets': 28, 'total_sheets': 45,
            'xml_structure': structure, 'native_excel_opened': False}


def scenarios():
    """Return fixture edits and additional assertions; no real source is edited."""
    def query(p, month):
        for name in ['예비비', '잔액']:
            p.write(name, 'B3', datetime(2026, month, 1))

    def use(p, row=180, uid='U180', fund='경조사비', when=datetime(2026, 10, 17), amount=210000):
        for col, value in [('A', uid), ('B', '검산 전용 카드 사용'), ('C', fund), ('D', when),
                           ('E', amount), ('F', None), ('G', '월기록'), ('H', '폐기할 검산 사본')]:
            p.write('예비비', f'{col}{row}', value)

    def payment(p, amount, when=datetime(2026, 10, 18), source=BILL, row=249, linkrow=440, uid='U180'):
        for col, value in [('A', when), ('B', '이체'), ('C', source), ('D', '외부'), ('E', amount),
                           ('F', '검산 전용 카드대금 이체'), ('G', '실제 원장에 없는 폐기용 입력'),
                           ('J', WRAP), ('K', '검산 전용 카드계좌'), ('L', 'WRAP계좌이체')]:
            p.write('잔액', f'{col}{row}', value)
        for col, value in [('B', uid), ('C', row), ('D', amount)]:
            p.write('예비비', f'{col}{linkrow}', value)

    def base(p):
        query(p, 10); use(p)

    def check(w, available=240000, bill=210000, bank=6669766, actual=449000):
        close(w['예비비']['I10'].value, available, 'scenario purpose available')
        close(w['잔액']['G13'].value, bill, 'scenario reserved bill account')
        close(w['잔액']['G16'].value, bank, 'scenario physical total')
        close(w['예비비']['F10'].value, actual, 'scenario actual including preserved birthday')

    def compose(*edits):
        def apply(p):
            for edit in edits:
                edit(p)
        return apply

    cases = []
    def add(name, edit, checkfn=lambda w: None, month=10, warnings=False):
        cases.append((name, month, edit, checkfn, warnings))
    add('card_use_reserves_bill_without_bank_outflow', base, check)
    add('partial_bill_payment_no_second_purpose_deduction', compose(base, lambda p: payment(p, 70000)),
        lambda w: check(w, bill=140000, bank=6599766))
    add('full_bill_payment_no_second_purpose_deduction', compose(base, lambda p: payment(p, 210000)),
        lambda w: check(w, bill=0, bank=6459766))
    add('original_purpose_payment_also_deducts_once', compose(base, lambda p: payment(p, 210000, source='경조사비')),
        lambda w: check(w, bill=0, bank=6459766))
    add('next_month_bill_payment_cannot_change_october', compose(base, lambda p: payment(p, 210000, datetime(2026, 11, 5))), check)
    add('next_month_payment_carries_correct_balance',
        compose(base, lambda p: payment(p, 210000, datetime(2026, 11, 5)), lambda p: query(p, 11)),
        lambda w: check(w, bill=0, bank=6459766, actual=0), month=11)
    add('bill_prepayment_before_use_reclassifies_funding',
        compose(base, lambda p: payment(p, 210000, datetime(2026, 9, 25)), lambda p: query(p, 9)),
        lambda w: check(w, bill=0, bank=8328603, actual=50000), month=9)
    add('bill_prepayment_then_use_no_second_deduction',
        compose(base, lambda p: payment(p, 210000, datetime(2026, 9, 25))),
        lambda w: check(w, bill=0, bank=6459766))
    add('future_use_does_not_reduce_current_balance',
        compose(lambda p: query(p, 10), lambda p: use(p, when=datetime(2026, 11, 5))),
        lambda w: check(w, available=450000, bill=0, actual=239000))
    add('duplicate_use_ids_fail_closed', compose(base, lambda p: use(p, row=181)),
        lambda w: check(w, available=450000, bill=0, actual=239000), warnings=True)
    add('duplicate_settlement_overallocation_fails_closed',
        compose(base, lambda p: payment(p, 210000), lambda p: payment(p, 210000, linkrow=441)),
        lambda w: check(w, bill=0, bank=6459766), warnings=True)
    add('wrong_funding_category_settlement_fails_closed',
        compose(base, lambda p: payment(p, 210000, source='여행비')),
        lambda w: check(w, bank=6459766), warnings=True)
    add('usage_plan_is_display_only_not_cash_or_actual',
        compose(lambda p: query(p, 10), lambda p: p.write('예비비', 'M36', 350000)),
        lambda w: (check(w, available=450000, bill=0, actual=239000), close(w['예비비']['C10'].value, 350000, '350k plan display')))
    add('explicit_zero_plan_is_not_unentered',
        compose(lambda p: query(p, 10), lambda p: p.write('예비비', 'M36', 0)),
        lambda w: (check(w, available=450000, bill=0, actual=239000), close(w['예비비']['C10'].value, 0, 'zero plan display')))
    add('350k_plan_210k_actual_only_actual_reserved',
        compose(lambda p: query(p, 10), lambda p: use(p, fund='여행비'), lambda p: p.write('예비비', 'M35', 350000)),
        lambda w: (close(w['예비비']['C9'].value, 350000, 'travel plan'),
                   close(w['예비비']['F9'].value, 210000, 'travel actual'),
                   close(w['예비비']['I9'].value, 1190000, 'travel available'),
                   close(w['잔액']['G9'].value, 1190000, 'travel balance account'),
                   close(w['잔액']['G13'].value, 210000, 'bill reserved'),
                   close(w['잔액']['G16'].value, 6669766, 'unchanged bank')))
    add('mixed_original_and_bill_partial_payments',
        compose(base, lambda p: payment(p, 60000, source='경조사비'),
                lambda p: payment(p, 70000, row=250, linkrow=441)),
        lambda w: check(w, bill=80000, bank=6539766))
    def internal_reclassification(p):
        base(p)
        payment(p, 210000, source='경조사비')
        for c, v in [('D', BILL), ('K', WRAP), ('L', 'WRAP내계정이동')]:
            p.write('잔액', f'{c}249', v)
    add('manual_internal_transfer_replaces_auto_allocation', internal_reclassification, check)
    add('cash_payment_after_manual_internal_transfer_not_new_use',
        compose(internal_reclassification, lambda p: payment(p, 210000, row=250, linkrow=441),
                lambda p: p.write('예비비', 'B441', None)),
        lambda w: check(w, bill=0, bank=6459766))
    add('month_end_timestamp_use_is_included',
        compose(lambda p: query(p, 10), lambda p: use(p, when=datetime(2026, 10, 31, 23, 59, 59))), check)
    add('month_end_timestamp_deposit_reaches_top_and_chart',
        compose(base, lambda p: payment(p, 10000, datetime(2026, 10, 31, 23, 59, 59), source='외부'),
                lambda p: p.write('잔액', 'D249', '경조사비'), lambda p: p.write('예비비', 'B440', None)),
        lambda w: check(w, available=250000, bank=6679766))
    add('prior_month_outstanding_reduces_following_month_available',
        compose(base, lambda p: query(p, 11)), lambda w: check(w, actual=0), month=11)
    add('insufficient_fund_stays_negative_not_hidden',
        compose(lambda p: query(p, 10), lambda p: use(p, fund='명절비')),
        lambda w: close(w['예비비']['I8'].value, -210000, 'insufficient allocation remains visible'))
    return cases


def run(baseline, candidate, output, workdir, report, workers=2, reuse=False):
    root = Path(workdir); root.mkdir(parents=True, exist_ok=True)
    fingerprint = sha256(Path(candidate).read_bytes()).hexdigest()
    manifest = root/'recalculation_manifest.json'
    if reuse:
        assert manifest.is_file() and json.loads(manifest.read_text())['candidate_sha256'] == fingerprint, 'Do not reuse calculations from a different candidate.'
    else:
        assert not manifest.exists(), 'Use a new workdir, or explicitly reuse this exact candidate.'
        manifest.write_text(json.dumps({'candidate_sha256': fingerprint})+'\n')
    jobs = []
    for month in range(1, 13):
        def edit(p, month=month):
            for name in ['예비비', '잔액']:
                p.write(name, 'B3', datetime(2026, month, 1))
        jobs.append((f'month_{month:02d}', month, edit, lambda w: None, False, True))
    jobs += [(name, month, edit, check, warning, False) for name, month, edit, check, warning in scenarios()]

    def execute(job):
        name, month, edit, check, warnings, is_month = job
        recalculated = root/name/'out/input.xlsx'
        if not (reuse and recalculated.is_file()):
            recalculated = recalculate(candidate, root/name, edit)
        w = read(recalculated)
        assert_no_errors(w)
        result = verify_month(w, month)
        check(w)
        if not warnings:
            assert not result['invalid_use_rows'] and not result['invalid_settlement_rows'], (name, result)
        else:
            assert result['invalid_use_rows'] or result['invalid_settlement_rows'], (name, 'warning was not surfaced')
        w.close(); del w; gc.collect()
        row = {'name': name, 'passed': True, 'expected_registration_warning': warnings, **result}
        print(json.dumps(row, ensure_ascii=False), flush=True)
        return row, recalculated, is_month

    months, cases, default = [], [], None
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(execute, job) for job in jobs]
        for future in as_completed(futures):
            result, recalculated, is_month = future.result()
            (months if is_month else cases).append(result)
            if result['name'] == 'month_10':
                default = recalculated
    assert default is not None
    count = cache_formulas(candidate, default, output, baseline)
    checks = verify_artifact(baseline, output)
    data = {'ui_updated_at': '2026-10-19', 'as_of': '2026-10-18', 'default_month': '2026-10', 'month_tests': sorted(months, key=lambda x: x['month']),
            'scenario_tests': sorted(cases, key=lambda x: x['name']), 'checks': checks, 'cached_formulas': count,
            'formula_errors': 0, 'workbook_sha256': sha256(Path(output).read_bytes()).hexdigest(),
            'baseline_sha256': sha256(Path(baseline).read_bytes()).hexdigest(), 'native_excel_opened': False,
            'recalculation': 'Isolated macro-free XLSX copies in LibreOffice; only default-scenario value caches imported into original XLSM package.',
            'independent_model': 'Python validation of source ledger and registers, then per-use paid-from-bill plus outstanding use allocation; no hidden allocation helper formulas reused.'}
    Path(report).write_text(json.dumps(data, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({'verified': True, 'months': len(months), 'scenarios': len(cases), 'sha256': data['workbook_sha256']}), flush=True)
    return data


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    for field in ['baseline', 'candidate', 'output', 'workdir', 'report']:
        parser.add_argument('--'+field, required=True)
    parser.add_argument('--workers', type=int, default=2, choices=[1, 2])
    parser.add_argument('--reuse-recalculations', action='store_true')
    args = parser.parse_args()
    run(args.baseline, args.candidate, args.output, args.workdir, args.report, args.workers, args.reuse_recalculations)
