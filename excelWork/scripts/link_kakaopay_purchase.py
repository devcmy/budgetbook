"""Record the user-confirmed charge/purchase link, without changing any money.

Only four text inputs and three deterministic text-display caches are patched.
The purchase is not added to W/X, which exclusively reconcile WRAP cash rows.
"""
import argparse
from copy import deepcopy
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl

from balance_selected_ledger import check_view
from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from validate_excel_structure import validate_package


LINK = 'KP20261017-01'


def apply(baseline, output, previous_audit, output_audit, report):
    assert Path(baseline).resolve() != Path(output).resolve()
    b = openpyxl.load_workbook(baseline, data_only=True)
    r, s = b['RAW'], b['잔액']
    assert (r['C2474'].value, r['F2474'].value, r['G2474'].value, r['H2474'].value) == (
        '카카오페이머니', '플라잉 더치', 17000, '지출')
    assert (r['C2475'].value, r['D2475'].value, r['G2475'].value, r['H2475'].value) == (
        '미래에셋_WRAP형CMA', '카카오페이머니', 15700, '이체지출')
    assert r['G2376'].value == -1300 and r['C2376'].value == '카카오페이머니'
    assert s['W248'].value == 2475 and s['X248'].value is None
    assert s['N38'].value + 63 == s['AA503'].value + 63 == 248
    assert LINK not in r['M2474'].value, 'Do not apply this migration twice.'
    p = ExtendedPackage(baseline)
    texts = {
        ('RAW', 'M2474'): f'{LINK} | 10/18 첨부 RAW 13행. 플라잉 더치 구매 17,000원 ↔ 충전 RAW 2475행 15,700원(잔액 원장 248행). 사용자 확인으로 연결. 소비는 구매에 1회, WRAP 현금 유출은 충전에 1회 반영. 차액 1,300원의 실제 재원은 미확정.',
        ('RAW', 'M2475'): f'{LINK} | 10/18 첨부 RAW 14행. 플라잉 더치 구매 RAW 2474행 17,000원에 연결된 충전 15,700원(사용자 확인). WRAP 원장 248행에서 현금 차감, 충전 자체는 소비 재합산 제외. 원본 입력 시각은 구매보다 늦으며 그대로 보존.',
        ('잔액', 'F248'): '카카오페이 충전 · 플라잉 더치',
        ('잔액', 'G248'): f'{LINK} · 플라잉 더치 구매 연결(사용자 확인). 충전 RAW 2475행 15,700원 ↔ 구매 RAW 2474행 17,000원. 차액 1,300원 재원 미확정. 현금은 충전에서만 차감, 소비는 구매에서만 집계.',
    }
    for (sheet, addr), value in texts.items():
        assert p.cell(sheet, addr).find(Q('f')) is None
        p.write(sheet, addr, value)
    name, memo = texts[('잔액', 'F248')], texts[('잔액', 'G248')]
    caches = {'K38': name, 'F503': name, 'G503': memo[:40] + ('…' if len(memo) > 40 else '')}
    for addr, value in caches.items():
        cell = p.cell('잔액', addr)
        formula = cell.find(Q('f'))
        assert formula is not None
        if addr == 'K38':
            assert 'INDEX($F$64:$F$500,$N38)' in formula.text
        elif addr == 'F503':
            assert 'INDEX($F$64:$F$500,MAX(1,$AA503))' in formula.text
        else:
            assert 'LEFT(INDEX($G$64:$G$500,MAX(1,$AA503)),40)' in formula.text
        for child in list(cell):
            if child.tag in [Q('v'), Q('is')]:
                cell.remove(child)
        cell.set('t', 'str')
        ET.SubElement(cell, Q('v')).text = value
    p.save(output)
    structure = validate_package(output)
    assert not structure['errors'], structure
    allowed = {'RAW': {'M2474', 'M2475'}, '잔액': {'F248', 'G248', *caches}}
    with ZipFile(baseline) as old, ZipFile(output) as new:
        assert old.namelist() == new.namelist() and new.testzip() is None
        changed = [n for n in old.namelist() if old.read(n) != new.read(n)]
        assert set(changed) == {p.sheets['RAW'], p.sheets['잔액']}, changed
        for sheet, addresses in allowed.items():
            before = ET.fromstring(old.read(p.sheets[sheet]))
            after = ET.fromstring(new.read(p.sheets[sheet]))
            old_cells = {c.get('r'): c for c in before.iter(Q('c'))}
            new_cells = {c.get('r'): c for c in after.iter(Q('c'))}
            assert old_cells.keys() == new_cells.keys()
            actual = {a for a in old_cells if ET.tostring(old_cells[a]) != ET.tostring(new_cells[a])}
            assert actual == addresses, (sheet, actual)
            for addr in addresses:
                old_formula, new_formula = old_cells[addr].find(Q('f')), new_cells[addr].find(Q('f'))
                assert (None if old_formula is None else ET.tostring(old_formula)) == (
                    None if new_formula is None else ET.tostring(new_formula))
                new_cells[addr].getparent().replace(new_cells[addr], deepcopy(old_cells[addr]))
            assert ET.tostring(before) == ET.tostring(after), sheet
    a = openpyxl.load_workbook(output, data_only=True)
    for sheet in b:
        for row in sheet:
            for cell in row:
                if cell.coordinate not in allowed.get(sheet.title, set()):
                    assert cell.value == a[sheet.title][cell.coordinate].value, (sheet.title, cell.coordinate)
    view = check_view(a, 10)
    assert a['잔액']['F503'].value == name and a['잔액']['G503'].value == caches['G503']
    assert [a['잔액'][c].value for c in ['G7', 'G15', 'G16']] == [4519766, 0, 6669766]
    assert a['10월']['I33'].value == 1476900 and a['10월계산기']['W12'].value == 1501000
    assert a['예비비']['L13'].value == 239000
    assert not [c.coordinate for row in a['검증'] for c in row if c.value is False]
    assert not [(sheet.title, c.coordinate) for sheet in a for row in sheet for c in row if c.data_type == 'e']
    result = {
        'date': '2026-10-18', 'link_id': LINK, 'status': 'user_confirmed',
        'confirmation': '그게 결국 카카오페이 충전한게 더치커피구매로 이어진글 체크해야해',
        'purchase': {'sheet': 'RAW', 'row': 2474, 'merchant': '플라잉 더치', 'amount': 17000,
                     'recorded_at': r['A2474'].value.isoformat(), 'consumption_count': 1, 'wrap_deduction': 0},
        'funding': {'sheet': 'RAW', 'row': 2475, 'ledger_row': 248, 'amount': 15700,
                    'recorded_at': r['A2475'].value.isoformat(), 'wrap_deduction': 15700, 'additional_consumption': 0},
        'amount_difference': 1300, 'difference_funding_confirmed': False,
        'possible_related_refund': {'raw_row': 2376, 'date': '2026-10-06', 'amount': -1300,
                                    'status': 'amount_consistent_only_not_verified_funding'},
        'time_note': 'Source entry timestamps put the charge after the purchase; preserved. Link confirmed by user, not inferred from chronology.',
        'financial_values_and_all_formulas_unchanged': True, 'all_other_zip_parts_byte_identical': True,
        'changed_zip_parts': changed, 'changed_cells': {s: sorted(v) for s, v in allowed.items()},
        'deterministic_text_cache_updates': list(caches), 'selected_month_view': view,
        'structure_validation': structure, 'no_full_recalculation': True, 'native_excel_opened': False,
        'baseline_sha256': sha256(Path(baseline).read_bytes()).hexdigest(),
        'workbook_sha256': sha256(Path(output).read_bytes()).hexdigest(),
    }
    Path(report).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    audit = json.loads(Path(previous_audit).read_text())
    assert audit['workbook_sha256'] == result['baseline_sha256']
    audit['workbook_sha256'] = result['workbook_sha256']
    audit['delta_20261018']['scope_note'] = 'RAW 반영 직후의 검증 이력. 금액은 유지하며 현재 거래 연결 표시·파일 해시는 kakaopay_purchase_link와 최상위 필드를 따릅니다.'
    audit['kakaopay_purchase_link'] = {**result, 'details': Path(report).name,
        'guide': '2026-10-18_카카오페이구매연결_검증.md', 'previous_version': 'history/2026-10-18/04'}
    Path(output_audit).write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'linked': LINK, 'only_changed_parts': changed,
                      'all_financial_values_and_formulas_unchanged': True,
                      'workbook_sha256': result['workbook_sha256']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['baseline', 'output', 'previous-audit', 'output-audit', 'report']:
        parser.add_argument('--' + name, required=True)
    args = parser.parse_args()
    apply(args.baseline, args.output, args.previous_audit, args.output_audit, args.report)
