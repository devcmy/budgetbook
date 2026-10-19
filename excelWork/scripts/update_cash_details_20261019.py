"""Name a recovered advance and display an already-counted prepaid purchase.

Patch the XLSM XML package; never save the delivery workbook through a
spreadsheet library. Recalculate a disposable macro-free copy and import only
changed formula caches. Per user confirmation, leave the combined original
purchase unchanged and preserve the -1300 recovery as a net expense offset.
"""
import argparse
from collections import Counter
from datetime import datetime
import gc
from hashlib import sha256
import json
import math
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import to_excel

from redesign_reserve import ExtendedPackage
from update_20261015 import CATEGORIES, Q
from validate_excel_structure import validate_package
from verify_reserve_cash_alignment import verify_month
from verify_reserve_redesign import assert_no_errors, recalculate


LINK = 'KP20261017-01'
ACTUAL, TETRIS = 1476900, 1501000


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def reimbursement_note():
    return (f'{LINK} | 10/19 사용자 확인: 양배추 대리구매 대납금 1,300원은 이미 회수했고 '
            '플라잉 더치 구매에 사용. 함께 산 원구매 금액은 수정하지 않으며, 기존 소비 상계 '
            '-1,300원을 유지. 신규 소득·추가 미회수 대납 아님.')


def paint_order():
    return [f'{get_column_letter(c)}{r}' for r0 in range(15, 145, 5)
            for c0 in range(2, 22, 2) for c in range(c0, c0 + 2)
            for r in range(r0, r0 + 5)]


def prepare(baseline, candidate):
    w = openpyxl.load_workbook(baseline, data_only=True)
    raw, month, calc = w['RAW'], w['10월'], w['10월계산기']
    assert raw['G2376'].value == raw['K2376'].value == -1300
    assert raw['I2376'].value == '양배추 대리구매'
    assert raw['F2474'].value == '플라잉 더치' and raw['K2474'].value == 17000
    assert raw['G2475'].value == 15700 and raw['H2475'].value == '이체지출'
    assert month['C53'].value is None and month['D53'].value is None
    assert month['I33'].value == 1476900 and calc['W12'].value == 1501000
    p = ExtendedPackage(baseline)
    edits = {
        ('RAW', 'F2376'): '양배추 대리구매',
        ('RAW', 'M2376'): reimbursement_note(),
        ('RAW', 'M2474'): f'{LINK} | 플라잉 더치 구매 17,000원 = 충전 RAW 2475행 15,700원 + 양배추 대납 회수 RAW 2376행 1,300원(10/19 사용자 확인). 소비는 구매에 1회, WRAP 출금은 충전에 1회. 10월 추가 실지출 표에는 이미 집계된 거래로 표시.',
        ('RAW', 'M2475'): f'{LINK} | 플라잉 더치 구매 RAW 2474행 17,000원의 재원: 이번 충전 15,700원 + 이미 회수한 양배추 대납금 1,300원(RAW 2376행). WRAP 원장 248행에서 충전만 차감하며 소비 재합산 제외. 원본 입력 시각 보존.',
        ('잔액', 'G248'): f'{LINK} · 플라잉 더치 구매 연결(사용자 확인). 충전 RAW 2475행 15,700원 + 양배추 대납 회수 RAW 2376행 1,300원 = 구매 RAW 2474행 17,000원. 현금은 충전에서만 차감, 소비는 구매에서만 집계.',
        ('10월', 'B53'): '선불금',
        ('10월', 'C53'): '플라잉 더치',
        ('10월', 'I53'): '카카오페이머니: 달력에 1회 집계, 추가 합산 0원.',
    }
    for (sheet, addr), value in edits.items():
        assert p.cell(sheet, addr).find(Q('f')) is None
        p.write(sheet, addr, value)
    for addr, formula in {
        'A53': '=INT(RAW!A2474)', 'D53': '=RAW!K2474', 'G53': '=RAW!K2474',
    }.items():
        p.write('10월', addr, formula, True)
    # Keep the F/H shared-formula groups, the right-hand savings totals and all
    # other row contents intact; copy only body styles for these new labels.
    for col in ['C', 'H', 'I']:
        p.cell('10월', f'{col}53').set('s', p.cell('10월', f'{col}52').get('s'))
    p.row_maps['10월'][53].set('ht', '44')
    p.row_maps['10월'][53].set('customHeight', '1')
    p.trees.pop('xl/workbook.xml')  # Workbook metadata is unchanged.
    p.save(candidate)
    w.close()
    gc.collect()
    return {'changed_input_cells': {name: sorted(a for s, a in edits if s == name)
                                    for name in {s for s, _ in edits}},
            'original_purchase_allocation_unchanged_per_user': True,
            'existing_recovery_expense_offset_preserved': True}


def import_changed_caches(candidate, calculated, output):
    p = ExtendedPackage(candidate)
    w = openpyxl.load_workbook(calculated, data_only=True)
    changed = {}
    for name in p.sheets:
        updates = []
        for cell in p.sheet(name).iter(Q('c')):
            if cell.find(Q('f')) is None:
                continue
            addr = cell.get('r'); source = w[name][addr]
            assert source.data_type != 'e', (name, addr, source.value)
            value = source.value
            if isinstance(value, datetime): value = to_excel(value)
            old = cell.find(Q('v')); old_text = None if old is None else old.text
            kind = cell.get('t')
            if kind == 'b' and value in [0, 1]: value = bool(value)
            if value is None:
                same = old_text is None
            elif isinstance(value, bool):
                same = kind == 'b' and old_text == ('1' if value else '0')
            elif isinstance(value, (int, float)):
                same = old_text is not None and kind not in ['str', 's', 'inlineStr', 'b'] and math.isclose(float(old_text), value, rel_tol=0, abs_tol=1e-7)
            else:
                same = kind == 'str' and old_text == str(value)
            if same: continue
            for node in list(cell):
                if node.tag in [Q('v'), Q('is')]: cell.remove(node)
            cell.attrib.pop('t', None)
            if value is not None:
                if isinstance(value, bool):
                    cell.set('t', 'b'); text = '1' if value else '0'
                elif isinstance(value, (int, float)): text = str(value)
                else: cell.set('t', 'str'); text = str(value)
                ET.SubElement(cell, Q('v')).text = text
            updates.append(addr)
        if updates: changed[name] = updates
        else: p.trees.pop(p.sheets[name])
    p.trees.pop('xl/workbook.xml')
    p.save(output)
    w.close(); gc.collect()
    return changed


def verify(baseline, final, metadata):
    w = openpyxl.load_workbook(final, data_only=True)
    raw, month, calc, ledger = [w[n] for n in ['RAW', '10월', '10월계산기', '잔액']]
    assert raw['F2376'].value == '양배추 대리구매'
    assert raw['G2376'].value == -1300 and raw['H2376'].value == '지출'
    assert raw['J2376'].value == '변동비' and raw['K2376'].value == -1300
    assert raw['L2376'].value == '본인'
    records = [r for r in range(2323, 2487) if raw[f'H{r}'].value == '지출'
               and raw[f'L{r}'].value == '본인' and raw[f'J{r}'].value in CATEGORIES]
    assert len(records) == 92 and 2376 in records and 2474 in records
    actual = sum(raw[f'K{r}'].value for r in records)
    tetris = sum(math.ceil(raw[f'K{r}'].value/1000)*1000 for r in records)
    assert actual == ACTUAL == month['I33'].value == month['H38'].value
    assert tetris == TETRIS == calc['W12'].value == month['I38'].value
    dates = []
    for dr in [3, 8, 13, 18, 23, 28]:
        for c in range(2, 9):
            dt = month.cell(dr, c).value
            if not isinstance(dt, datetime) or dt.month != 10: continue
            rows = [r for r in records if raw[f'A{r}'].value.date() == dt.date()]
            assert (month.cell(dr+2, c).value or 0) == sum(raw[f'K{r}'].value for r in rows), dt
            assert (month.cell(dr+1, c).value or 0) == sum(math.ceil(raw[f'K{r}'].value/1000)*1000 for r in rows), dt
            dates.append(dt.date().isoformat())
    assert len(dates) == 31
    assert month['C53'].value == '플라잉 더치'
    assert month['A53'].value == datetime(2026, 10, 17)
    assert [month[f'{c}53'].value for c in ['D', 'F', 'G', 'H']] == [17000, 17000, 17000, 0]
    assert month['I62'].value == month['H39'].value == 0
    assert month['C52'].value == '두끼떡볶이' and month['H52'].value == 0
    assert month['E80'].value == 103600
    assert all(month[a].value == 0 for a in ['E83', 'E84', 'E85'])
    assert all(bool(month[f'U{r}'].value) for r in [28, 29, 30, 31])
    assert raw.max_row == 2486 and ledger['E248'].value == 15700
    assert ledger['G7'].value == 4519766 and ledger['G16'].value == 6669766
    assert ledger['G13'].value == 0 and w['예비비']['I10'].value == 450000
    assert w['예비비']['F13'].value == 239000
    palette = {calc[f'B{r}'].fill.fgColor.rgb: r for r in range(3, 11)}
    counts = Counter()
    painted = []
    for addr in paint_order():
        cell = calc[addr]; solid = cell.fill.patternType == 'solid'; painted.append(solid)
        if solid: counts[palette[cell.fill.fgColor.rgb]] += 1
    assert painted == [True]*1501 + [False]*(2600-1501)
    for r in range(3, 11): assert counts[r] == calc[f'C{r}'].value == calc[f'AA{r}'].value
    assert counts[3] == 323 and sum(counts.values()) == 1501
    reserve = verify_month(w, 10)
    assert_no_errors(w)
    w.close(); gc.collect()
    # Original dates, assets, amounts and source types are immutable. Only
    # F2376 is a user-authorized correction to an original descriptive field.
    allowed_raw = {'F2376', 'M2376', 'M2474', 'M2475'}
    old, new = ExtendedPackage(baseline), ExtendedPackage(final)
    for sheet in old.sheets:
        before = {c.get('r'): c for c in old.sheet(sheet).iter(Q('c'))}
        after = {c.get('r'): c for c in new.sheet(sheet).iter(Q('c'))}
        assert before.keys() == after.keys(), ('cell identities', sheet)
        for addr in before:
            a, b = before[addr], after[addr]
            if sheet == 'RAW' and addr not in allowed_raw:
                assert ET.tostring(a) == ET.tostring(b), ('RAW preservation', addr)
            if sheet == '잔액' and 64 <= int(''.join(x for x in addr if x.isdigit())) <= 500 and addr != 'G248':
                assert ET.tostring(a) == ET.tostring(b), ('physical ledger preservation', addr)
            if sheet == '잔액' and addr != 'G248':
                assert ET.tostring(a) == ET.tostring(b), ('balance view preservation', addr)
            if sheet == '10월' and addr not in {f'{c}53' for c in 'ABCDGHI'} | {'F53'}:
                assert ET.tostring(a) == ET.tostring(b), ('monthly existing cells', addr)
            if sheet not in ['RAW', '10월', '잔액']:
                assert ET.tostring(a) == ET.tostring(b), ('unrelated worksheet cell', sheet, addr)
    assert old.sheets == new.sheets and len(new.sheets) == 45 and '검증' not in new.sheets
    with ZipFile(baseline) as b, ZipFile(final) as a:
        assert b.namelist() == a.namelist() and a.testzip() is None
        changed_parts = [x for x in b.namelist() if b.read(x) != a.read(x)]
        allowed_parts = {new.sheets[s] for s in ['RAW', '10월', '잔액']}
        assert set(changed_parts) <= allowed_parts, changed_parts
        for name in b.namelist():
            if name.startswith(('xl/drawings/', 'xl/media/', 'xl/ctrlProps/', 'xl/activeX/')) or name == 'xl/vbaProject.bin':
                assert b.read(name) == a.read(name), ('VBA/control preservation', name)
    structure = validate_package(final)
    assert not structure['errors'], structure
    return {'date': '2026-10-19', 'data_as_of': '2026-10-18',
            'baseline_sha256': digest(baseline), 'workbook_sha256': digest(final),
            'original_raw_amount_and_kind_preserved': True,
            'recovered_advance': {'row': 2376, 'name': '양배추 대리구매', 'source_amount': -1300,
                                  'existing_net_expense_offset': -1300, 'recovered': True,
                                  'combined_original_purchase_unchanged_per_user': True,
                                  'used_for_purchase_row': 2474},
            'purchase_funding': {'purchase': 17000, 'wrap_topup': 15700, 'recovered_advance': 1300, 'confirmed': True},
            'additional_expense_display': {'row': 53, 'actual': 17000, 'already_counted': 17000, 'additional': 0},
            'calendar_basis_unchanged': 'personal_variable_consumption_all_payment_methods',
            'daily_calendar_checks': len(dates), 'personal_transactions': 92,
            'personal_actual': ACTUAL, 'personal_tetris': TETRIS,
            'food_blocks': counts[3], 'total_blocks': sum(counts.values()),
            'wrap_total': 6669766, 'living_balance': 4519766,
            'outstanding_advances': 103600, 'birthday_actual_use': 239000,
            'reserve_alignment': reserve, 'changed_package_parts': changed_parts,
            'raw_and_physical_ledger_preservation_passed': True,
            'unrelated_sheet_cells_preserved': True, 'vba_and_controls_byte_identical': True,
            'structure_validation': structure, 'formula_errors': 0,
            'recalculation': 'disposable_macro_free_LibreOffice_copy_changed_caches_only',
            'native_excel_opened': False, 'sharepoint_file_updated': False, **metadata}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for arg in ['baseline', 'workdir', 'output', 'report']:
        parser.add_argument('--'+arg, required=True)
    a = parser.parse_args()
    folder = Path(a.workdir); candidate = folder/'candidate.xlsm'
    metadata = prepare(a.baseline, candidate)
    print('Prepared XML changes; recalculating disposable copy.', flush=True)
    calculated = recalculate(candidate, folder/'recalc')
    metadata['changed_formula_caches'] = import_changed_caches(candidate, calculated, a.output)
    result = verify(a.baseline, a.output, metadata)
    Path(a.report).write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps(result, ensure_ascii=False), flush=True)
