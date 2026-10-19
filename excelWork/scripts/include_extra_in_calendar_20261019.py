"""Route all September–December extra expenses through the dated calendar once.

F is net expense, G is its existing base-calendar coverage, and H=F-G is
connected to the calendar. No raw record, cash ledger or reserve is changed.
Preserve old calendar bases, old calculator paint (including September gaps),
and both fixed-cost perspectives. Recalculate only disposable macro-free copies.
"""
import argparse
from collections import Counter
from copy import deepcopy
from datetime import datetime
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl

from monthly_usage_summary import MONTHS, advance_total, digest, fixed_total
from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from update_cash_details_20261019 import import_changed_caches, paint_order
from validate_excel_structure import validate_package
from verify_reserve_redesign import assert_no_errors, recalculate

PALETTES = {(9, 53): 6, (9, 54): 8, (10, 54): 4}


def rounded(amount):
    return math.ceil((amount or 0) / 1000) * 1000


def extra_delta(net, covered):
    return (net or 0) - (covered or 0), rounded(net) - rounded(covered)


def usage_totals(fixed, calendar_with_extras, advances):
    """Extras are already in the calendar; both visible usage totals agree."""
    total = fixed + calendar_with_extras + advances
    return total, total


def grid(sheet, month):
    for r in [3, 8, 13, 18, 23, 28]:
        for c in range(2, 9):
            date = sheet.cell(r, c).value
            if isinstance(date, datetime) and date.month == month:
                yield date, sheet.cell(r, c).coordinate, sheet.cell(r+1, c).coordinate, sheet.cell(r+2, c).coordinate


def daily_formula(base, date_cell, source):
    delta = f'SUMIFS(${source}$52:${source}$61,$A$52:$A$61,">="&{date_cell},$A$52:$A$61,"<"&{date_cell}+1)'
    if base is None:
        count = f'COUNTIFS($A$52:$A$61,">="&{date_cell},$A$52:$A$61,"<"&{date_cell}+1,$D$52:$D$61,"<>")'
        return f'=IF({count}=0,"",{delta})'
    expression = base.lstrip('=') if isinstance(base, str) and base.startswith('=') else str(base)
    return f'=({expression})+{delta}'


def variance_formula(tetris, actual):
    # Formula-produced empty strings must not count as entered transactions.
    return f'=IF(COUNT({tetris},{actual})=0,"",SUM({tetris})-SUM({actual}))'


def prepare(baseline, candidate):
    w = openpyxl.load_workbook(baseline)
    d = openpyxl.load_workbook(baseline, data_only=True)
    p = ExtendedPackage(baseline)
    xfs = p.tree('xl/styles.xml').find(Q('cellXfs'))
    style_cache = {ET.tostring(x): i for i, x in enumerate(xfs)}
    metadata = {}
    for month in MONTHS:
        name = f'{month}월'; s = w[name]; values = d[name]
        assert values['AB51'].value is None, 'Already migrated; do not add the same delta twice.'
        extra_rows, paint = [], []
        for r in range(52, 62):
            if values[f'D{r}'].value is not None:
                date, net, covered = [values[f'{c}{r}'].value for c in ['A', 'F', 'G']]
                assert isinstance(date, datetime) and date.year == 2026 and date.month == month
                actual, tetris = extra_delta(net, covered)
                extra_rows.append(dict(row=r, date=date.isoformat(), name=values[f'C{r}'].value,
                                       net=net, covered=covered or 0, added_actual=actual, added_tetris=tetris))
                if tetris:
                    assert tetris > 0 and (month, r) in PALETTES
                    paint.append((PALETTES[(month, r)], tetris // 1000))
            p.write(name, f'AB{r}', f'=IF(D{r}="","",CEILING(F{r},1000)-IF(G{r}="",0,CEILING(G{r},1000)))', True)
        p.write(name, 'AB51', '추가 실지출 테트리스 연결액')
        ET.SubElement(p.sheet(name).find(Q('cols')), Q('col'), min='28', max='28', hidden='1', width='13', customWidth='1')
        p.sheet(name).find(Q('dimension')).set('ref', 'A1:AB79')
        for _, date, tetris, actual in grid(values, month):
            p.write(name, actual, daily_formula(s[actual].value, date, 'H'), True)
            p.write(name, tetris, daily_formula(s[tetris].value, date, 'AB'), True)
            variance = re.sub(r'\d+$', lambda m: str(int(m[0])+1), actual)
            p.write(name, variance, variance_formula(tetris, actual), True)
        for addr, value in {
            'A33': '달력 지출 합계', 'E38': '달력지출', 'E39': '추가실지출(참고)',
            'A50': '추가 실지출 · 전액 달력 포함', 'G51': '기본 달력액', 'H51': '달력 연결액',
            'A62': '추가 실지출 합계 · 달력 포함',
            'A42': '추가 실지출은 사용일 달력에 전액 포함합니다.\n39행은 포함분 참고이며 다시 더하지 않습니다.\n목적자금 사용·적립과 대리결제는 기존 기준을 유지합니다.',
            'J76': '추가 실지출 전액 포함 · 대리·미리결제는 아래에서 따로 더합니다.',
            'J79': '추가 실지출은 달력에 이미 전액 포함 · 위 실 사용금액을 그대로 표시하며 다시 더하지 않습니다.',
        }.items():
            p.write(name, addr, value)
        for r in [item['row'] for item in extra_rows]:
            p.write(name, f'I{r}', '사용일 달력·테트리스에 전액 포함. 월 합계에 재가산하지 않음.')
        for addr, formula in {
            'I62': '=SUM(F52:F61)',
            'I39': '=' + '+'.join(f'IF(F{r}="",0,CEILING(F{r},1000))' for r in range(52, 62)),
            'H44': '=SUM(H37:H38,H40:H42)+H47-H43',
            'I44': '=SUM(I37:I38,I40:I42)+I47-I43',
            'E79': '=E78',
        }.items():
            p.write(name, addr, formula, True)
        calc_name = f'{month}월계산기'; calc = d[calc_name]; order = paint_order()
        occupied = [i for i, a in enumerate(order) if calc[a].fill.patternType == 'solid']
        next_slot = max(occupied, default=-1) + 1
        added_counts = Counter()
        painted = []
        for palette, units in paint:
            for _ in range(units):
                assert next_slot < len(order)
                addr = order[next_slot]
                assert calc[addr].fill.patternType != 'solid'
                cell = p.cell(calc_name, addr)
                style = deepcopy(xfs[int(cell.get('s', '0'))])
                legend = p.cell(calc_name, f'B{palette}')
                style.set('fillId', xfs[int(legend.get('s'))].get('fillId'))
                key = ET.tostring(style)
                if key not in style_cache:
                    style_cache[key] = len(xfs); xfs.append(style)
                cell.set('s', str(style_cache[key]))
                painted.append(addr); added_counts[palette] += 1; next_slot += 1
        for palette, count in added_counts.items():
            p.write(calc_name, f'AA{palette}', calc[f'AA{palette}'].value + count)
        if paint:
            p.write(calc_name, 'A154', str(calc['A154'].value).split('추가 실지출 통합:')[0] +
                    f' 추가 실지출 통합: 사용일 달력 포함 누락분 {sum(added_counts.values()):,}칸 보완. 기존 색상·목적재원 조정 유지.')
        if month == 10:
            p.write(calc_name, 'A154', '10/19까지 추가 실지출 보증보험료 110,075원(111칸)을 달력·계산기에 포함. 충전·대납·이미 반영된 두끼/더치 중복 제외.')
        metadata[name] = dict(rows=extra_rows, painted_cells=painted, added_color_counts=dict(added_counts))
    xfs.set('count', str(len(xfs)))
    p.trees.pop('xl/workbook.xml')
    p.save(candidate)
    w.close(); d.close()
    return metadata


def verify(baseline, final, metadata):
    b = openpyxl.load_workbook(baseline, data_only=True)
    w = openpyxl.load_workbook(final, data_only=True)
    assert w.sheetnames == b.sheetnames and '검증' not in w.sheetnames
    results = {}; raw_rows = list(w['RAW'].values)[1:]; checks = 0
    for month in MONTHS:
        name = f'{month}월'; s = w[name]; old = b[name]
        for date, _, tetris, actual in grid(old, month):
            expenses = [r for r in range(52, 62) if isinstance(s[f'A{r}'].value, datetime)
                        and s[f'A{r}'].value.date() == date.date() and s[f'D{r}'].value is not None]
            deltas = [extra_delta(s[f'F{r}'].value, s[f'G{r}'].value) for r in expenses]
            assert (s[actual].value or 0) == (old[actual].value or 0) + sum(x[0] for x in deltas), (name, date, actual)
            assert (s[tetris].value or 0) == (old[tetris].value or 0) + sum(x[1] for x in deltas), (name, date, tetris)
            checks += 1
        net = sum(s[f'F{r}'].value or 0 for r in range(52, 62))
        added = sum(s[f'H{r}'].value or 0 for r in range(52, 62))
        added_tetris = sum(s[f'AB{r}'].value or 0 for r in range(52, 62))
        assert s['I33'].value == old['I33'].value + added
        assert s['H39'].value == s['I62'].value == net
        assert s['H44'].value == old['H44'].value and s['I44'].value == old['I44'].value
        fixed = fixed_total([(s[f'U{r}'].value, s[f'Q{r}'].value) for r in range(4, 25)])
        advances = advance_total(raw_rows, s['T1'].value)
        totals = usage_totals(fixed, s['I33'].value, advances)
        assert (s['E78'].value, s['E79'].value) == totals
        assert s['E79'].value == old['E79'].value
        for a in ['E73', 'E74', 'E75', 'N69']:
            assert s[a].value == old[a].value
        for r in range(52, 62):
            for c in 'ABCDEFGH':
                assert s[f'{c}{r}'].value == old[f'{c}{r}'].value, (name, c, r)
        calc = w[f'{month}월계산기']; previous = b[calc.title]
        palette = {calc[f'B{r}'].fill.fgColor.rgb: r for r in range(3, 11)}
        counts = Counter(palette[calc[a].fill.fgColor.rgb] for a in paint_order() if calc[a].fill.patternType == 'solid')
        for r in range(3, 11):
            assert counts[r] == calc[f'C{r}'].value == calc[f'AA{r}'].value
        for a in paint_order():
            if a not in metadata[name]['painted_cells']:
                assert calc[a]._style == previous[a]._style, (month, 'original paint', a)
        assert calc['W12'].value == previous['W12'].value + added_tetris == s['E77'].value
        assert calc['W13'].value == s['I38'].value
        if month in [9, 10]:
            assert all(s[f'U{r}'].value is True for r in range(28, 32))
        else:
            assert s['AA1'].value is False and s['E79'].value == 0
        results[name] = dict(calendar_actual=s['I33'].value, calendar_tetris=calc['W12'].value,
                            extra_net=net, newly_in_calendar=added, newly_in_tetris=added_tetris,
                            actual_usage=s['E78'].value, final_usage=s['E79'].value,
                            fixed_actual=s['E74'].value, fixed_with_planned=s['E75'].value,
                            allocation_actual=s['H44'].value, allocation_tetris=s['I44'].value,
                            colored_cells=sum(counts.values()), counts=dict(counts),
                            salary_burden=calc['W13'].value)
    assert results['9월']['calendar_actual'] == 2073851 and results['9월']['calendar_tetris'] == 2116000
    assert results['10월']['calendar_actual'] == 1692425 and results['10월']['calendar_tetris'] == 1718000
    assert w['잔액']['G16'].value == 6559691
    assert w['잔액']['B248'].value == '출금' and w['RAW']['G2376'].value == -1300
    assert_no_errors(w)
    structure = validate_package(final); assert not structure['errors'], structure
    before, after = ExtendedPackage(baseline), ExtendedPackage(final)
    allowed = {before.sheets[f'{m}월'] for m in MONTHS} | {before.sheets[f'{m}월계산기'] for m in [9, 10]} | {'xl/styles.xml'}
    with ZipFile(baseline) as bz, ZipFile(final) as az:
        assert bz.namelist() == az.namelist() and az.testzip() is None
        changed = [n for n in bz.namelist() if bz.read(n) != az.read(n)]
        assert set(changed) <= allowed, changed
        oldstyle, newstyle = before.tree('xl/styles.xml'), after.tree('xl/styles.xml')
        for component in oldstyle:
            target = newstyle.find(component.tag)
            if component.tag == Q('cellXfs'):
                assert [ET.tostring(x) for x in component] == [ET.tostring(x) for x in list(target)[:len(component)]]
            else:
                assert ET.tostring(component) == ET.tostring(target)
    for item in [b, w]: item.close()
    return dict(date='2026-10-19', baseline_sha256=digest(baseline), workbook_sha256=digest(final),
                months=results, daily_checks=checks, formula_errors=0, structure_validation=structure,
                raw_ledger_reserve_prior_months_vba_controls_byte_identical=True,
                existing_paint_and_style_ids_preserved=True, changed_parts=changed,
                native_excel_opened=False, sharepoint_file_updated=False, details=metadata)


def mutation_regression(final, folder):
    def edit(p):
        for addr, value in {'A52': datetime(2026, 11, 5, 22, 30), 'D52': 5300, 'E52': 1100,
                            'A53': datetime(2026, 11, 5, 9), 'D53': 1700,
                            'A54': datetime(2026, 11, 6), 'D54': 1200,
                            'A55': datetime(2026, 11, 7), 'D55': -1300}.items(): p.write('11월', addr, value)
        p.write('12월', 'A52', datetime(2026, 12, 1, 20))
        p.write('12월', 'D52', 9900)
        # Already counted Dutch purchase now receives a 1,000 refund: adjust only once.
        p.write('10월', 'E53', 1000)
    path = recalculate(final, folder, edit)
    return verify_mutation_result(path)


def verify_mutation_result(path):
    w = openpyxl.load_workbook(path, data_only=True)
    s = w['11월']
    assert s['I33'].value == 5800 and s['E78'].value == s['E79'].value == 5800
    assert s['E77'].value == 8000  # 5,000+2,000+2,000−1,000, per transaction.
    assert s['H39'].value == s['I62'].value == 5800
    assert s['H44'].value == 1234406+5800
    assert w['12월']['I33'].value == w['12월']['E79'].value == 9900
    assert w['12월']['E77'].value == 10000
    assert w['10월']['I33'].value == 1691425 and w['10월']['E79'].value == 2265047
    assert w['10월']['E77'].value == 1717000
    # The established calculator needs matching new paint for new manual inputs;
    # it must flag this rather than silently claiming unpainted blocks agree.
    assert w['11월']['U28'].value == False  # LibreOffice may store logical FALSE as numeric 0.
    assert_no_errors(w); w.close()
    return dict(future_dated_inputs_automatic=True, same_day_transactions_round_separately=True,
                refund_net_expense=True, negative_adjustment_rounded_up=True, month_filter=True,
                already_covered_refund_deducted_once=True, usage_totals_no_second_addition=True,
                unpainted_new_entries_flag_existing_calculator_check=True)


def update_audit(previous, final, report, output):
    a = json.loads(Path(previous).read_text()); w = openpyxl.load_workbook(final, data_only=True)
    a['current_verification'] = 'extra_calendar_20261019_validation.json'
    a['workbook_sha256'] = report['workbook_sha256']
    for entry in a['calculators']:
        result = report['months'].get(f'{entry["month"]}월')
        if result is not None:
            entry.update(color_total=result['colored_cells']*1000, counts=result['counts'],
                         display=result['calendar_tetris'], salary_burden=result['salary_burden'])
    a['october_block_layout'].update(colored_cells=1718, full_blocks=171, partial_block_cells=8)
    a['october_lifestyle'].update(calendar_actual=1692425, calendar_tetris=1718000,
        scope='actual/tetris는 RAW 본인 변동지출 98건. calendar_actual/calendar_tetris는 추가 실지출 전액 포함 합계입니다.')
    for addr in a['october_summary']:
        if re.fullmatch(r'[A-Z]+[0-9]+', addr): a['october_summary'][addr] = w['10월'][addr].value
    a['extra_calendar_20261019'] = {**report, 'guide': '2026-10-19_추가실지출_달력통합_검증.md',
        'details': 'extra_calendar_20261019_validation.json', 'previous_version': 'history/2026-10-19/07'}
    a['fixed_cost_split_20261019']['scope'] = '고정비 두 기준 표시 당시 이력. 고정비는 그대로이며 현재 달력·합계·해시는 extra_calendar_20261019_validation.json을 따릅니다.'
    w.close(); Path(output).write_text(json.dumps(a, ensure_ascii=False, indent=2)+'\n')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for key in ['baseline', 'workdir', 'output', 'report', 'audit-input', 'audit-output']:
        parser.add_argument('--'+key, required=True)
    args = parser.parse_args(); folder = Path(args.workdir); candidate = folder/'candidate.xlsm'
    metadata = prepare(args.baseline, candidate)
    print('Extra expenses connected to the dated calendar; recalculating disposable copy.', flush=True)
    calculated = recalculate(candidate, folder/'default')
    caches = import_changed_caches(candidate, calculated, args.output)
    result = verify(args.baseline, args.output, metadata); result['changed_formula_caches'] = caches
    result['mutation_regression'] = mutation_regression(args.output, folder/'mutations')
    Path(args.report).write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    update_audit(args.audit_input, args.output, result, args.audit_output)
    print(json.dumps({k:v for k,v in result.items() if k not in ['details','changed_formula_caches']}, ensure_ascii=False, indent=2), flush=True)
