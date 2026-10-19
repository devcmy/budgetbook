"""Move one funded purchase's WRAP outflow from transfers to cash-out history.

Original RAW still records the real wallet top-up as a transfer. Only the WRAP
management ledger is classified as cash out; consumption is not counted again.
Patch XML and import changed caches from disposable macro-free recalculations.
"""
import argparse
from datetime import datetime
import gc
from hashlib import sha256
import json
import math
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.utils.datetime import to_excel

from balance_selected_ledger import check_view
from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from update_cash_details_20261019 import import_changed_caches
from validate_excel_structure import validate_package
from verify_reserve_cash_alignment import independent_balances, verify_month
from verify_reserve_redesign import assert_no_errors, recalculate


INPUT_EDITS = {'B248', 'D248', 'G248', 'L248'}


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def prepare(baseline, output):
    p = ExtendedPackage(baseline)
    old = p.cell('잔액', 'G248').find(Q('is')).find(Q('t')).text
    # P/Q/S/T are existing shared formulas, not static inputs. Preserve their
    # definitions and let the new source type drive group/account/signed amount.
    changes = {'B248': '출금', 'D248': '지출', 'L248': 'WRAP입출금',
               'G248': old + ' 사용자 요청: WRAP 관리상 출금으로 분류하여 입금·사용 표에 표시. RAW 충전 이체는 보존하고 소비는 더치 구매에만 1회 반영.'}
    for addr, value in changes.items(): p.write('잔액', addr, value)
    p.trees.pop('xl/workbook.xml')
    p.save(output)


def check_month(w, month):
    s = w['잔액']
    cutoff = datetime(2027, 1, 1) if month == 12 else datetime(2026, month+1, 1)
    balances, cash, _ = independent_balances(w, cutoff)
    for r in range(7, 16): assert s[f'G{r}'].value == balances[s[f'A{r}'].value]
    assert s['G16'].value == cash
    view = check_view(w, month)
    expected = {kind: sorted([r for r in range(64, 501)
                             if isinstance(s[f'A{r}'].value, datetime)
                             and s[f'A{r}'].value.year == 2026
                             and s[f'A{r}'].value.month == month
                             and s[f'P{r}'].value == kind],
                            key=lambda r: (s[f'A{r}'].value, -r), reverse=True)
                for kind in ['내역', '이체']}
    for group, helper, date_col, name_col, amount_col in [
        ('내역', 'M', 'A', 'C', 'F'), ('이체', 'N', 'G', 'K', 'J')]:
        actual = [int(s[f'{helper}{r}'].value)+63 for r in range(38, 60) if s[f'{helper}{r}'].value]
        assert actual == expected[group][:22], (month, group, actual, expected[group])
        for r, source in enumerate(actual, 38):
            # The upper tables display normalized U, not the synthetic month
            # start A used for historical records with an unknown actual day.
            shown_date = s[f'{date_col}{r}'].value
            source_date = s[f'U{source}'].value
            if isinstance(shown_date, datetime): shown_date = to_excel(shown_date)
            if isinstance(source_date, datetime): source_date = to_excel(source_date)
            if isinstance(shown_date, (int, float)) and isinstance(source_date, (int, float)):
                assert math.isclose(shown_date, source_date, rel_tol=0, abs_tol=1e-8)
            else:
                assert shown_date == source_date
            assert s[f'{name_col}{r}'].value == s[f'F{source}'].value
            assert s[f'{amount_col}{r}'].value == s[f'T{source}'].value
    if month == 10:
        assert len(expected['내역']) == 4 and len(expected['이체']) == 12
        assert expected['내역'][0] == 248 and 248 not in expected['이체']
        assert s['F38'].value == -15700 and s['B38'].value == '생활비'
        assert s['D7'].value == 27600 and s['E7'].value == 2537839
        assert s['B503'].value == '출금'
        assert s['G7'].value == 4519766 and cash == 6669766
        assert w['10월']['I33'].value == 1476900
        assert w['10월계산기']['W12'].value == 1501000
        assert w['10월']['H53'].value == 0
        verify_month(w, 10)
    assert_no_errors(w)
    return {'month': month, 'cash_total': cash, 'cash_history_count': len(expected['내역']),
            'transfer_count': len(expected['이체']), 'cash_history_source_rows': expected['내역'],
            'transfer_source_rows': expected['이체'], 'selected_month_ledger': view}


def verify(baseline, final):
    before, after = ExtendedPackage(baseline), ExtendedPackage(final)
    assert before.sheets == after.sheets and len(after.sheets) == 45
    part = before.sheets['잔액']
    with ZipFile(baseline) as a, ZipFile(final) as b:
        assert a.namelist() == b.namelist() and b.testzip() is None
        parts = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert parts == [part], parts
    old = {c.get('r'): c for c in before.sheet('잔액').iter(Q('c'))}
    new = {c.get('r'): c for c in after.sheet('잔액').iter(Q('c'))}
    assert old.keys() == new.keys()
    changed = []
    for addr, cell in old.items():
        target = new[addr]
        f, g = cell.find(Q('f')), target.find(Q('f'))
        assert (None if f is None else ET.tostring(f)) == (None if g is None else ET.tostring(g)), ('formula', addr)
        assert cell.get('s') == target.get('s'), ('style', addr)
        if ET.tostring(cell) != ET.tostring(target):
            changed.append(addr)
            assert addr in INPUT_EDITS or f is not None, ('unexpected literal change', addr)
        # Preserve all other input literals. Existing shared normalization and
        # ranking formulas may acquire caches after one item changes groups.
        row = int(''.join(c for c in addr if c.isdigit()))
        if 64 <= row <= 500 and addr not in INPUT_EDITS and f is None:
            assert ET.tostring(cell) == ET.tostring(target), ('ledger preservation', addr)
    w = openpyxl.load_workbook(final, data_only=True)
    s = w['잔액']
    assert s['B248'].value == '출금' and s['E248'].value == s['Y248'].value == 15700
    assert s['W248'].value == 2475 and w['RAW']['H2475'].value == '이체지출'
    assert w['RAW']['J2475'].value == '내부이체'
    assert w['RAW']['K2474'].value == 17000 and w['RAW']['K2376'].value == -1300
    assert s['I248'].value == 'OK'
    monthly = check_month(w, 10)
    w.close(); gc.collect()
    structure = validate_package(final)
    assert not structure['errors'], structure
    return {'date': '2026-10-19', 'data_as_of': '2026-10-18',
            'baseline_sha256': digest(baseline), 'workbook_sha256': digest(final),
            'new_ledger_type': '출금', 'new_display_group': 'WRAP 입금·사용',
            'input_edits': sorted(INPUT_EDITS), 'changed_cells_including_caches': changed,
            'changed_package_parts': parts, 'all_formulas_and_styles_unchanged': True,
            'other_transaction_inputs_unchanged': True,
            'raw_other_sheets_vba_controls_byte_identical': True,
            'wrap_cash_outflow': 15700, 'consumer_purchase': 17000,
            'additional_consumer_expense': 0, 'wrap_total_preserved': 6669766,
            'living_balance_preserved': 4519766, 'personal_actual_preserved': 1476900,
            'personal_tetris_preserved': 1501000, 'monthly_default_check': monthly,
            'structure_validation': structure, 'formula_errors': 0,
            'recalculation': 'disposable_macro_free_LibreOffice_changed_caches_only',
            'native_excel_opened': False, 'sharepoint_file_updated': False}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['baseline', 'workdir', 'output', 'report']: parser.add_argument('--'+name, required=True)
    a = parser.parse_args(); folder = Path(a.workdir)
    candidate = folder/'candidate.xlsm'
    prepare(a.baseline, candidate)
    print('Reclassified source; recalculating cash-out history and transfer views.', flush=True)
    calculated = recalculate(candidate, folder/'default')
    caches = import_changed_caches(candidate, calculated, a.output)
    report = verify(a.baseline, a.output)
    report['changed_formula_caches'] = caches
    report['month_regressions'] = []
    for month in [9, 11]:
        def edit(p, m=month): p.write('잔액', 'B3', datetime(2026, m, 1))
        result = recalculate(candidate, folder/f'month_{month}', edit)
        w = openpyxl.load_workbook(result, data_only=True)
        report['month_regressions'].append(check_month(w, month))
        w.close(); gc.collect()
    Path(a.report).write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['changed_formula_caches', 'changed_cells_including_caches']}, ensure_ascii=False), flush=True)
