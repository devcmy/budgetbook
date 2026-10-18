"""Link balance detail/ledger views to B3 without moving stable source rows.

The full input ledger remains at A64:Y500, folded by default. A formula-only
monthly ledger below it follows B3 automatically without VBA or a stale filter.
"""
import argparse
from copy import deepcopy
from datetime import datetime
import gc
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.utils.datetime import to_excel

from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from balance_month_selector import expected_summary
from verify_reserve_redesign import read, recalculate, assert_no_errors, cache_original
from validate_excel_structure import validate_package

FIRST, LAST, SIZE = 503, 524, 22


def prepare(source, output):
    p=ExtendedPackage(source)
    w=openpyxl.load_workbook(source,data_only=False)
    assert w['잔액']['A503'].value is None and w['잔액']['Z64'].value is None, 'Use the unmodified pre-update snapshot; do not apply this migration twice.'
    s=p.sheet('잔액')
    def write(addr,value,template=None):
        if template:
            p.cell('잔액',addr).set('s',p.cell('잔액',template).get('s','0'))
        p.write('잔액',addr,value,isinstance(value,str) and value.startswith('='))
    write('A18','조회월(B3)은 잔액·내역·하단 원장에 함께 적용됩니다. 이체 보냄·받음에는 WRAP 내부 계정 이동도 포함됩니다.')
    write('A35','내역 월(연동)')
    write('B35','=MONTH($B$3)','F35')
    write('V35','=$B$3')
    write('AA35','=MIN(MAX(1,IFERROR(INT($E$35),1)),$F$35)')
    write('AB35','=MIN(MAX(1,IFERROR(INT($H$35),1)),$I$35)')
    write('J35','="표시: 내역 "&$AA$35&"/"&$F$35&" · 이체 "&$AB$35&"/"&$I$35')
    for r in range(38,60):
        write(f'M{r}',w['잔액'][f'M{r}'].value.replace('$E$35','$AA$35'))
        write(f'N{r}',w['잔액'][f'N{r}'].value.replace('$H$35','$AB$35'))
    write('A60','내역·원장 조회는 위쪽 조회월(B3)에 자동 연동됩니다. 조회표는 직접 수정하지 말고, 전체 원장 입력자료는 아래 왼쪽 +를 펼쳐 편집하세요.')
    write('A62','선택월 WRAP 원장 | 조회월(B3) 자동 연동 · 전체 입력자료는 왼쪽 +로 펼침')

    validations=s.find(Q('dataValidations'))
    for v in list(validations):
        target=v.get('sqref','')
        if target=='B35':validations.remove(v)
        elif target=='E35':v.find(Q('formula2')).text='$F$35'
        elif target=='H35':v.find(Q('formula2')).text='$I$35'
        elif target=='B3':
            v.set('prompt','조회월을 고르면 상단 잔액·내역과 하단 WRAP 원장이 같은 월로 바뀝니다.')
    v=ET.SubElement(validations,Q('dataValidation'),type='whole',operator='between',
                    sqref='B501',allowBlank='0',showInputMessage='1',showErrorMessage='1',
                    errorStyle='stop',promptTitle='선택월 원장 페이지',
                    prompt='표시된 페이지 범위 안의 정수를 입력하세요.',
                    errorTitle='페이지 확인',error='현재 월의 페이지 범위 안에서 선택하세요.')
    ET.SubElement(v,Q('formula1')).text='1'
    ET.SubElement(v,Q('formula2')).text='$AB$501'
    validations.set('count',str(len(validations)))

    # Rank every dated source row, including inherited movements whose P/Q are
    # blank. Only A64:Y500 remains the financial/input source for the workbook.
    for r in range(64,501):
        write(f'Z{r}',f'=IFERROR(IF(AND(ISNUMBER(A{r}),A{r}>=$B$3,A{r}<EDATE($B$3,1)),'
              f'COUNTIFS($A$64:$A$500,">"&A{r},$A$64:$A$500,"<"&EDATE($B$3,1))+'
              f'COUNTIF($A$64:A{r},A{r}),0),0)')
    write('A501','원장 페이지','A35')
    write('B501',1,'E35')
    write('AC501','=COUNTIF($Z$64:$Z$500,">0")')
    write('AB501','=MAX(1,ROUNDUP($AC$501/22,0))')
    write('AA501','=MIN(MAX(1,IFERROR(INT($B$501),1)),$AB$501)')
    write('C501','=$AC$501&"건 · 표시 페이지 "&$AA$501&" / "&$AB$501&" · 최신 기록순"','F35')
    write('I501','=HYPERLINK("#\'잔액\'!A63",IF(COUNTIF($I$64:$I$500,"*확인*")>0,'
          '"전체 원장 확인 필요: "&COUNTIF($I$64:$I$500,"*확인*")&"건 (+로 펼침)","전체 원장 입력은 +로 펼침"))','B38')
    headings=['날짜/월','유형','보내는 곳','받는 곳','금액','내용','메모 요약','월','검증','보낸 계좌','받은 계좌','원본 행']
    for c,title in enumerate(headings,1):
        letter=openpyxl.utils.get_column_letter(c)
        write(f'{letter}502',title,f'{letter}63')
    for r in range(FIRST,LAST+1):
        ix=f'$AA{r}'
        write(f'AA{r}',f'=IFERROR(MATCH(($AA$501-1)*22+{r-FIRST+1},$Z$64:$Z$500,0),0)')
        def pick(c):return f'INDEX(${c}$64:${c}$500,MAX(1,{ix}))'
        write(f'A{r}',f'=IF({ix}=0,"",TEXT({pick("A")},IF({pick("L")}="원본WRAP","yyyy-mm","m/d hh:mm")))','B38')
        for col in 'BCDFIJK':
            write(f'{col}{r}',f'=IF({ix}=0,"",IF({pick(col)}="","",{pick(col)}))','B38')
        write(f'E{r}',f'=IF({ix}=0,"",{pick("E")})','G7')
        write(f'G{r}',f'=IF({ix}=0,"",LEFT({pick("G")},40)&IF(LEN({pick("G")})>40,"…",""))','G64')
        write(f'H{r}',f'=IF({ix}=0,"",TEXT({pick("H")},"yyyy-mm"))','B38')
        write(f'L{r}',f'=IF({ix}=0,"",HYPERLINK("#\'잔액\'!A"&({ix}+63),"원장 "&({ix}+63)&"행 ↗"))','B38')
    write('A525','선택월 조회용 표입니다. 메모 전문·원장 편집·새 거래 입력은 왼쪽 +를 펼친 뒤 원본 행에서 확인합니다. 날짜 미기록 원장은 월만 표시합니다.','A60')
    merges=s.find(Q('mergeCells'))
    for area in ['C501:H501','I501:L501','A525:L525']:
        ET.SubElement(merges,Q('mergeCell'),ref=area)
    merges.set('count',str(len(merges)))
    props=s.find(Q('sheetPr'))
    outline=props.find(Q('outlinePr'))
    if outline is None:
        outline=ET.Element(Q('outlinePr'))
        page_setup=props.find(Q('pageSetUpPr'))
        props.insert(list(props).index(page_setup) if page_setup is not None else len(props),outline)
    outline.set('summaryBelow','1');outline.set('showOutlineSymbols','1')
    for view in s.find(Q('sheetViews')):
        view.set('showOutlineSymbols','1')
    for r in range(63,501):
        p.cell('잔액',f'A{r}')
        row=p.row_maps['잔액'][r]
        row.set('hidden','1');row.set('outlineLevel','1');row.attrib.pop('collapsed',None)
    for r,height in [(501,32),(502,32),*[(r,54) for r in range(FIRST,LAST+1)],(525,40)]:
        row=p.row_maps['잔액'][r]
        row.set('hidden','0');row.set('outlineLevel','0');row.set('ht',str(height));row.set('customHeight','1')
    p.row_maps['잔액'][501].set('collapsed','1')
    ET.SubElement(s.find(Q('cols')),Q('col'),min='26',max='29',width='18',hidden='1',customWidth='1')
    s.find(Q('dimension')).set('ref','A1:AC525')
    workbook=p.tree('xl/workbook.xml')
    for name in workbook.find(Q('definedNames')):
        if name.get('name')=='_xlnm.Print_Area' and name.get('localSheetId')==str(w.sheetnames.index('잔액')):
            name.text="'잔액'!$A$1:$L$525"
    calc=workbook.find(Q('calcPr'))
    calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
    p.write('사용안내','A14','잔액 탭 조회월(B3)은 상단 잔액·내역과 하단 WRAP 원장에 함께 적용됩니다. 전체 입력 원장은 같은 탭의 +를 펼쳐 편집합니다.')
    p.save(output)
    structure=validate_package(output)
    assert not structure['errors'],structure['errors']


def expected_rows(sheet,month):
    start=datetime(2026,month,1)
    end=datetime(2027,1,1) if month==12 else datetime(2026,month+1,1)
    rows=[r for r in range(64,501) if isinstance(sheet[f'A{r}'].value,datetime) and start<=sheet[f'A{r}'].value<end]
    return sorted(rows,key=lambda r:(-to_excel(sheet[f'A{r}'].value),r))


def check_view(w,month,page=1):
    s=w['잔액'];rows=expected_rows(s,month)
    last=max(1,(len(rows)+21)//22);actual_page=min(max(1,page),last)
    assert s['B35'].value==month and s['V35'].value==datetime(2026,month,1)
    assert s['AC501'].value==len(rows) and s['AA501'].value==actual_page
    shown=[int(s[f'AA{r}'].value)+63 for r in range(FIRST,LAST+1) if s[f'AA{r}'].value]
    assert shown==rows[(actual_page-1)*22:actual_page*22],(month,page,shown,rows)
    for r in range(FIRST,LAST+1):
        ix=s[f'AA{r}'].value
        if not ix:
            assert all(s[f'{col}{r}'].value in [None,''] for col in 'ABCDEFGHIJKL')
        else:
            src=int(ix)+63
            assert s[f'E{r}'].value==s[f'E{src}'].value
            assert s[f'I{r}'].value==s[f'I{src}'].value
            assert str(src) in s[f'L{r}'].value
    for col in ['M','N']:
        for r in range(38,60):
            if s[f'{col}{r}'].value:
                src=int(s[f'{col}{r}'].value)+63
                assert s[f'A{src}'].value.month==month
    return {'month':month,'all_ledger_rows':len(rows),'displayed_page':actual_page,'displayed_rows':shown}


def verify_artifact(baseline,final):
    before,after=read(baseline),read(final)
    bf=openpyxl.load_workbook(baseline);af=openpyxl.load_workbook(final)
    assert before.sheetnames==after.sheetnames
    assert [(s.title,s.sheet_state) for s in before]==[(s.title,s.sheet_state) for s in after]
    for row in before['잔액'].iter_rows(min_row=64,max_row=500,max_col=25):
        for c in row:
            assert c.value==after['잔액'][c.coordinate].value,('ledger',c.coordinate)
            assert bf['잔액'][c.coordinate].value==af['잔액'][c.coordinate].value,('source formula',c.coordinate)
    for row in before['잔액'].iter_rows(min_row=7,max_row=32,max_col=12):
        for c in row:
            if c.coordinate=='A18':continue
            assert c.value==after['잔액'][c.coordinate].value,('balance',c.coordinate)
    for s in before:
        if s.title=='잔액':continue
        for row in s:
            for c in row:
                if (s.title,c.coordinate)==('사용안내','A14'):continue
                actual=after[s.title][c.coordinate].value
                if c.value in [None,''] and actual in [None,'']:continue
                assert c.value==actual,(s.title,c.coordinate,c.value,actual)
    assert all(af['잔액'].row_dimensions[r].hidden for r in range(63,501))
    assert all(not af['잔액'].row_dimensions[r].hidden for r in range(501,526))
    assert af['잔액'].row_dimensions[501].collapsed
    assert not any('B35' in v.sqref for v in af['잔액'].data_validations.dataValidation)
    assert set(map(str,bf['잔액'].merged_cells.ranges))<=set(map(str,af['잔액'].merged_cells.ranges))
    assert_no_errors(after)
    for row in after['검증']:
        for c in row:
            if isinstance(c.value,bool):assert c.value is True,c.coordinate
    with ZipFile(baseline) as a,ZipFile(final) as b:
        assert set(a.namelist())==set(b.namelist()) and b.testzip() is None
        for name in a.namelist():
            if name in ['xl/styles.xml','xl/vbaProject.bin'] or name.endswith('.rels') or name.startswith(('xl/drawings/','xl/media/','xl/ctrlProps/','xl/activeX/')):
                assert a.read(name)==b.read(name),name
    structure=validate_package(final)
    assert not structure['errors'],structure['errors']
    return {'source_ledger_addresses_values_formulas_preserved':True,'all_financial_values_preserved':True,
            'source_folded_and_month_view_visible_by_default':True,'sheet_order_and_visibility_preserved':True,
            'styles_vba_controls_drawings_preserved':True,'all_validation_booleans_true':True}


def run(a):
    root=Path(a.workdir);root.mkdir(parents=True,exist_ok=True)
    candidate=root/'candidate.xlsm';prepare(a.source,candidate)
    original=read(a.source);months=[];default=None
    for month in range(3,13):
        path=recalculate(candidate,root/f'month_{month}',lambda p,m=month:p.write('잔액','B3',datetime(2026,m,1)))
        w=read(path);assert_no_errors(w)
        for addr,value in expected_summary(original['잔액'],month).items():
            assert w['잔액'][addr].value==value,(month,addr)
        months.append(check_view(w,month))
        if month==10:default=path
        print(json.dumps({'month_passed':month,'rows':months[-1]['all_ledger_rows']},ensure_ascii=False),flush=True)
        del w;gc.collect()
    cases=[]
    for month in [7,8]:
        def edit(p,m=month):
            p.write('잔액','B3',datetime(2026,m,1));p.write('잔액','B501',2);p.write('잔액','E35',2)
        path=recalculate(candidate,root/f'page2_{month}',edit)
        w=read(path);assert_no_errors(w);result=check_view(w,month,2)
        cases.append({'name':f'page_two_{month}','passed':True,**result});del w;gc.collect()
    def stale(p):
        p.write('잔액','B3',datetime(2026,3,1))
        for c in ['B501','E35','H35']:p.write('잔액',c,99)
    path=recalculate(candidate,root/'stale_page',stale)
    w=read(path);assert_no_errors(w);check_view(w,3,99)
    assert w['잔액']['AA35'].value==w['잔액']['AB35'].value==1
    cases.append({'name':'stale_pages_clamped_on_month_change','passed':True});del w;gc.collect()
    def monthend(p):p.write('잔액','A247',datetime(2026,10,31,23,59,59))
    path=recalculate(candidate,root/'month_end',monthend)
    w=read(path);assert_no_errors(w);check_view(w,10)
    assert w['잔액']['AA503'].value+63==247
    cases.append({'name':'month_end_timestamp_included','passed':True});del w;gc.collect()
    cached=cache_original(candidate,default,a.output,a.source)
    checks=verify_artifact(a.source,a.output)
    report={'ui_updated_at':'2026-10-18','data_as_of':'2026-10-15','month_cell':'잔액!B3',
            'detail_month_cell':'잔액!B35','source_ledger':'잔액!A64:Y500',
            'selected_month_ledger':'잔액!A501:L525','page_cell':'잔액!B501',
            'summary_transfer_columns_are_selected_month_only':True,
            'summary_transfer_columns_include_internal_account_moves':True,
            'october_transfer_sent':2598839,'october_transfer_received':748510,
            'wrap_total':6685466,'living_balance':4535466,'temporary_balance':0,
            'selected_month_usage':239000,'month_tests':months,'scenario_tests':cases,
            'checks':checks,'formula_errors':0,'cached_formulas':cached,
            'baseline_sha256':sha256(Path(a.source).read_bytes()).hexdigest(),
            'workbook_sha256':sha256(Path(a.output).read_bytes()).hexdigest(),
            'recalculation':'Disposable macro-free XLSX in LibreOffice; original XLSM package retained.',
            'excel_ui_and_vba_executed':False}
    Path(a.report).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'verified':True,'months':len(months),'scenarios':len(cases),'hash':report['workbook_sha256']}),flush=True)


def final_display_recheck(a):
    """Reuse temporal/page tests only after proving only memo/style cells differ."""
    root=Path(a.workdir);root.mkdir(parents=True,exist_ok=True)
    candidate=root/'candidate.xlsm';prepare(a.source,candidate)
    old,new=ExtendedPackage(a.previous_candidate),ExtendedPackage(candidate)
    part=new.sheets['잔액']
    assert set(old.parts)==set(new.parts)
    for name in old.parts:
        if name!=part:assert old.parts[name]==new.parts[name],name
    oldtree,newtree=old.sheet('잔액'),new.sheet('잔액')
    originals={c.get('r'):c for c in oldtree.iter(Q('c'))}
    allowed={'G502','A525'}|{f'{col}{r}' for r in range(FIRST,LAST+1) for col in ['E','G']}
    changed=[]
    for cell in list(newtree.iter(Q('c'))):
        addr=cell.get('r');prior=originals[addr]
        if ET.tostring(cell)!=ET.tostring(prior):
            assert addr in allowed,addr
            changed.append(addr);cell.getparent().replace(cell,deepcopy(prior))
    assert changed and ET.tostring(newtree)==ET.tostring(oldtree)
    del old,new;gc.collect()
    path=recalculate(candidate,root/'default')
    w=read(path);assert_no_errors(w);check_view(w,10)
    for addr,value in expected_summary(read(a.source)['잔액'],10).items():
        assert w['잔액'][addr].value==value,addr
    for r in range(FIRST,LAST+1):
        if w['잔액'][f'AA{r}'].value:assert len(w['잔액'][f'G{r}'].value or '')<=41
    del w;gc.collect()
    cached=cache_original(candidate,path,a.output,a.source)
    checks=verify_artifact(a.source,a.output)
    result=json.loads(Path(a.previous_report).read_text())
    result.update(checks=checks,cached_formulas=cached,workbook_sha256=sha256(Path(a.output).read_bytes()).hexdigest())
    result['targeted_display_recheck']={'only_changed_cells':changed,
        'change':'Summarize long memos and use the existing lighter amount style in the monthly view.',
        'all_financial_filter_pagination_formulas_and_other_parts_byte_identical':True,
        'previous_candidate_sha256':sha256(Path(a.previous_candidate).read_bytes()).hexdigest(),
        'default_month_recalculated':True}
    Path(a.report).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'verified':True,'months':len(result['month_tests']),'scenarios':len(result['scenario_tests']),'hash':result['workbook_sha256']}),flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for name in ['source','output','workdir','report']:parser.add_argument('--'+name,required=True)
    parser.add_argument('--previous-candidate');parser.add_argument('--previous-report')
    args=parser.parse_args()
    if args.previous_candidate:
        assert args.previous_report
        final_display_recheck(args)
    else:run(args)
