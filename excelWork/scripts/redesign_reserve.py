"""A single reserve workspace: status, plan, linked cash movements.

The original XLSM package is patched, never round-tripped through openpyxl.
openpyxl is used only to author a new worksheet and read the existing workbook.
Legacy budget calculations are retained in one hidden sheet. Input registers
have stable row identities; month-filtered displays are never editable sources.
"""
from copy import deepcopy
from datetime import datetime
import argparse
import json
from pathlib import Path
import re
from zipfile import ZipFile, ZIP_DEFLATED

from lxml import etree as ET
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

from update_20261015 import Package, Q, NS, RNS, PNS

BACK = '_예비비계산'
FUNDS = ['명절비', '여행비', '경조사비', '상환비', '예비비']
U0, U1 = 100, 399
S0, S1 = 430, 729
UR = lambda c: f"예비비!${c}${U0}:${c}${U1}"
SR = lambda c: f"예비비!${c}${S0}:${c}${S1}"
DAY = 'yyyy-mm-dd'
MONEY = '#,##0;[Red]-#,##0;–'
MONTH = 'yyyy-mm'
BLUE, INK, MUTED = '25476A', '26384A', '66788A'
PALE, YELLOW, GREEN = 'EDF3F8', 'FFF4D6', 'E8F3ED'
VARIABLE = ['변동비', '담배', '일회성전자', '경조·선물', '여행', '약속·만남']


class ExtendedPackage(Package):
    def save(self, path):
        for name, tree in self.trees.items():
            self.parts[name] = ET.tostring(tree, encoding='UTF-8', xml_declaration=True)
        old = {i.filename for i in self.infos}
        with ZipFile(path, 'w', compression=ZIP_DEFLATED) as z:
            for info in self.infos:
                if info.filename in self.parts:
                    z.writestr(info, self.parts[info.filename])
            for name in self.parts.keys() - old:
                z.writestr(name, self.parts[name])


def clone_calculations(p):
    tree = deepcopy(p.sheet('예비비'))
    # No duplicated VBA codenames, hyperlinks or external relationship IDs.
    pr = tree.find(Q('sheetPr'))
    if pr is not None:
        pr.attrib.pop('codeName', None)
    for child in list(tree):
        if child.tag in [Q('hyperlinks'), Q('drawing'), Q('legacyDrawing')]:
            tree.remove(child)
    book = p.tree('xl/workbook.xml')
    sheets = book.find(Q('sheets'))
    sheetid = max(int(s.get('sheetId')) for s in sheets) + 1
    part = f'xl/worksheets/sheet{sheetid}.xml'
    rels = p.tree('xl/_rels/workbook.xml.rels')
    rid = 'rIdReserveCalculation'
    assert not any(r.get('Id') == rid for r in rels)
    ET.SubElement(rels, '{'+PNS+'}Relationship', Id=rid,
                  Type=RNS+'/worksheet', Target=f'worksheets/sheet{sheetid}.xml')
    ET.SubElement(sheets, Q('sheet'), name=BACK, sheetId=str(sheetid), state='hidden',
                  attrib={'{'+RNS+'}id': rid})
    types = p.tree('[Content_Types].xml')
    ET.SubElement(types, '{'+ET.QName(types).namespace+'}Override',
                  PartName='/'+part,
                  ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
    p.parts[part] = b''
    p.trees[part] = tree
    p.sheets[BACK] = part
    # Remove a stale dependency chain; Excel rebuilds it on automatic recalc.
    for r in list(rels):
        if r.get('Type', '').endswith('/calcChain'):
            rels.remove(r)
    for t in list(types):
        if t.get('PartName') == '/xl/calcChain.xml':
            types.remove(t)
    p.parts.pop('xl/calcChain.xml', None)
    p.trees.pop('xl/calcChain.xml', None)
    app = p.tree('docProps/app.xml')
    ans = {'a': ET.QName(app).namespace, 'v': 'http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'}
    count = app.find('a:HeadingPairs/v:vector/v:variant[2]/v:i4', ans)
    if count is not None:
        count.text = str(int(count.text) + 1)
    vector = app.find('a:TitlesOfParts/v:vector', ans)
    if vector is not None:
        node = ET.Element('{'+ans['v']+'}lpstr'); node.text = BACK
        vector.insert(sheetid-1, node)
        vector.set('size', str(len(vector)))


def append_styles(p, authored_zip):
    """Append style components; all original style IDs remain unchanged."""
    source = ET.fromstring(authored_zip.read('xl/styles.xml'))
    target = p.tree('xl/styles.xml')
    maps = {}
    for component in ['fonts', 'fills', 'borders']:
        dst, src = target.find(Q(component)), source.find(Q(component))
        maps[component] = {}
        for i, node in enumerate(src):
            maps[component][i] = len(dst)
            dst.append(deepcopy(node))
        dst.set('count', str(len(dst)))
    nums = target.find(Q('numFmts'))
    fmtmap = {}
    nextid = max([163] + [int(x.get('numFmtId')) for x in nums]) + 1
    for node in source.find(Q('numFmts')):
        oldid = int(node.get('numFmtId')); copied = deepcopy(node)
        copied.set('numFmtId', str(nextid)); fmtmap[oldid] = nextid
        nums.append(copied); nextid += 1
    nums.set('count', str(len(nums)))
    dst = target.find(Q('cellXfs')); stylemap = {}
    for i, node in enumerate(source.find(Q('cellXfs'))):
        copied = deepcopy(node)
        for attr, component in [('fontId', 'fonts'), ('fillId', 'fills'), ('borderId', 'borders')]:
            copied.set(attr, str(maps[component][int(node.get(attr, '0'))]))
        numid = int(node.get('numFmtId', '0'))
        copied.set('numFmtId', str(fmtmap.get(numid, numid)))
        copied.set('xfId', '0')
        stylemap[i] = len(dst); dst.append(copied)
    dst.set('count', str(len(dst)))
    return stylemap


def seed_records(d):
    # Use history is retained, but internal reallocations are not consumption.
    uses = [
        ('U001', '여동생 생일', '경조사비', datetime(2026,1,1), 50000, 0, '월기록', '1월 사용 기록 · 일자 미기록'),
        ('U002', '설 명절', '명절비', datetime(2026,2,1), 100000, 0, '월기록', '기존 예비비 2월 사용 · 원장 시작 전'),
        ('U003', '휴가비지원사업', '예비비', d['RAW']['A356'].value, 200000, 356, '기존반영', '사용월은 RAW 실제 2/19. 기존 계획의 1월과 구분'),
        ('U004', '어버이날', '명절비', datetime(2026,5,1), 148900, 0, '월기록', '생활비 이전 51,100원은 사용에서 제외'),
        ('U005', '한국장학재단 상환', '상환비', datetime(2026,5,1), 313980, 0, '월기록', '생활비 이전 36,020원은 사용에서 제외'),
        ('U006', '하계휴양 숙소', '여행비', datetime(2026,8,1), 300000, 0, '월기록', '원본 8월 지급 기록 · 실제 사용일 미기록'),
        ('U007', '막내 생일', '경조사비', d['RAW']['A2203'].value, 50000, 2203, '기존반영', '9월 월급 부담에서 이미 제외한 50,000원'),
        ('U008', '부모 생신 본인 부담', '경조사비', datetime(2026,10,1), 239000, 0, '월기록', '본인 납입 300,000 − 잔금 회수 61,000 = 239,000. 명절비 200,000 + 임시보관 39,000. 식사 139,000원 중복 제외'),
    ]
    seen = {u[5] for u in uses if u[5]}
    for row in d['RAW'].iter_rows(min_row=2):
        v = [c.value for c in row]
        if row[0].row in seen:
            continue
        if v[7] == '지출' and v[11] == '본인' and v[9] in {'여행','경조·선물','일회성전자'} and v[10] > 0:
            uses.append((f'R{row[0].row}', v[5], '생활비', v[0], v[10], row[0].row, '신규',
                         '기존 비고정 지출 유지. 예비비 부담으로 확인한 경우에만 항목 변경'))
    # Pre-ledger settlements are opening-balance evidence, not invented dates.
    settlements = [
        ('U001', 0, 50000, '기초반영', '3월 기초 이전 사용·정산. 실제 송금일 미기록'),
        ('U002', 0, 100000, '기초반영', '3월 기초에 반영된 명절 사용. 실제 송금일 미기록'),
        ('U003', 0, 200000, '기초반영', '3월 기초 이전 사용. 송금일을 2/19로 추정하지 않음'),
        ('U004', 102, 148900, '원장', '명절비 실제 사용'),
        ('U005', 106, 313980, '원장', '상환 실제 사용'),
        ('U006', 188, 300000, '원장', '숙소 지급'),
        ('U007', 222, 50000, '원장', '9월 생일 사용'),
        ('U008', 223, 200000, '원장', '확정된 교차계정 선이체: 명절비 → 가족통장 → 10월 경조 사용'),
        ('U008', 208, 39000, '원장', '임시보관 납입 100,000 중 행사 부담 39,000. 나머지 원금 61,000은 10/15 WRAP 회수'),
    ]
    return uses, settlements


def settlement_permission_formula(r):
    """Keep the two evidenced U008 sources narrow; never allow arbitrary cash links."""
    ledger = f'MAX(64,C{r})'
    allocated = f'SUMIF($C${S0}:$C${S1},C{r},$D${S0}:$D${S1})'
    inherited = f'AND(B{r}="U008",C{r}=223,J{r}="경조사비",F{r}="명절비")'
    family_cash = (f'AND(B{r}="U008",C{r}=208,J{r}="경조사비",'
                   f'F{r}="임시보관",{allocated}<=39000)')
    opening = (f'AND(H{r}="기초반영",OR(B{r}="U001",B{r}="U002",B{r}="U003"),'
               f'C{r}="",E{r}=DATE(2026,3,1))')
    ledgerok = (f'AND(H{r}="원장",C{r}>=64,C{r}<=500,'
                f'INDEX(잔액!$I:$I,{ledger})="OK",'
                f'E{r}=INDEX(잔액!$A:$A,{ledger}),F{r}=INDEX(잔액!$C:$C,{ledger}),'
                f'OR(F{r}=J{r},{inherited},{family_cash}),'
                f'OR(COUNTIF($A$8:$A$12,F{r})=1,{family_cash}),'
                f'COUNTIF($A$8:$A$12,INDEX(잔액!$D:$D,{ledger}))=0,'
                f'{allocated}<=INDEX(잔액!$Y:$Y,{ledger}))')
    return f'=IF(B{r}="","",IFERROR(IF(OR({opening},{ledgerok}),1,0),0))'


def author_ui(d, path):
    book = openpyxl.Workbook(); s = book.active; s.title = '예비비'
    s.sheet_view.showGridLines = False; s.sheet_view.zoomScale = 80
    s.freeze_panes = 'D8'
    s.sheet_properties.outlinePr.summaryBelow = False
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.column_dimensions['A'].width = 15
    s.column_dimensions['B'].width = 25
    s.column_dimensions['C'].width = 14
    for c in range(4,16): s.column_dimensions[col(c)].width = 10.5
    for c in range(16,35): s.column_dimensions[col(c)].hidden = True
    line = Side(style='hair', color='D9E3EB')
    for row in s.iter_rows(min_row=1,max_row=69,max_col=15):
        for cell in row:
            cell.font = Font(name='맑은 고딕', size=10, color=INK)
            cell.alignment = Alignment(vertical='center')
        s.row_dimensions[row[0].row].height = 24
    def text(addr, value, fmt=None, input=False):
        c=s[addr]; c.value=value
        c.font=Font(name='맑은 고딕',size=10,color=INK)
        c.alignment=Alignment(vertical='center',wrap_text=True)
        if fmt: c.number_format=fmt
        if input: c.fill=PatternFill('solid',fgColor=YELLOW)
        return c
    def span(range_, value, dark=False):
        s.merge_cells(range_); cell=s[range_.split(':')[0]]
        text(cell.coordinate,value)
        cell.font=Font(name='맑은 고딕',size=11,bold=True,color='FFFFFF' if dark else BLUE)
        for row in s[range_]:
            for c in row:c.fill=PatternFill('solid',fgColor=BLUE if dark else PALE)
        return cell
    def header(row, fields):
        for region,label in fields:
            span(f'{region[0]}{row}:{region[1]}{row}',label)
    def dv(range_, formula, type='list', low=None):
        v=DataValidation(type=type,formula1=formula,formula2=low,
                         allow_blank=True,showErrorMessage=True,showInputMessage=True)
        v.showDropDown=False;v.errorStyle='stop';v.errorTitle='입력 확인';v.error='목록 또는 허용 범위의 값을 입력하세요.'
        s.add_data_validation(v);v.add(range_)
    span('A1:O1','예비비 | 현황 · 계획 · 이체 내역',True);s.row_dimensions[1].height=34
    span('A2:O2','노란색만 입력합니다. 사용은 한 번, 나중 이체는 정산으로 연결합니다. 원장잔액과 사용가능액은 서로 다른 숫자입니다.')
    text('A3','조회월');text('B3',datetime(2026,10,1),MONTH,True)
    span('D3:O3','=TEXT($B$3,"yyyy년 m월")&IF($B$3>MAX(잔액!$A$64:$A$500)," · 미래 월: 현재 기록 이월 / 예측 아님"," · 선택월 말까지 기록 기준")')
    span('A4:O4','=IF($B$3=잔액!$B$3,"잔액 탭과 같은 조회월입니다.","잔액 탭 조회월은 "&TEXT(잔액!$B$3,"yyyy-mm")&"입니다. 아래 금액은 선택한 월의 동일 원장으로 계산합니다.")')
    for i in range(1,13):text(f'AD{i}',datetime(2026,i,1),MONTH)
    dv('B3','$AD$1:$AD$12')
    span('A6:O6','1  현황 | 원장잔액 − 정산대기 = 더 쓸 수 있는 돈',True)
    groups=[(('A','B'),'목적'),(('C','E'),'원장잔액'),(('F','H'),'정산대기'),(('I','K'),'사용가능액'),(('L','M'),'선택월 예비비 사용액'),(('N','O'),'연결 확인')]
    header(7,groups)
    for r,fund in enumerate(FUNDS,8):
        span(f'A{r}:B{r}',fund)
        for a,b in [('C','E'),('F','H'),('I','K'),('L','M'),('N','O')]:s.merge_cells(f'{a}{r}:{b}{r}')
        cash=f'SUMIFS(잔액!$Y$64:$Y$500,잔액!$D$64:$D$500,$A{r},잔액!$A$64:$A$500,"<"&EDATE($B$3,1))-SUMIFS(잔액!$Y$64:$Y$500,잔액!$C$64:$C$500,$A{r},잔액!$A$64:$A$500,"<"&EDATE($B$3,1))'
        text(f'C{r}',f'=IF($B$3<DATE(2026,3,1),"원장 없음",{cash})',MONEY)
        text(f'F{r}',f'=IF($B$3<DATE(2026,3,1),"원장 없음",SUMIFS($U${U0}:$U${U1},$C${U0}:$C${U1},$A{r}))',MONEY)
        text(f'I{r}',f'=IF(ISNUMBER(C{r}),C{r}-F{r},"원장 없음")',MONEY)
        text(f'L{r}',f'=SUMIFS($E${U0}:$E${U1},$C${U0}:$C${U1},$A{r},$D${U0}:$D${U1},">="&$B$3,$D${U0}:$D${U1},"<"&EDATE($B$3,1),$R${U0}:$R${U1},"OK")',MONEY)
        text(f'N{r}',f'=IF($B$3<DATE(2026,3,1),"원장 없음",IF($B$3=잔액!$B$3,IF(C{r}=INDEX(잔액!$G$8:$G$12,MATCH(A{r},잔액!$A$8:$A$12,0)),"일치","확인"),"동일월 원장"))')
        s[f'I{r}'].fill=PatternFill('solid',fgColor=GREEN)
    for a,b in [('A','B'),('C','E'),('F','H'),('I','K'),('L','M'),('N','O')]:span(f'{a}13:{b}13','합계' if a=='A' else '')
    for c in ['C','F','I','L']:text(c+'13',f'=IF($B$3<DATE(2026,3,1),"원장 없음",SUM({c}8:{c}12))',MONEY)
    text('L13','=SUM(L8:L12)',MONEY)
    text('N13','=IF($B$3<DATE(2026,3,1),"원장 없음",IF(C13=F13+I13,"차이 0","확인"))')
    span('A14:O14','명절비는 앞으로 본가 방문(숙박·교통 포함) 준비용입니다. 10월 숙박·본인 철도 181,800원은 비고정 지출로 유지했습니다.')
    span('A15:O15',f'=IF(COUNTIF($R${U0}:$R${U1},"확인*")+COUNTIF($N${S0}:$N${S1},"확인*")>0,"확인할 등록·연결 내역이 있습니다. 아래 등록 자료를 확인하세요.",IF(MIN($I$8:$I$12)<0,"사용가능액이 부족한 항목이 있습니다. 부담할 예비비 항목 또는 추가 적립을 확인하세요.","기존 경조사비 적립 기록 차이 100,000원은 검증 자료에 보존했습니다. 원장 금액을 임의로 맞추지 않았습니다."))')
    s['A15'].font=Font(name='맑은 고딕',size=9,color=MUTED)
    span('A16:H16','=HYPERLINK("#\'"&MONTH($B$3)&"월\'!K46","선택월 월간 내역 열기 ↗")')
    span('I16:O16','=HYPERLINK("#\'예비비\'!A98","사용·정산 등록 자료 펼치기 ↓")')
    span('A18:O18','2  계획 | 목적별 월 적립액 · 노란 칸 입력',True)
    text('A19','분류');text('B19','목적');text('C19','예정일')
    for m in range(1,13):text(f'{col(m+3)}19',datetime(2026,m,1),'m"월"')
    for i,oldrow in enumerate(range(5,30,3),20):
        text(f'A{i}',d['예비비'][f'A{oldrow}'].value)
        name=d['예비비'][f'B{oldrow}'].value
        text(f'B{i}','본가 방문·명절' if oldrow==17 else name)
        text(f'C{i}',d['예비비'][f'D{oldrow}'].value or '필요시')
        for m in range(1,13):text(f'{col(m+3)}{i}',d['예비비'].cell(oldrow,m+7).value or 0,MONEY,True)
    text('A29','계획 합계')
    for c in range(4,16):text(f'{col(c)}29',f'=SUM({col(c)}20:{col(c)}28)',MONEY)
    span('A30:O30','계획은 잔액을 늘리지 않습니다. 실제 적립은 잔액 원장에서 자동 반영합니다. 2025.12·2027.1 기존 계획은 계산 자료에 보존했습니다.')
    span('A32:O32','3  이체 내역 | 카드 사용과 정산을 한 화면에서 확인',True)
    text('A33','표시');text('B33','조회월',input=True);dv('B33','"조회월,미정산,전체"')
    span('D33:O33','조회월에는 해당 월 사용·정산과 이전 달 미정산을 함께 표시합니다. 선이체는 사용 시 다시 차감하지 않습니다.')
    text('A34','사용 페이지');text('B34',1,'0',True);dv('B34','1','whole','999')
    text('AC34',f'=MIN(MAX(1,B34),MAX(1,ROUNDUP(COUNTIF($X${U0}:$X${U1},1)/12,0)))')
    span('D34:O34',f'=COUNTIF($X${U0}:$X${U1},1)&"건 · 표시 페이지 "&AC34&" / "&MAX(1,ROUNDUP(COUNTIF($X${U0}:$X${U1},1)/12,0))')
    header(35,[(('A','A'),'사용일/월'),(('B','C'),'사용 내용'),(('D','D'),'항목'),(('E','E'),'사용액'),(('F','F'),'당월 정산'),(('G','G'),'누적 정산'),(('H','H'),'정산대기'),(('I','J'),'상태'),(('K','L'),'연결'),(('M','O'),'원본·등록')])
    for r in range(36,48):
        text(f'AB{r}',f'=IFERROR(MATCH(($AC$34-1)*12+{r-35},$Y${U0}:$Y${U1},0),0)')
        ix=f'$AB{r}'
        def pick(c):return f'INDEX(${c}${U0}:${c}${U1},{ix})'
        text(f'A{r}',f'=IF({ix}=0,"",IF({pick("G")}="월기록",TEXT({pick("D")},"yyyy-mm"),TEXT({pick("D")},"m/d")))')
        for a,b in [('B','C'),('I','J'),('K','L'),('M','O')]:s.merge_cells(f'{a}{r}:{b}{r}')
        for a,src in [('B','B'),('D','C'),('E','E'),('F','Z'),('G','T'),('H','U')]:
            text(f'{a}{r}',f'=IF({ix}=0,"",{pick(src)})',MONEY if a in 'EFGH' else None)
        text(f'I{r}',f'=IF({ix}=0,"",IF($B$3<DATE(2026,3,1),"정산일 미기록",IF({pick("D")}>=EDATE($B$3,1),"선이체",IF({pick("U")}=0,"정산완료",IF({pick("T")}>0,"부분정산","미정산")))))')
        text(f'K{r}',f'=IF({ix}=0,"",COUNTIFS($B${S0}:$B${S1},{pick("A")},$N${S0}:$N${S1},"OK",$E${S0}:$E${S1},"<"&EDATE($B$3,1))&"건")')
        text(f'M{r}',f'=IF({ix}=0,"",HYPERLINK("#\'예비비\'!A98","등록 "&({ix}+{U0-1})&"행 확인"))')
    span('A48:O48','누적 정산은 조회월 말까지입니다. 정산대기는 카드 미납액이 아니라 예비비에서 아직 이동하지 않은 사용분입니다.')
    text('A50','이체 페이지');text('B50',1,'0',True);dv('B50','1','whole','999')
    text('AC50',f'=MIN(MAX(1,B50),MAX(1,ROUNDUP(COUNTIF(\'{BACK}\'!$BH$120:$BH$556,">0")/12,0)))')
    span('D50:O50',f'=COUNTIF(\'{BACK}\'!$BH$120:$BH$556,">0")&"건 · 해당 월 원장 변동 + 선택월 사용에 연결된 선이체"')
    header(51,[(('A','A'),'원장일/월'),(('B','C'),'내용'),(('D','E'),'보낸 계정'),(('F','H'),'받는 계정/계좌'),(('I','J'),'원장 금액'),(('K','L'),'연결 상태'),(('M','O'),'원장 근거')])
    for r in range(52,64):
        text(f'AB{r}',f'=IFERROR(MATCH(($AC$50-1)*12+{r-51},\'{BACK}\'!$BH$120:$BH$556,0),0)')
        idx=f'$AB{r}';lr=f'INDEX(\'{BACK}\'!$BA$120:$BA$556,{idx})'
        for a,b in [('B','C'),('D','E'),('F','H'),('I','J'),('K','L'),('M','O')]:s.merge_cells(f'{a}{r}:{b}{r}')
        text(f'A{r}',f'=IF({idx}=0,"",TEXT(INDEX(잔액!$A:$A,{lr}),IF(INDEX(잔액!$L:$L,{lr})="원본WRAP","yyyy-mm","yyyy-mm-dd")))')
        for a,src in [('B','F'),('D','C')]:text(f'{a}{r}',f'=IF({idx}=0,"",INDEX(잔액!${src}:${src},{lr}))')
        text(f'F{r}',f'=IF({idx}=0,"",IF(INDEX(잔액!$K:$K,{lr})<>"",INDEX(잔액!$K:$K,{lr}),INDEX(잔액!$D:$D,{lr})))')
        text(f'I{r}',f'=IF({idx}=0,"",INDEX(잔액!$Y:$Y,{lr}))',MONEY)
        text(f'K{r}',f'=IF({idx}=0,"",IF(COUNTIFS($C${S0}:$C${S1},{lr},$N${S0}:$N${S1},"OK")>0,"정산 연결",INDEX(잔액!$B:$B,{lr})))')
        text(f'M{r}',f'=IF({idx}=0,"",HYPERLINK("#\'잔액\'!A"&{lr},"원장 "&{lr}&"행 ↗"))')
    span('A65:O65','새 사용은 아래 사용등록에 RAW행·예비비 항목만 지정합니다. 정산은 사용ID·원장행·금액을 연결합니다. 조회표에 직접 입력하지 않습니다.')
    span('A66:O66','기초·분류이동·계좌이체는 새 소비가 아닙니다. 조회월을 바꿔도 고정 등록 행의 사용ID와 연결은 유지됩니다.')
    span('A68:O68','=HYPERLINK("#\'예비비\'!A98","사용 등록 자료 펼치기 ↓")')
    span('A69:O69','=HYPERLINK("#\'예비비\'!A428","이체 연결 자료 펼치기 ↓")')
    # Stable registers: no month-dependent editable cells.
    span('A98:O98',f'="사용 등록 ("&COUNTIF($A${U0}:$A${U1},"?*")&"/300) | 왼쪽 +로 펼침. RAW행(F)·항목(C) 지정, 필요시 부담액(E) 수정."',True)
    names=['사용ID','내용','예비비 항목','사용일/월','본인 부담액','RAW행','반영 구분','메모']
    for c,name in enumerate(names,1):text(f'{col(c)}99',name)
    uses, settlements=seed_records(d)
    for r in range(U0,U1+1):
        for c in ['C','E','F']:text(c+str(r),None,MONEY if c=='E' else None,True)
        text(f'A{r}',f'=IF(F{r}="","",IF(F{r}>0,"R"&F{r},""))')
        text(f'B{r}',f'=IF(F{r}="","",IFERROR(INDEX(RAW!$F:$F,F{r}),""))')
        text(f'D{r}',f'=IF(F{r}="","",IFERROR(INDEX(RAW!$A:$A,F{r}),""))',DAY)
        text(f'E{r}',f'=IF(F{r}="","",IFERROR(INDEX(RAW!$K:$K,F{r}),0))',MONEY,True)
        text(f'G{r}','신규')
        if r-U0<len(uses):
            uid,title,fund,date,amount,raw,flag,note=uses[r-U0]
            for c,v in [('A',uid),('B',title),('C',fund),('D',date),('E',amount),('F',raw or None),('G',flag),('H',note)]:text(c+str(r),v,MONEY if c=='E' else DAY if c=='D' else None,c in ['C','E','F'])
            if raw:
                if uid.startswith('R'):text(f'A{r}',f'=IF(F{r}="","","R"&F{r})')
                for c,src in [('B','F'),('D','A'),('E','K')]:
                    text(c+str(r),f'=IFERROR(INDEX(RAW!${src}:${src},F{r}),"")',DAY if c=='D' else MONEY if c=='E' else None,c=='E')
        # Every active use has exactly one stable ID; unknown / overallocated inputs fail closed.
        var='+'.join(f'(INDEX(RAW!$J:$J,F{r})="{v}")' for v in VARIABLE)
        eligible_source=f'OR(AND(G{r}="기존반영",OR(A{r}="U003",A{r}="U007")),AND(INDEX(RAW!$L:$L,F{r})="본인",({var})>0))'
        fixed_sources=f'AND(OR(A{r}<>"U003",F{r}=356),OR(A{r}<>"U007",F{r}=2203),OR(G{r}<>"기존반영",A{r}="U003",A{r}="U007"))'
        rawok=f'IF(F{r}="",AND(LEFT(A{r},1)="U",G{r}="월기록"),AND(F{r}>=2,COUNTIF($F${U0}:$F${U1},F{r})=1,INDEX(RAW!$H:$H,F{r})="지출",D{r}=INDEX(RAW!$A:$A,F{r}),E{r}<=INDEX(RAW!$K:$K,F{r}),{eligible_source},{fixed_sources}))'
        text(f'R{r}',f'=IF(A{r}="","",IF(C{r}="생활비","생활비",IF(AND(COUNTIF($A$8:$A$12,C{r})=1,ISNUMBER(D{r}),E{r}>0,COUNTIF($A${U0}:$A${U1},A{r})=1,IFERROR({rawok},FALSE)),"OK","확인: 사용등록")))')
        text(f'S{r}',f'=IF(ISNUMBER(D{r}),DATE(YEAR(D{r}),MONTH(D{r}),1),0)')
        text(f'T{r}',f'=IF(R{r}="OK",SUMIFS($O${S0}:$O${S1},$B${S0}:$B${S1},A{r},$E${S0}:$E${S1},"<"&EDATE($B$3,1)),0)',MONEY)
        text(f'U{r}',f'=IF(AND(R{r}="OK",D{r}<EDATE($B$3,1),$B$3>=DATE(2026,3,1)),MAX(0,E{r}-T{r}),0)',MONEY)
        eligible=f'AND(F{r}>=2,G{r}<>"기존반영",INDEX(RAW!$L:$L,F{r})="본인",({var})>0)'
        text(f'V{r}',f'=IFERROR(IF(AND(R{r}="OK",{eligible}),E{r},0),0)',MONEY)
        text(f'W{r}',f'=IF(V{r}>0,ROUNDUP(INDEX(RAW!$K:$K,F{r})/1000,0)*1000-ROUNDUP((INDEX(RAW!$K:$K,F{r})-V{r})/1000,0)*1000,0)',MONEY)
        text(f'Z{r}',f'=IF(R{r}="OK",SUMIFS($O${S0}:$O${S1},$B${S0}:$B${S1},A{r},$E${S0}:$E${S1},">="&$B$3,$E${S0}:$E${S1},"<"&EDATE($B$3,1)),0)',MONEY)
        monthview=f'OR(S{r}=$B$3,Z{r}>0,U{r}>0)'
        text(f'X{r}',f'=IF(AND(R{r}="OK",IF($B$33="전체",D{r}<EDATE($B$3,1),IF($B$33="미정산",U{r}>0,{monthview}))),1,0)')
        text(f'Y{r}',f'=IF(X{r}=1,COUNTIF($X${U0}:X{r},1),0)')
    dv(f'C{U0}:C{U1}','"생활비,명절비,여행비,경조사비,상환비,예비비"')
    dv(f'F{U0}:F{U1}','2','whole','100000')
    dv(f'E{U0}:E{U1}','0','whole','1000000000')
    span('A400:O400','=HYPERLINK("#\'예비비\'!A1","현황으로 돌아가기 ↑")')
    s.row_dimensions.group(U0,U1,outline_level=1,hidden=True)
    s.row_dimensions[99].collapsed=True
    span('A428:O428',f'="이체 연결 ("&COUNTIF($B${S0}:$B${S1},"?*")&"/300) | 사용ID(B)·원장행(C)·금액(D) 입력. 한 이체를 여러 사용에 나눠 연결 가능."',True)
    for c,name in enumerate(['연결ID','사용ID','원장행','연결금액','원장일/월','출금계정','받는 곳','근거 유형','메모','사용분류','사용일','사용월','허용 확인','연결 확인','유효금액'],1):text(f'{col(c)}429',name)
    for r in range(S0,S1+1):
        text(f'A{r}',f'=IF(B{r}="","","S"&ROW())')
        for c in ['B','C','D']:text(c+str(r),None,MONEY if c=='D' else None,True)
        text(f'D{r}',f'=IF(C{r}="","",IFERROR(INDEX(잔액!$Y:$Y,C{r}),0))',MONEY,True)
        text(f'E{r}',f'=IF(C{r}="","",IFERROR(INDEX(잔액!$A:$A,C{r}),""))',DAY)
        text(f'F{r}',f'=IF(C{r}="","",IFERROR(INDEX(잔액!$C:$C,C{r}),""))')
        text(f'G{r}',f'=IF(C{r}="","",IF(INDEX(잔액!$K:$K,C{r})<>"",INDEX(잔액!$K:$K,C{r}),INDEX(잔액!$D:$D,C{r})))')
        text(f'H{r}','원장')
        if r-S0<len(settlements):
            uid,ledger,amount,kind,note=settlements[r-S0]
            for c,v in [('B',uid),('C',ledger or None),('D',amount),('H',kind),('I',note)]:text(c+str(r),v,MONEY if c=='D' else None,c in ['B','C','D'])
            if kind=='기초반영':
                text(f'E{r}',datetime(2026,3,1),MONTH);text(f'F{r}','기초 이전');text(f'G{r}','송금일 미기록')
        match=f'MATCH(B{r},$A${U0}:$A${U1},0)'
        for target,source in [('J','C'),('K','D'),('L','S')]:text(f'{target}{r}',f'=IFERROR(INDEX(${source}${U0}:${source}${U1},{match}),"")',DAY if target=='K' else None)
        text(f'M{r}',settlement_permission_formula(r))
        useok=f'INDEX($R${U0}:$R${U1},{match})="OK"'
        cap=f'SUMIF($B${S0}:$B${S1},B{r},$D${S0}:$D${S1})<=INDEX($E${U0}:$E${U1},{match})'
        text(f'N{r}',f'=IF(B{r}="","",IFERROR(IF(AND(COUNTIF($A${U0}:$A${U1},B{r})=1,{useok},D{r}>0,ISNUMBER(E{r}),M{r}=1,{cap}),"OK","확인: 정산연결"),"확인: 정산연결"))')
        text(f'O{r}',f'=IF(N{r}="OK",D{r},0)',MONEY)
    dv(f'B{S0}:B{S1}',f'$A${U0}:$A${U1}')
    dv(f'C{S0}:C{S1}','64','whole','500')
    dv(f'D{S0}:D{S1}','1','whole','1000000000')
    s.row_dimensions.group(S0,S1,outline_level=1,hidden=True)
    s.row_dimensions[429].collapsed=True
    span('A730:O730','=HYPERLINK("#\'예비비\'!A1","현황으로 돌아가기 ↑")')
    # Register warnings, rather than silently ignoring bad links.
    for r in range(U0,U1+1):
        s[f'O{r}']=f'=IF(LEFT(R{r},2)="확인",R{r},"")'
    for r in [7,19,35,51,99,429]:
        for cell in s[r][:15]:
            cell.fill=PatternFill('solid',fgColor=PALE)
            cell.font=Font(name='맑은 고딕',size=9,bold=True,color=BLUE)
            cell.alignment=Alignment(vertical='center',wrap_text=True)
        s.row_dimensions[r].height=30
    s.row_dimensions[7].height=42
    for r in list(range(8,14))+list(range(20,30))+list(range(36,48))+list(range(52,64)):
        for cell in s[r][:15]:cell.border=Border(bottom=line)
    for r in range(8,14):
        for c in ['C','F','I','L']:s[f'{c}{r}'].alignment=Alignment(horizontal='right',vertical='center',indent=1)
        s[f'N{r}'].alignment=Alignment(horizontal='center',vertical='center')
    s.print_area='A1:O69';s.print_options.horizontalCentered=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=2
    s.sheet_view.selection[0].activeCell='B3';s.sheet_view.selection[0].sqref='B3'
    book.save(path)
    return {'uses':len(uses),'settlements':len(settlements),'default_month':'2026-10',
            'confirmed_living_raw_rows':[2379,2380,2433,2459]}


def build(source, output, tempdir, manifest):
    root=Path(tempdir);root.mkdir(parents=True,exist_ok=True)
    w=openpyxl.load_workbook(source,data_only=False)
    d=openpyxl.load_workbook(source,data_only=True)
    p=ExtendedPackage(source)
    clone_calculations(p)
    # Preserve old planning and audit contracts, but remove their visible clutter.
    for sheet in w:
        if sheet.title=='예비비':continue
        for row in sheet:
            for cell in row:
                if cell.data_type=='f' and ('예비비!' in cell.value or "'예비비'!" in cell.value):
                    f=cell.value.replace("'예비비'!",f"'{BACK}'!").replace('예비비!',f"'{BACK}'!")
                    p.write(sheet.title,cell.coordinate,f,True)
    ui=root/'authored.xlsx'
    info=author_ui(d,ui)
    with ZipFile(ui) as z:
        stylemap=append_styles(p,z)
        tree=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
    for cell in tree.iter(Q('c')):
        cell.set('s',str(stylemap[int(cell.get('s','0'))]))
    p.trees[p.sheets['예비비']]=tree
    p.cell_maps.pop('예비비',None);p.row_maps.pop('예비비',None)
    # Guard empty source rows even in Excel's eagerly evaluated AND/OR branches.
    # INDEX(whole_column, 0) denotes an entire column, not an empty value.
    for cell in tree.iter(Q('c')):
        formula=cell.find(Q('f'))
        if formula is not None and formula.text:
            formula.text=re.sub(r'INDEX\((RAW!\$[A-Z]+:\$[A-Z]+),F(\d+)\)',r'INDEX(\1,MAX(2,F\2))',formula.text)
            formula.text=re.sub(r'INDEX\((잔액!\$[A-Z]+:\$[A-Z]+),C(\d+)\)',r'INDEX(\1,MAX(64,C\2))',formula.text)
    booktree=p.tree('xl/workbook.xml')
    names=booktree.find(Q('definedNames'))
    if names is None:
        names=ET.Element(Q('definedNames'))
        booktree.insert(list(booktree).index(booktree.find(Q('calcPr'))),names)
    sheetindex=w.sheetnames.index('예비비')
    for node in list(names):
        if node.get('name')=='_xlnm.Print_Area' and node.get('localSheetId')==str(sheetindex):names.remove(node)
    ET.SubElement(names,Q('definedName'),name='_xlnm.Print_Area',localSheetId=str(sheetindex)).text="'예비비'!$A$1:$O$69"
    # Editable year plan feeds all 12 monthly sheets and the balance plan columns.
    for visible,oldrow in enumerate(range(5,30,3),20):
        for month in range(1,13):
            p.write(BACK,f'{col(month+7)}{oldrow}',f"='예비비'!{col(month+3)}{visible}",True)
    # Monthly cache-independent summaries. Expenses are recognized on use date;
    # transfers only settle them. Existing calendar/painted totals are retained.
    for month in range(1,13):
        rr=month+1
        p.write(BACK,f'AA{rr}',datetime(2026,month,1))
        for c,fund in enumerate(FUNDS,28):
            f=f'=SUMIFS({UR("E")},{UR("C")},"{fund}",{UR("D")},">="&$AA{rr},{UR("D")},"<"&EDATE($AA{rr},1),{UR("R")},"OK")'
            p.write(BACK,f'{col(c)}{rr}',f,True)
        for c,src in [('AG','V'),('AH','W')]:
            p.write(BACK,f'{c}{rr}',f'=SUMIFS({UR(src)},{UR("D")},">="&$AA{rr},{UR("D")},"<"&EDATE($AA{rr},1))',True)
        name=f'{month}월';calc=f'{month}월계산기'
        for r,c in zip(range(64,69),range(28,33)):
            p.write(name,f'N{r}',f"='{BACK}'!{col(c)}{rr}",True)
        # All original living charges stay unchanged unless a transaction is
        # explicitly assigned to reserve funding in the stable use register.
        p.write(name,'H38',f'={w[name]["H38"].value.lstrip("=")}-\'{BACK}\'!AG{rr}',True)
        p.write(name,'I38',f'={w[name]["I38"].value.lstrip("=")}-\'{BACK}\'!AH{rr}',True)
        old=w[calc]['W14'].value
        oldf=old.lstrip('=') if isinstance(old,str) and old.startswith('=') else str(old or 0)
        p.write(calc,'W14',f'={oldf}+\'{BACK}\'!AH{rr}',True)
    # The fixed dated financial audit must remain independent of UI month.
    for r in range(120,557):
        ledger=r-56
        p.write(BACK,f'BA{r}',ledger)
        p.write(BACK,f'BB{r}',f'=IFERROR(INDEX(잔액!$A:$A,BA{r}),0)',True)
        rel=(f'OR(COUNTIFS({SR("C")},BA{r},{SR("N")},"OK")>0,'
             f'AND(INDEX(잔액!$J:$J,BA{r})="가족통장",INDEX(잔액!$K:$K,BA{r})="미래에셋_WRAP형CMA"),'
             f'COUNTIF(예비비!$A$8:$A$12,INDEX(잔액!$C:$C,BA{r}))=1,'
             f'COUNTIF(예비비!$A$8:$A$12,INDEX(잔액!$D:$D,BA{r}))=1)')
        period=f'AND(BB{r}>=예비비!$B$3,BB{r}<EDATE(예비비!$B$3,1))'
        linked=f'AND(BB{r}<EDATE(예비비!$B$3,1),COUNTIFS({SR("C")},BA{r},{SR("L")},예비비!$B$3,{SR("N")},"OK")>0)'
        show=f'AND(BB{r}>0,{rel},OR({period},{linked}))'
        preceding='0' if r==120 else f'MAX($BH$120:BH{r-1})'
        p.write(BACK,f'BH{r}',f'=IFERROR(IF({show},{preceding}+1,0),0)',True)
    p.sheet(BACK).find(Q('dimension')).set('ref','A1:BH556')
    checks=[
        ('사용·정산 등록 오류',f'=COUNTIF(예비비!$R${U0}:$R${U1},"확인*")+COUNTIF(예비비!$N${S0}:$N${S1},"확인*")','=I193=0'),
        ('원장잔액 = 사용가능 + 정산대기','=IF(ISNUMBER(예비비!C13),예비비!C13-예비비!F13-예비비!I13,0)','=I194=0'),
        ('예비비 월 사용 ↔ 월간 내역','=예비비!L13-SUM(INDIRECT("\'"&MONTH(예비비!$B$3)&"월\'!N64:N68"))','=I195=0'),
        ('예비비 월 계획 ↔ 월간 계획','=SUM(INDEX(예비비!$D$20:$O$28,0,MONTH(예비비!$B$3)))-INDIRECT("\'"&MONTH(예비비!$B$3)&"월\'!M53")','=I196=0'),
    ]
    p.write('검증','A191','예비비 현황·계획·이체 연결 검증')
    for rr,(label,value,check) in enumerate(checks,193):
        p.clone_row('검증',168,rr)
        p.write('검증',f'A{rr}',label);p.write('검증',f'I{rr}',value,True);p.write('검증',f'M{rr}',check,True)
    p.sheet('검증').find(Q('dimension')).set('ref','A1:Y196')
    p.write('사용안내','A11','예비비는 현황·월별 적립계획·이체 내역을 한 화면에서 봅니다. 원장잔액은 잔액 탭과 같은 월 기준, 사용가능액은 미정산 사용분을 뺀 금액입니다.')
    p.write('사용안내','A27','예비비의 조회표는 자동 계산입니다. 사용등록의 RAW행·항목을 지정하고 정산등록에서 사용ID·원장행·금액을 연결합니다. 카드 사용과 다음 달 이체를 중복 소비로 합산하지 않습니다.')
    calc=p.tree('xl/workbook.xml').find(Q('calcPr'))
    calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')
    p.save(output)
    info.update(backend=BACK,usage_range=[U0,U1],settlement_range=[S0,S1],
                unchanged_living_travel_amount=181800,
                prior_budget_usage_adjustments='Internal reallocations 51,100 and 36,020 are not new consumption; vacation actual date follows RAW 2/19, not January plan.')
    Path(manifest).write_text(json.dumps(info,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps(info,ensure_ascii=False))


if __name__=='__main__':
    cli=argparse.ArgumentParser()
    for name in ['source','output','tempdir','manifest']:cli.add_argument('--'+name,required=True)
    args=cli.parse_args();build(args.source,args.output,args.tempdir,args.manifest)
