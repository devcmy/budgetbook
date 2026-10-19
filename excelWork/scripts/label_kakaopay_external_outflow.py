"""Label one WRAP outflow; inspect, but do not modify, other October transfers.

Only F248 and its two deterministic INDEX display caches may change. Bank
amounts, all formulas, source RAW, styles and every other ZIP part stay intact.
"""
import argparse
from collections import Counter
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl

from redesign_reserve import ExtendedPackage
from update_20261015 import Q
from validate_excel_structure import validate_package


LABEL = '외부출금 · 카카오페이 충전'
ALLOWED = {'F248', 'K38', 'F503'}
CARD_PAYMENTS = {232: 2462, 233: 2456, 234: 2460, 235: 2461, 237: 2455, 239: 2441}


def review(w):
    s, raw = w['잔액'], w['RAW']
    selected = []
    for r in range(64, 501):
        when = s[f'A{r}'].value
        if not isinstance(when, datetime) or when.year != 2026 or when.month != 10:
            continue
        if s[f'L{r}'].value not in ['WRAP계좌이체', 'WRAP내계정이동']:
            continue
        source, destination, amount = [s[f'{c}{r}'].value for c in 'CDE']
        direction = '외부입금' if source == '외부' else '외부출금' if destination == '외부' else '내부분류'
        assert s[f'I{r}'].value == 'OK'
        raw_row = s[f'W{r}'].value
        if raw_row is not None:
            assert raw[f'G{int(raw_row)}'].value == amount
        else:
            assert r == 247 and direction == '내부분류'
        followup = None
        if r in CARD_PAYMENTS:
            followup = CARD_PAYMENTS[r]
            assert raw[f'C{followup}'].value == s[f'K{r}'].value
            assert raw[f'G{followup}'].value == amount
            assert raw[f'H{followup}'].value == '이체지출'
            assert raw[f'F{followup}'].value == '카드정산 결제'
            assert raw[f'R{followup}'].value == '타계좌'
            assert raw[f'A{followup}'].value >= when
            finding = '같은 상대계좌·금액의 후속 카드대금 납부 확인. WRAP 출금 1회, 새 소비로 합산하지 않음.'
        elif r == 240:
            followup = 2444
            assert raw['C2444'].value == '신한은행' and raw['G2444'].value == amount == 90410
            assert raw['J2444'].value == '고정비' and raw['H2444'].value == '지출'
            assert w['10월']['O13'].value == 90410
            finding = '대출이자 준비 이체. 실제 이자는 RAW 2444행과 고정비에 반영되며 이체를 추가 소비로 합산하지 않음.'
        elif r == 241:
            assert amount == w['10월']['N57'].value == 50000
            finding = '모임통장 본인 납입. 공동계좌 적립 50,000원에 반영되며 모임통장 전체 사용을 다시 본인 소비에 더하지 않음.'
        elif r in [231, 236]:
            finding = '다른 은행에서 WRAP으로 들어온 이체. WRAP 잔액은 증가하되 새 급여·소득으로 재분류하지 않음.'
        elif r == 246:
            assert amount == 61000 and source == '외부' and destination == '임시보관'
            finding = '가족통장의 생신 미사용 원금 회수. 실제 WRAP 입금이며 새 소득이 아님.'
        elif r == 247:
            assert amount == 61000 and source == '임시보관' and destination == '생활비'
            finding = 'WRAP 안의 계정 분류 이동. 실제 은행 송금·통장 총액 변화 없음. 화면 표기 구분은 개선 후보이며 이번에는 수정하지 않음.'
        else:
            assert r == 248 and amount == 15700
            finding = '카카오페이머니로 외부출금. WRAP에서 15,700원 차감, 더치 구매 17,000원은 소비에 이미 1회 반영. 이 건의 표시만 수정.'
        selected.append({'ledger_row': r, 'raw_row': raw_row, 'date': when.isoformat(),
                         'sender': s[f'J{r}'].value, 'recipient': s[f'K{r}'].value,
                         'amount': amount, 'direction': direction, 'finding': finding,
                         'followup_raw_row': followup, 'modified': r == 248})
    assert len(selected) == 13
    counts = Counter(x['direction'] for x in selected)
    assert counts == Counter({'외부출금': 9, '외부입금': 3, '내부분류': 1})
    totals = {k: sum(x['amount'] for x in selected if x['direction'] == k) for k in counts}
    assert totals == {'외부출금': 2553539, '외부입금': 687510, '내부분류': 61000}
    # Independent WRAP cash roll-forward, excluding internal account moves.
    assert 8538603 + totals['외부입금'] + 8557 + 535 - totals['외부출금'] - 11900 == s['G16'].value == 6669766
    return {'selected_month': '2026-10', 'transfer_count': 13,
            'other_transfers_reviewed_only': 12, 'direction_counts': dict(counts),
            'direction_amounts': totals, 'card_payment_preparations': len(CARD_PAYMENTS),
            'reviewed_transfers': selected}


def apply(baseline, output):
    p = ExtendedPackage(baseline)
    assert p.cell('잔액', 'F248').find(Q('f')) is None
    p.write('잔액', 'F248', LABEL)
    for addr, expected in [('K38', 'INDEX($F$64:$F$500,$N38)'),
                           ('F503', 'INDEX($F$64:$F$500,MAX(1,$AA503))')]:
        cell = p.cell('잔액', addr)
        assert expected in cell.find(Q('f')).text
        for node in list(cell):
            if node.tag in [Q('v'), Q('is')]: cell.remove(node)
        cell.set('t', 'str')
        ET.SubElement(cell, Q('v')).text = LABEL
    p.trees.pop('xl/workbook.xml')
    p.save(output)


def verify(baseline, output):
    old, new = ExtendedPackage(baseline), ExtendedPackage(output)
    part = old.sheets['잔액']
    assert old.sheets == new.sheets and len(new.sheets) == 45
    with ZipFile(baseline) as a, ZipFile(output) as b:
        assert a.namelist() == b.namelist() and b.testzip() is None
        changed_parts = [n for n in a.namelist() if a.read(n) != b.read(n)]
        assert changed_parts == [part], changed_parts
        before, after = ET.fromstring(a.read(part)), ET.fromstring(b.read(part))
        bc = {c.get('r'): c for c in before.iter(Q('c'))}
        ac = {c.get('r'): c for c in after.iter(Q('c'))}
        assert bc.keys() == ac.keys()
        changed = {addr for addr in bc if ET.tostring(bc[addr]) != ET.tostring(ac[addr])}
        assert changed == ALLOWED, changed
        for addr in changed:
            oldf, newf = bc[addr].find(Q('f')), ac[addr].find(Q('f'))
            assert (None if oldf is None else ET.tostring(oldf)) == (None if newf is None else ET.tostring(newf))
            ac[addr].getparent().replace(ac[addr], deepcopy(bc[addr]))
        assert ET.tostring(before) == ET.tostring(after)
    w = openpyxl.load_workbook(output, data_only=True)
    assert all(w['잔액'][a].value == LABEL for a in ALLOWED)
    assert w['잔액']['N38'].value + 63 == w['잔액']['AA503'].value + 63 == 248
    assert w['잔액']['B248'].value == '이체' and w['잔액']['E248'].value == 15700
    assert w['잔액']['G7'].value == 4519766 and w['잔액']['G16'].value == 6669766
    assert w['10월']['I33'].value == 1476900 and w['10월계산기']['W12'].value == 1501000
    assert w['10월']['H53'].value == 0 and w['예비비']['F13'].value == 239000
    audit = review(w)
    structure = validate_package(output)
    assert not structure['errors'], structure
    w.close()
    return {'date': '2026-10-19', 'data_as_of': '2026-10-18',
            'baseline_sha256': sha256(Path(baseline).read_bytes()).hexdigest(),
            'workbook_sha256': sha256(Path(output).read_bytes()).hexdigest(),
            'new_label': LABEL, 'changed_cells': sorted(ALLOWED),
            'changed_package_parts': changed_parts, 'all_formulas_unchanged': True,
            'all_amounts_and_other_transactions_unchanged': True,
            'raw_styles_vba_controls_and_other_parts_byte_identical': True,
            'no_full_recalculation': True,
            'deterministic_index_display_caches_updated': ['K38', 'F503'],
            'structure_validation': structure, 'native_excel_opened': False,
            'sharepoint_file_updated': False, 'transfer_review': audit}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['baseline', 'output', 'report']: parser.add_argument('--'+name, required=True)
    a = parser.parse_args()
    apply(a.baseline, a.output)
    result = verify(a.baseline, a.output)
    Path(a.report).write_text(json.dumps(result, ensure_ascii=False, indent=2)+'\n')
    print(json.dumps({k:v for k,v in result.items() if k != 'transfer_review'}, ensure_ascii=False), flush=True)
