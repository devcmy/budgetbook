"""Independently verify the 144 -> 164 October import and refresh live audit.

Reads, but never saves, Excel files with openpyxl. Native Excel/VBA execution is
not claimed. Historical audit blocks keep their original dated evidence.
"""
import argparse
from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from hashlib import sha256
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

import openpyxl
from lxml import etree as ET

from balance_selected_ledger import check_view
from update_20261015 import Q
from validate_excel_structure import validate_package
from verify_reserve_redesign import assert_no_errors, verify_month


VARIABLE = {'변동비', '담배', '일회성전자', '경조·선물', '여행', '약속·만남'}
START, OLD_END, NEW_END = 2323, 2466, 2486
ASOF = datetime(2026, 10, 18)
TEXT_EDITS = {
    'RAW': {'Q1', 'Q2', 'Q3', 'Q4', 'Q5', 'Q6'},
    '사용안내': {'A16', 'A22', 'A23'},
    '잔액': {'D3'},
    '10월': {'A70', 'A72', 'J75', 'J80', 'J82', 'J83', 'J86'}
           | {f'V{r}' for r in range(4, 25)},
    '검증': {'Q14', 'A24', 'A25', 'A59', 'A62', 'A70', 'A71', 'A72'},
    '10월계산기': {'A154'},
}
FORMULA_EDITS = {
    '10월': {'E75', 'E83', 'E84', 'E85'}
           | {f'{col}{r}' for col in 'FGH' for r in [14, 15]},
    '검증': {'I44', 'I70'} | {f'M{r}' for r in [35, 63, 64, 65, 66, 70, 169, 170, 171, 172, 174]},
}


def digest(path):
    return sha256(Path(path).read_bytes()).hexdigest()


def same(a, b):
    if a in [None, ''] and b in [None, '']:
        return True
    if (isinstance(a, (int, float)) and not isinstance(a, bool)
            and isinstance(b, (int, float)) and not isinstance(b, bool)):
        return math.isclose(a, b, rel_tol=1e-12, abs_tol=1e-6)
    return a == b


def asset(value):
    return '모임통장' if value == '20260606' else value


def source_key(values):
    return values[0], asset(values[1]), values[5], values[6]


def raw_key(sheet, row):
    return (sheet[f'A{row}'].value, asset(sheet[f'C{row}'].value),
            sheet[f'G{row}'].value, sheet[f'H{row}'].value)


def verify(baseline, candidate, source, summary, previous_audit, output_audit, report):
    before = openpyxl.load_workbook(baseline, data_only=False)
    old = openpyxl.load_workbook(baseline, data_only=True)
    formulas = openpyxl.load_workbook(candidate, data_only=False)
    after = openpyxl.load_workbook(candidate, data_only=True)
    uploaded = openpyxl.load_workbook(source, data_only=True).active
    metadata = json.loads(Path(summary).read_text())
    for key, path in [('source_sha256', source), ('baseline_sha256', baseline)]:
        if key in metadata:
            assert metadata[key] == digest(path), (key, 'summary hash mismatch')
    assert old['RAW'].max_row == OLD_END and after['RAW'].max_row == NEW_END
    assert uploaded.max_row == 165
    assert list(next(uploaded.values)) == ['날짜', '자산', '분류', '소분류', '내용',
                                            '금액(원)', '수입/지출', '메모', '금액', '화폐']
    incoming = {r: tuple(c.value for c in uploaded[r]) for r in range(2, 166)}
    assert all(isinstance(v[0], datetime) and v[0].year == 2026 and v[0].month == 10
               and v[5] == v[8] and v[9] == 'KRW' for v in incoming.values())
    old_keys = Counter(raw_key(old['RAW'], r) for r in range(START, OLD_END + 1))
    source_keys = Counter(source_key(v) for v in incoming.values())
    final_keys = Counter(raw_key(after['RAW'], r) for r in range(START, NEW_END + 1))
    assert source_keys == final_keys
    assert not old_keys - source_keys
    assert sum((source_keys - old_keys).values()) == 20
    key_rows = defaultdict(list)
    for r in range(START, NEW_END + 1):
        key_rows[raw_key(after['RAW'], r)].append(r)
    assert all(len(rows) == 1 for rows in key_rows.values()), 'ambiguous input identities'
    mapping = [{'source_row': sr, 'workbook_row': key_rows[source_key(v)][0],
                'status': 'added' if key_rows[source_key(v)][0] > OLD_END else 'preserved'}
               for sr, v in incoming.items()]
    if 'source_mapping' in metadata:
        assert sorted(metadata['source_mapping'], key=lambda x: x['source_row']) == mapping
    for item in mapping:
        if item['status'] != 'added':
            continue
        v, r = incoming[item['source_row']], item['workbook_row']
        expected = {'A': v[0], 'B': datetime(2026, 10, 1), 'C': asset(v[1]),
                    'D': v[2], 'E': v[3], 'F': v[4], 'G': v[5], 'H': v[6], 'I': v[7]}
        for col, value in expected.items():
            assert same(after['RAW'][f'{col}{r}'].value, value), ('added source data', r, col)
    for r in range(2, OLD_END + 1):
        for col in list(range(1, 14)) + [18, 19]:
            assert same(old['RAW'].cell(r, col).value, after['RAW'].cell(r, col).value), ('old RAW', r, col)
        assert before['RAW'][f'N{r}'].value == formulas['RAW'][f'N{r}'].value

    assert before.sheetnames == formulas.sheetnames
    assert [(s.title, s.sheet_state) for s in before] == [(s.title, s.sheet_state) for s in formulas]
    for month in range(1, 10):
        for title in [f'{month}월', f'{month}월계산기']:
            for row in before[title]:
                for c in row:
                    assert c.value == formulas[title][c.coordinate].value, ('past formula/literal', title, c.coordinate)
                    assert same(old[title][c.coordinate].value, after[title][c.coordinate].value), ('past cache', title, c.coordinate)
    # The new calendar dates and audit expectations are intentional. Other
    # existing formulas may only grow the known current RAW range endpoint.
    literal_changes, formula_changes = [], []
    paint_order = [f'{openpyxl.utils.get_column_letter(c)}{r}'
                   for r0 in range(15, 145, 5) for c0 in range(2, 22, 2)
                   for c in range(c0, c0 + 2) for r in range(r0, r0 + 5)]
    repainted = set(paint_order[1321:1501])
    for s in before:
        target = formulas[s.title]
        assert set(map(str, s.merged_cells.ranges)) == set(map(str, target.merged_cells.ranges)), ('merges', s.title)
        assert str(s.data_validations) == str(target.data_validations), ('validation', s.title)
        for row in s:
            for c in row:
                t = target[c.coordinate]
                # A newly materialized unstyled cell and a missing cell both
                # use Excel's default style, despite None vs all-zero arrays.
                if c._style != t._style and (c.has_style or t.has_style):
                    assert s.title == '10월계산기' and c.coordinate in repainted, ('unexpected style', s.title, c.coordinate)
                    for attr in ['font', 'border', 'alignment', 'number_format', 'protection']:
                        assert copy(getattr(c, attr)) == copy(getattr(t, attr)), ('non-fill style', c.coordinate, attr)
                if c.value == t.value:
                    continue
                if s.title == '잔액' and c.row == 248:
                    continue
                if s.title == '10월' and c.coordinate == 'U1':
                    assert t.value == ASOF
                    continue
                if s.title == '10월계산기' and c.coordinate in {f'AA{r}' for r in range(3, 11)}:
                    continue
                if c.coordinate in FORMULA_EDITS.get(s.title, set()):
                    assert t.data_type == 'f', ('expected formula', s.title, c.coordinate)
                    formula_changes.append((s.title, c.coordinate))
                    continue
                if c.data_type == 'f' and t.data_type == 'f':
                    expected = re.sub(r'(RAW!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?)2466\b', r'\g<1>2486', c.value)
                    assert t.value == expected, ('unexpected formula', s.title, c.coordinate, c.value, t.value)
                    formula_changes.append((s.title, c.coordinate))
                    continue
                assert c.coordinate in TEXT_EDITS.get(s.title, set()), ('unexpected literal', s.title, c.coordinate, c.value, t.value)
                assert isinstance(t.value, str) and not t.value.startswith('='), ('expected guide text', s.title, c.coordinate)
                literal_changes.append((s.title, c.coordinate))

    raw = after['RAW']
    value = lambda col, r: raw[f'{col}{r}'].value
    living = [r for r in range(START, NEW_END + 1)
              if value('H', r) == '지출' and value('L', r) == '본인' and value('J', r) in VARIABLE]
    actual = sum(value('K', r) for r in living)
    tetris = sum(math.ceil(value('K', r) / 1000) * 1000 for r in living)
    assert len(living) == 92 and actual == 1476900 and tetris == 1501000
    added_personal = [r for r in living if r > OLD_END]
    assert len(added_personal) == 17
    assert sum(value('K', r) for r in added_personal) == 177160
    assert sum(math.ceil(value('K', r) / 1000) * 1000 for r in added_personal) == 180000
    scope = Counter(value('R', r) for r in range(START, NEW_END + 1))
    assert scope == Counter({'WRAP': 15, '카드사용': 101, '타계좌': 48})
    new_common = [r for r in range(OLD_END + 1, NEW_END + 1) if value('C', r) == '모임통장']
    assert len(new_common) == 2 and sum(value('G', r) for r in new_common) == 36800
    assert all(value('L', r) != '본인' and value('R', r) == '타계좌' for r in new_common)
    days = []
    month = after['10월']
    for date_row in [3, 8, 13, 18, 23, 28]:
        for col in range(2, 9):
            day = month.cell(date_row, col).value
            if not isinstance(day, datetime) or day.month != 10 or day > ASOF:
                continue
            records = [r for r in living if value('A', r).date() == day.date()]
            da = sum(value('K', r) for r in records)
            dt = sum(math.ceil(value('K', r) / 1000) * 1000 for r in records)
            assert same(month.cell(date_row + 1, col).value or 0, dt), ('daily tetris', day)
            assert same(month.cell(date_row + 2, col).value or 0, da), ('daily actual', day)
            days.append({'day': day.date().isoformat(), 'actual': da, 'tetris': dt})
    assert len(days) == 18
    assert month['I33'].value == month['H38'].value == actual
    assert sum(month[f'I{r}'].value or 0 for r in [4, 9, 14, 19, 24, 29]) == tetris
    assert month['I38'].value == tetris
    assert month['E75'].value == 357230
    assert month['E80'].value == 103600
    for addr in ['E83', 'E84', 'E85']:
        assert month[addr].value == 0
    calc = formulas['10월계산기']
    palette = {calc[f'B{r}'].fill.fgColor.rgb: r for r in range(3, 11)}
    counts = Counter()
    filled = []
    for addr in paint_order:
        c = calc[addr]
        filled.append(c.fill.patternType == 'solid')
        if filled[-1]:
            assert c.fill.fgColor.rgb in palette, ('unknown color', addr)
            counts[palette[c.fill.fgColor.rgb]] += 1
    assert filled == [True] * 1501 + [False] * (len(filled) - 1501)
    for r in range(3, 11):
        assert counts[r] == after['10월계산기'][f'C{r}'].value == after['10월계산기'][f'AA{r}'].value
    assert after['10월계산기']['W12'].value == after['10월계산기']['W13'].value == tetris
    for addr in paint_order[:1321]:
        assert before['10월계산기'][addr]._style == formulas['10월계산기'][addr]._style

    ledger = after['잔액']
    rows = [r for r in range(64, 501) if isinstance(ledger[f'A{r}'].value, datetime)
            and ledger[f'A{r}'].value.year == 2026 and ledger[f'A{r}'].value.month == 10]
    assert len(rows) == 16
    links = Counter(ledger[f'{c}{r}'].value for r in rows for c in ['W', 'X'] if ledger[f'{c}{r}'].value is not None)
    wrap_rows = {r for r in range(START, NEW_END + 1) if value('R', r) == 'WRAP'}
    assert set(links) == wrap_rows and set(links.values()) == {1}
    for r in rows:
        assert ledger[f'I{r}'].value == 'OK', ('ledger invalid', r)
        for c in ['W', 'X']:
            rr = ledger[f'{c}{r}'].value
            if rr is not None:
                assert ledger[f'E{r}'].value == value('G', int(rr)), ('linked amount', r, rr)
    for addr, expected in {'G7': 4519766, 'G15': 0, 'G16': 6669766, 'G17': 2150000}.items():
        assert ledger[addr].value == expected, (addr, ledger[addr].value)
    assert ledger['E248'].value == 15700
    assert ledger['C248'].value == '생활비' and ledger['D248'].value == '외부'
    assert value('G', int(ledger['W248'].value)) == 15700
    assert value('C', int(ledger['W248'].value)) == '미래에셋_WRAP형CMA'
    for r in [246, 247]:
        for col in range(1, 26):
            # Date-ranked display keys legitimately move when a newer WRAP
            # transfer is appended. Preserve their formulas, not stale ranks.
            prior = before['잔액'].cell(r, col)
            current = formulas['잔액'].cell(r, col)
            assert prior.value == current.value, ('family ledger formula/input', r, col)
            if prior.data_type != 'f':
                assert same(old['잔액'].cell(r, col).value, ledger.cell(r, col).value), ('family ledger preservation', r, col)
    assert after['예비비']['L7'].value == '선택월 예비비 사용액'
    for addr, expected in {'L13': 239000, 'F13': 0, 'I13': 2150000, 'E107': 239000}.items():
        assert after['예비비'][addr].value == expected
    assert after['예비비']['F438'].value == old['예비비']['F438'].value
    reserve = verify_month(after, 10)
    view = check_view(after, 10)
    assert len(view['displayed_rows']) == 16 and view['displayed_rows'][0] == 248
    assert all(formulas['잔액'].row_dimensions[r].hidden for r in range(63, 501))
    assert all(not formulas['잔액'].row_dimensions[r].hidden for r in range(501, 526))
    assert_no_errors(after)
    false_cells = [c.coordinate for row in after['검증'] for c in row if c.value is False]
    assert not false_cells, false_cells
    for addr in ['U28', 'U29', 'U30', 'U31']:
        assert month[addr].value is True
    for addr in ['O149', 'O150']:
        assert after['10월계산기'][addr].value is True
    assert not [(s.title, c.coordinate) for s in formulas for row in s for c in row
                if c.data_type == 'f' and '#REF!' in c.value]
    structure = validate_package(candidate)
    assert not structure['errors'], structure['errors']
    preserved = []
    with ZipFile(baseline) as a, ZipFile(candidate) as b:
        assert a.namelist() == b.namelist() and b.testzip() is None
        for name in a.namelist():
            if name != 'xl/styles.xml' and name != 'xl/workbook.xml' and not name.startswith('xl/worksheets/sheet'):
                assert a.read(name) == b.read(name), ('package part changed', name)
                preserved.append(name)
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in b.read('[Content_Types].xml')
        props = ET.fromstring(b.read('xl/worksheets/sheet4.xml')).find(Q('sheetPr'))
        assert [ET.QName(c).localname for c in props] == ['tabColor', 'outlinePr', 'pageSetUpPr']
        for section in ['fonts', 'fills', 'borders', 'cellXfs', 'cellStyleXfs', 'numFmts']:
            x = ET.fromstring(a.read('xl/styles.xml')).find(Q(section))
            y = ET.fromstring(b.read('xl/styles.xml')).find(Q(section))
            assert len(y) >= len(x)
            assert all(ET.tostring(u) == ET.tostring(v) for u, v in zip(x, y)), ('old style component', section)

    inflows = sum(value('G', r) for r in wrap_rows if value('S', r) in ['입금', '이체입금'])
    outflows = sum(value('G', r) for r in wrap_rows if value('S', r) in ['출금', '이체출금'])
    assert inflows == 696602 and outflows == 2565439
    assert 8538603 + inflows - outflows == ledger['G16'].value
    audit = json.loads(Path(previous_audit).read_text())
    audit['as_of'] = '2026-10-18'
    audit['workbook_sha256'] = digest(candidate)
    audit['raw']['count'] = NEW_END - 1
    audit['raw']['end'] = max(value('A', r) for r in range(2, NEW_END + 1)).isoformat()
    audit['raw']['by_month']['2026-10'] = 164
    audit['raw']['october_scope'] = dict(scope)
    audit['october_ledger'].update(rows=16, raw_linked_rows=15, internal_classification_rows=1,
                                   raw_link_missing=[], raw_link_duplicates={}, unexpected_raw_links=[],
                                   invalid_rows=[], amount_mismatches=[])
    accounts = ['생활비', '명절비', '여행비', '경조사비', '상환비', '예비비', '결제대금비', '금융', '임시보관']
    audit['wrap'].update(balances={fund: ledger[f'G{r}'].value for r, fund in enumerate(accounts, 7)},
                         total=ledger['G16'].value, inflows=inflows, outflows=outflows, cached_matches=True)
    for entry in audit['calculators']:
        if entry['month'] == 10:
            entry.update(color_total=tetris, counts=dict(counts), display=tetris, salary_burden=tetris,
                         cached_counts_match=True, expected_display_matches=True, month_display_matches=True)
    audit['october_block_layout'].update(colored_cells=1501, full_blocks=150, partial_block_cells=1)
    audit['october_lifestyle'].update(records=len(living), actual=actual, tetris=tetris,
                                    calendar_actual=month['I33'].value, calendar_tetris=tetris)
    audit['october_summary'] = {k: month[k].value for k in audit['october_summary'] if k != 'fixed_independent'}
    audit['october_summary']['fixed_independent'] = sum(month[f'Q{r}'].value or 0 for r in range(4, 25)
                                                       if month[f'U{r}'].value in ['확정', '결제'])
    audit['vba'].update(binary_identical_to_previous_version=True, executed=False)
    for name in ['update_checks', 'temporary_to_living', 'balance_selected_month', 'excel_xml_repair']:
        previous_scope = audit[name].get('scope', '')
        audit[name]['scope'] = (previous_scope + ' 10/18 신규 RAW 반영 전의 이력이며 현재 수치·해시는 delta_20261018 및 최상위 필드를 따릅니다.').strip()
    checks = {'source_164_rows_matched_once': True, 'old_144_october_and_all_earlier_raw_preserved': True,
              'past_month_literals_formulas_and_values_preserved': True, 'daily_calendar_independently_matched': True,
              'existing_paint_layout_preserved': True, 'independent_color_counts_matched': True,
              'wrap_15_rows_linked_once': True, 'family_settlement_and_61000_reclassification_preserved': True,
              'reserve_matches_balance_sheet': True, 'selected_month_view_preserved': True,
              'old_styles_merges_and_validation_preserved': True, 'vba_controls_relationships_byte_identical': True,
              'sheet4_property_order_correct': True, 'formula_errors': 0, 'false_validation_cells': []}
    result = {'as_of': '2026-10-18', 'baseline_sha256': digest(baseline), 'source_sha256': digest(source),
              'workbook_sha256': digest(candidate), 'october_count': 164, 'new_rows': 20,
              'raw_total': NEW_END - 1, 'source_mapping': mapping, 'scope': dict(scope),
              'living_records': len(living), 'living_actual': actual, 'living_tetris': tetris,
              'new_personal_records': 17, 'new_personal_actual': 177160, 'new_personal_tetris': 180000,
              'new_shared_spending': 36800, 'shared_october_spending': month['E75'].value,
              'new_wrap_outflow': 15700, 'wrap_total': ledger['G16'].value, 'living_balance': ledger['G7'].value,
              'temporary_balance': 0, 'reserve_cash': 2150000, 'selected_month_reserve_use': 239000,
              'outstanding_reserve_settlement': 0, 'outstanding_advances': 103600,
              'daily_calendar': days, 'color_counts': dict(counts), 'reserve_month': reserve,
              'selected_month_ledger': view, 'structure_validation': structure,
              'preserved_package_parts': len(preserved), 'checks': checks,
              'intentional_literal_changes': literal_changes, 'intentional_formula_changes': formula_changes,
              'native_excel_opened': False, 'vba_executed': False,
              'recalculation': 'Disposable macro-free LibreOffice copy; original XLSM package/formulas retained, caches only imported.',
              'validation_scope': 'Focused worksheet-properties XSD and structural/accounting invariants; not full ECMA-376 validation.'}
    audit['delta_20261018'] = {k: v for k, v in result.items()
                              if k not in ['source_mapping', 'daily_calendar', 'intentional_literal_changes', 'intentional_formula_changes']}
    audit['delta_20261018'].update(details=Path(report).name, guide='2026-10-18_RAW반영_검증.md',
                                  source='raw/2026-10-18/2026-10-01_2026-10-31.xlsx',
                                  previous_version='history/2026-10-18/03')
    Path(report).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
    Path(output_audit).write_text(json.dumps(audit, ensure_ascii=False, indent=2) + '\n')
    print(json.dumps({'verified': True, 'october_count': 164, 'living_actual': actual,
                      'tetris': tetris, 'wrap_total': ledger['G16'].value,
                      'workbook_sha256': result['workbook_sha256']}, ensure_ascii=False), flush=True)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for option in ['baseline', 'candidate', 'source', 'summary', 'previous-audit', 'output-audit', 'report']:
        parser.add_argument('--' + option, required=True)
    a = parser.parse_args()
    verify(a.baseline, a.candidate, a.source, a.summary, a.previous_audit, a.output_audit, a.report)
