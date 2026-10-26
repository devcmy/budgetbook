"""Separate mutable RAW notes from accounting and mirror notes into live views.

One-time ZIP/XML migration of the RAW automation workbook. Memo display uses
I only; classification and duplicate identity must not depend on its text.
The historical anchor/reset limitations are intentionally not redesigned here.
"""
from collections import defaultdict
from copy import deepcopy
import argparse
import json
from pathlib import Path
import re
from zipfile import ZipFile

import openpyxl
from lxml import etree as ET
from openpyxl.styles import Alignment, Font, PatternFill, Protection

from redesign_reserve import ExtendedPackage, append_styles
from raw_auto_link_20261026 import (TABLE, BACK, PREPARED_END, default_class,
                                  insert_before_conditional)
from update_20261015 import Q
from update_cash_details_20261019 import import_changed_caches
from verify_reserve_redesign import recalculate
from validate_excel_structure import validate_package

DISCRIMINATOR='동일거래 구분'
MEMO='연동메모'
MEMO_INDEX='메모순번'
RR=lambda title:f'{TABLE}[{title}]'


def classification(r,safe_end):
    """Keep the old business rules except the memo-driven advance inference."""
    formula=default_class(r,safe_end)
    old=(f'IF(OR(ISNUMBER(SEARCH("대리결제",I{r})),ISNUMBER(SEARCH("대납",I{r}))),"대납",'
         f'IF(F{r}="보증보험료","이사부대",IF(OR(C{r}="가족통장",C{r}="모임통장",C{r}="20260606"),"공금","변동비")))')
    new=f'IF(F{r}="보증보험료","이사부대",IF(OR(C{r}="가족통장",C{r}="모임통장",C{r}="20260606"),"공금","변동비"))'
    assert old in formula
    result=formula.replace(old,new)
    assert not re.search(r'\b[IM]'+str(r)+r'\b',result)
    return result


def key_formula(r):
    # Memo is deliberately absent. M is an explicit business discriminator.
    parts=[f'TEXT(O{r},"0.0000000000")']+[f'LEN({c}{r})&":"&{c}{r}' for c in ['C','D','E','F']]
    parts += [f'V{r}',f'LEN(H{r})&":"&H{r}',f'LEN(M{r})&":"&M{r}']
    return f'=IF(O{r}="","",'+ '&"|"&'.join(parts)+')'


def duplicate_formula(r):
    criterion=f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(P{r},"~","~~"),"*","~*"),"?","~?")'
    checks='*'.join(f'(${c}$2:{c}{r}={c}{r})' for c in ['O','C','D','E','F','V','H','M'])
    return f'=IF(P{r}="","",IF(LEN({criterion})<=250,COUNTIF($P$2:P{r},{criterion}),SUMPRODUCT({checks})))'


def memo_formula(r):
    return f'=IF(OR(A{r}="",I{r}=""),"",I{r}&"")'


def memo_index_formula(r):
    return f'=IF(OR(B{r}="",AN{r}=""),"",IF(YEAR(B{r})<>2026,"",MONTH(B{r})*10000+COUNTIFS($B$2:B{r},B{r},$AN$2:AN{r},"?*")))'


def joined_rows(rows):
    cells=[f'IF(RAW!A{r}="","",RAW!AN{r})' for r in sorted(set(rows))]
    if not cells:return '=""'
    if len(cells)==1:return '='+cells[0]
    return '=IFERROR(_xlfn.TEXTJOIN(CHAR(10),TRUE,'+','.join(cells)+'),"메모 길이 초과 · RAW에서 확인")'


def dynamic_note(row_expression):
    return (f'=IFERROR(IF(OR({row_expression}="",{row_expression}<2),"",'
            f'INDEX({RR(MEMO)},{row_expression}-1)),"")')


def ledger_note(r):
    a=f'IF(W{r}<2,"",INDEX({RR(MEMO)},W{r}-1))'
    b=f'IF(OR(X{r}<2,X{r}=W{r}),"",INDEX({RR(MEMO)},X{r}-1))'
    return f'=IFERROR(_xlfn.TEXTJOIN(CHAR(10),TRUE,{a},{b}),"")'


def fixed_note(slot):
    condition=f'({RR("월")}=$T$1)*({RR("고정항목")}={slot})*({RR("적용분류")}="고정비")*({RR("동일거래순번")}=1)'
    return f'=IFERROR(_xlfn.TEXTJOIN(CHAR(10),TRUE,IF({condition},{RR(MEMO)},"")),"메모 길이 초과 · RAW에서 확인")'


def make_styles(p,folder):
    wb=openpyxl.Workbook();s=wb.active
    for cell in ['A1','A2','A3','A4']:
        s[cell]='style';s[cell].font=Font(name='맑은 고딕',size=10,color='243B53')
        s[cell].alignment=Alignment(wrap_text=True,vertical='center')
    s['A1'].fill=PatternFill('solid',fgColor='17324D');s['A1'].font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
    s['A2'].fill=PatternFill('solid',fgColor='F5F8FC');s['A2'].protection=Protection(locked=False)
    s['A3'].fill=PatternFill('solid',fgColor='EDF7F2')
    s['A4'].number_format='yyyy-mm-dd'
    path=folder/'memo_styles.xlsx';wb.save(path);wb.close()
    with ZipFile(path) as z:
        mapping=append_styles(p,z);root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        return {c.get('r'):mapping[int(c.get('s','0'))] for c in root.iter(Q('c'))}


def merge(p,name,ref):
    sheet=p.sheet(name);ranges=sheet.find(Q('mergeCells'))
    if ranges is None:ranges=ET.Element(Q('mergeCells'));insert_before_conditional(sheet,ranges)
    assert not any(c.get('ref')==ref for c in ranges)
    ET.SubElement(ranges,Q('mergeCell'),ref=ref);ranges.set('count',str(len(ranges)))


def column_layout(p,name,first,last,intervals):
    columns=p.sheet(name).find(Q('cols'))
    assert columns is not None
    for node in list(columns):
        start,end=int(node.get('min')),int(node.get('max'))
        if end<first or start>last:continue
        columns.remove(node)
        if start<first:
            left=deepcopy(node);left.set('max',str(first-1));columns.append(left)
        if end>last:
            right=deepcopy(node);right.set('min',str(last+1));columns.append(right)
    for start,end,width,hidden in intervals:
        ET.SubElement(columns,Q('col'),min=str(start),max=str(end),width=str(width),customWidth='1',hidden=str(int(hidden)))
    ordered=sorted(list(columns),key=lambda x:int(x.get('min')))
    for node in list(columns):columns.remove(node)
    columns.extend(ordered)


def raw_table(p,w,styles):
    raw=w['RAW'];tablepart=next(n for n,b in p.parts.items() if n.startswith('xl/tables/') and TABLE.encode() in b)
    table=p.tree(tablepart);columns=table.find(Q('tableColumns'))
    assert len(columns)==39,'Already migrated or unknown input table'
    for i,title in enumerate([MEMO,MEMO_INDEX],40):
        ET.SubElement(columns,Q('tableColumn'),id=str(i),name=title)
    columns[8].set('name','메모');columns[12].set('name',DISCRIMINATOR)
    columns.set('count',str(len(columns)));table.set('ref',f'A1:AO{PREPARED_END}')
    table.find(Q('autoFilter')).set('ref',table.get('ref'))
    safe_end=max(c.row for c in w[BACK]['AN'] if c.value is not None)
    groups=defaultdict(lambda:defaultdict(list))
    for r in range(2,2494):
        core=tuple(raw[f'{c}{r}'].value for c in ['A','C','D','E','F','G','H'])
        groups[core][raw[f'P{r}'].value].append(r)
    discriminators={}
    for group in groups.values():
        if len(group)>1:
            for number,rows in enumerate(group.values(),1):
                for r in rows:discriminators[r]='' if number==1 else f'기존별도-{number}'
    formulas={}
    for r in range(2,PREPARED_END+1):
        forms={'P':key_formula(r),'Q':duplicate_formula(r),'AJ':classification(r,safe_end),
               'AN':memo_formula(r),'AO':memo_index_formula(r)}
        for c,f in forms.items():p.write('RAW',f'{c}{r}',f,True)
        old_notes=[str(raw[f'{c}{r}'].value) for c in ['I','M'] if raw[f'{c}{r}'].value not in [None,'']]
        merged='\n'.join(dict.fromkeys(old_notes)) or None
        assert merged is None or len(merged)<=32767,('Memo length exceeds Excel capacity',r)
        p.write('RAW',f'I{r}',merged)
        p.write('RAW',f'M{r}',discriminators.get(r))
        p.cell('RAW',f'M{r}').set('s',str(styles['A2']))
        if r==2:formulas=forms
    for col,title in [('I','메모'),('M',DISCRIMINATOR),('AN',MEMO),('AO',MEMO_INDEX)]:
        p.write('RAW',col+'1',title);p.cell('RAW',col+'1').set('s',str(styles['A1']))
    for index,column in enumerate(columns,1):
        col=openpyxl.utils.get_column_letter(index)
        if col in formulas:
            for child in list(column):
                if child.tag==Q('calculatedColumnFormula'):column.remove(child)
            ET.SubElement(column,Q('calculatedColumnFormula')).text=formulas[col].lstrip('=')
    cols=p.sheet('RAW').find(Q('cols'))
    for old in list(cols):
        if int(old.get('min'))==15 and int(old.get('max'))==42:cols.remove(old)
    # Rebuild sorted column intervals, retaining all unrelated dimensions.
    for attrs in [dict(min='15',max='42',hidden='1',width='12',customWidth='1')]:
        ET.SubElement(cols,Q('col'),**attrs)
    ordered=sorted(list(cols),key=lambda c:int(c.get('min')))
    for node in list(cols):cols.remove(node)
    cols.extend(ordered)
    guide={4:('메모는 I열 하나','이전 추가입력·비고는 I 메모로 합쳤습니다. 분류·중복 판정과 무관합니다.'),
           8:('분류 수정','대납 등은 J 재분류에서 직접 선택합니다.'),
           9:('동일 거래 구별','실제 별개 결제만 M 동일거래 구분을 다르게 입력합니다.'),
           17:('메모 수정 위치','RAW I를 수정·삭제하면 연결된 모든 메모 보기도 바뀝니다.')}
    for r,(label,value) in guide.items():p.write('RAW',f'AQ{r}',label);p.write('RAW',f'AR{r}',value)
    column_layout(p,'RAW',9,9,[(9,9,45,False)])
    column_layout(p,'RAW',13,13,[(13,13,18,False)])
    return {'existing_explicit_discriminators':discriminators,'table_range':table.get('ref')}


def prepare(source,candidate,folder):
    folder=Path(folder);w=openpyxl.load_workbook(source,data_only=True);f=openpyxl.load_workbook(source)
    p=ExtendedPackage(source);styles=make_styles(p,folder);meta=raw_table(p,w,styles)
    destinations=[];replaced=[]
    def note(name,addr,formula,array=False):
        old=w[name][addr].value
        if old not in [None,'']:replaced.append({'sheet':name,'cell':addr,'previous':old})
        p.write(name,addr,formula,True)
        # Old internal links on a now-empty memo cell can make readers display
        # the link location instead of blank. Memo views need no legacy links.
        links=p.sheet(name).find(Q('hyperlinks'))
        if links is not None:
            for link in list(links):
                if link.get('ref','').replace('$','')==addr:links.remove(link)
            if not len(links):p.sheet(name).remove(links)
        if array:p.cell(name,addr).find(Q('f')).attrib.update({'t':'array','ref':addr})
        destinations.append((name,addr))
    for m in range(1,13):
        name=f'{m}월';calc=f'{m}월계산기'
        p.write(name,'V3','RAW 메모')
        merge(p,name,'V3:AA3')
        column_layout(p,name,22,29,[(22,22,14,False),(23,23,18,False),(24,26,9,False),(27,29,12,True)])
        column_layout(p,calc,1,1,[(1,1,12,False)])
        for r in list(range(4,25))+[26]:
            note(name,f'V{r}',fixed_note(r),True);p.cell(name,f'V{r}').set('s',str(styles['A3']))
            merge(p,name,f'V{r}:AA{r}')
        for r in range(52,62):
            formula=str(f[name][f'D{r}'].value or '')
            rows=[int(x) for x in re.findall(r'RAW!\$?[GK]\$?(\d+)',formula)]
            if (m,r)==(9,52):rows=[2157]
            if rows:note(name,f'I{r}',joined_rows(rows))
        # Compact five-row memo browser; calculators mirror the same page.
        p.write(name,'V34','RAW 거래 메모 · 수정은 RAW I에서');merge(p,name,'V34:AA34')
        p.cell(name,'V34').set('s',str(styles['A1']))
        p.write(name,'AC40',f'=COUNTIFS({RR("월")},$T$1,{RR(MEMO)},"?*")',True)
        p.write(name,'V40','메모 페이지');p.write(name,'W40',1);p.cell(name,'W40').set('s',str(styles['A2']))
        p.write(name,'X40','=AC40&"건 · 페이지 변경은 왼쪽 숫자 입력"',True);merge(p,name,'X40:AA40')
        page='IFERROR(MAX(1,MIN(INT($W$40),MAX(1,ROUNDUP($AC$40/5,0)))),1)'
        p.write(calc,'A157','RAW 거래 메모 · 월간 메모 페이지와 연동');merge(p,calc,'A157:U157')
        p.cell(calc,'A157').set('s',str(styles['A1']))
        for k in range(1,6):
            r=34+k;lookup=f'MATCH(MONTH($T$1)*10000+({page}-1)*5+{k},{RR(MEMO_INDEX)},0)'
            note(name,f'W{r}',f'=IFERROR(INDEX({RR(MEMO)},{lookup}),"")')
            p.write(name,f'V{r}',f'=IFERROR(INT(INDEX({RR("유효일시")},{lookup})),"")',True)
            merge(p,name,f'W{r}:AA{r}');p.cell(name,f'W{r}').set('s',str(styles['A3']))
            p.cell(name,f'V{r}').set('s',str(styles['A4']))
            cr=157+k
            note(calc,f'B{cr}',f'=IF(\'{name}\'!W{r}="","",\'{name}\'!W{r})')
            p.write(calc,f'A{cr}',f'=IF(\'{name}\'!V{r}="","",\'{name}\'!V{r})',True)
            merge(p,calc,f'B{cr}:U{cr}');p.cell(calc,f'B{cr}').set('s',str(styles['A3']))
            p.cell(calc,f'A{cr}').set('s',str(styles['A4']))
            for sheet,row,height in [(name,r,42),(calc,cr,34)]:
                p.row_maps[sheet][row].set('ht',str(height));p.row_maps[sheet][row].set('customHeight','1')
    # Existing RAW row links remain the authority; notes cannot change them.
    for r in range(64,501):
        rows=[int(w['잔액'][f'{c}{r}'].value) for c in ['W','X'] if isinstance(w['잔액'][f'{c}{r}'].value,(int,float)) and w['잔액'][f'{c}{r}'].value>=2]
        if rows or (not w['잔액'][f'A{r}'].value and not w['잔액'][f'G{r}'].value):
            note('잔액',f'G{r}',ledger_note(r))
    for r in range(100,400):
        # All registrations including future ones display the selected RAW row's note.
        note('예비비',f'H{r}',dynamic_note(f'F{r}'))
        # Unlinked handwritten historical records have independent notes, not RAW copies.
        if not w['예비비'][f'F{r}'].value:
            p.write('예비비',f'H{r}',w['예비비'][f'H{r}'].value)
            destinations.remove(('예비비',f'H{r}'))
            if not w['예비비'][f'A{r}'].value:note('예비비',f'H{r}',dynamic_note(f'F{r}'))
    for r in range(430,730):
        if w['예비비'][f'C{r}'].value or not w['예비비'][f'B{r}'].value:
            note('예비비',f'I{r}',f'=IFERROR(IF(C{r}="","",IF(INDEX(잔액!$G$64:$G$500,C{r}-63)="","",INDEX(잔액!$G$64:$G$500,C{r}-63))),"")')
    for addr,text in {
        'A19':'RAW A 일시·C~I 거래를 추가합니다. I 메모는 메모로만 사용하고 J 분류·K 본인부담·L 결제구분과 분리합니다.',
        'A22':'메모 수정은 RAW I에서 합니다. 이전 I 추가입력/M 비고를 I 메모로 합쳤습니다. 연결된 월간·계산기·잔액/예비비 메모는 자동 변경·삭제됩니다.',
        'A24':'대납은 J 재분류에서 선택합니다. 메모의 대납/대리결제 단어로 분류하지 않습니다. 같은 일시의 별개 거래는 M 동일거래 구분으로 구별합니다.',
        'A25':'RAW 전체 삭제·교체 초기화 기능은 아닙니다. 기존 행에 다른 거래를 덮어쓸 때 J/K/L의 이전 보정이 남을 수 있으므로 새 빈 행을 사용합니다.',
    }.items():p.write('사용안내',addr,text)
    p.save(candidate);w.close();f.close()
    replaced=[item for item in replaced if p.cell(item['sheet'],item['cell']).find(Q('f')) is not None]
    meta.update(memo_destinations=destinations,replaced_legacy_notes=replaced)
    (folder/'memo_migration_metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2,default=str)+'\n')
    return meta


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for arg in ['source','output','workdir']:parser.add_argument('--'+arg,required=True)
    args=parser.parse_args();folder=Path(args.workdir);folder.mkdir(parents=True,exist_ok=True)
    prepare(args.source,folder/'candidate.xlsm',folder)
    result=recalculate(folder/'candidate.xlsm',folder/'calculation')
    import_changed_caches(folder/'candidate.xlsm',result,args.output)
    report=validate_package(args.output);assert not report['errors'],report
    print(json.dumps(report,ensure_ascii=False,indent=2))
