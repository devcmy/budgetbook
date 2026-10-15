"""Independent evidence checks for the 2026-10-15 workbook update."""
from collections import Counter
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import argparse
import json
import math
import openpyxl
from lxml import etree as ET


def verify(baseline, candidate, source, summary, previous_audit, output):
    b=openpyxl.load_workbook(baseline,data_only=False)
    bd=openpyxl.load_workbook(baseline,data_only=True)
    w=openpyxl.load_workbook(candidate,data_only=False)
    d=openpyxl.load_workbook(candidate,data_only=True)
    s=json.loads(Path(summary).read_text())
    assert w.sheetnames==b.sheetnames
    for ws in w:
        assert ws.sheet_state==b[ws.title].sheet_state
        assert {str(x) for x in ws.merged_cells.ranges}=={str(x) for x in b[ws.title].merged_cells.ranges}
        for row in b[ws.title]:
            for c in row:
                target=ws[c.coordinate]
                assert getattr(c._style,'borderId',0)==getattr(target._style,'borderId',0),(ws.title,c.coordinate,'border')
    with ZipFile(baseline) as z0,ZipFile(candidate) as z1:
        assert z0.namelist()==z1.namelist()
        preserved=[n for n in z0.namelist() if n!='xl/workbook.xml' and n!='xl/styles.xml' and not n.startswith('xl/worksheets/sheet')]
        for name in preserved:assert z0.read(name)==z1.read(name),name
        vba_same=z0.read('xl/vbaProject.bin')==z1.read('xl/vbaProject.bin')
        assert vba_same
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in z1.read('[Content_Types].xml')
        assert z1.testzip() is None
    incoming=openpyxl.load_workbook(source,data_only=True).active
    assert incoming.max_row-1==143 and d['RAW'].max_row==2465
    mapping={x['source_row']:x['workbook_row'] for x in s['source_mapping']}
    assert len(mapping)==143 and len(set(mapping.values()))==143
    for sr,rr in mapping.items():
        assert incoming.cell(sr,1).value==d['RAW'][f'A{rr}'].value
        assert incoming.cell(sr,6).value==d['RAW'][f'G{rr}'].value
        assert incoming.cell(sr,7).value==d['RAW'][f'H{rr}'].value
    for r in range(2,2454):
        for c in list(range(1,14))+[18,19]:
            assert bd['RAW'].cell(r,c).value==d['RAW'].cell(r,c).value,(r,c,'baseline data')
    for month in range(1,10):
        for name in [f'{month}월',f'{month}월계산기']:
            for row in bd[name]:
                for c in row:
                    new=d[name][c.coordinate].value
                    if c.value is None and new in [None,'']:continue
                    if isinstance(c.value,(int,float)) and not isinstance(c.value,bool) and isinstance(new,(int,float)) and not isinstance(new,bool):
                        assert math.isclose(c.value,new,abs_tol=1e-6),(name,c.coordinate,c.value,new)
                    else:assert c.value==new,(name,c.coordinate,c.value,new)
    def rv(c,r):return d['RAW'][f'{c}{r}'].value
    scope=Counter(rv('R',r) for r in range(2323,2466))
    assert scope==Counter({'WRAP':13,'카드사용':85,'타계좌':45})
    ledger_rows=[r for r in range(64,501) if isinstance(d['잔액'][f'A{r}'].value,datetime) and d['잔액'][f'A{r}'].value.month==10]
    assert len(ledger_rows)==13
    links=Counter(d['잔액'][f'{col}{r}'].value for r in ledger_rows for col in ['W','X'] if d['잔액'][f'{col}{r}'].value is not None)
    assert set(links)=={r for r in range(2323,2466) if rv('R',r)=='WRAP'}
    assert all(n==1 for n in links.values())
    for row in ledger_rows:
        assert d['잔액'][f'I{row}'].value=='OK'
        rr=d['잔액'][f'W{row}'].value
        assert d['잔액'][f'E{row}'].value==rv('G',rr)
    lifestyle=[r for r in range(2323,2466) if rv('H',r)=='지출' and rv('L',r)=='본인' and rv('J',r) in {'변동비','담배','일회성전자','경조·선물','여행','약속·만남'}]
    actual=sum(rv('K',r) for r in lifestyle)
    tetris=sum(math.ceil(rv('K',r)/1000)*1000 for r in lifestyle)
    assert len(lifestyle)==75 and actual==1299740 and tetris==1321000
    calc=w['10월계산기'];palette={calc[f'B{r}'].fill.fgColor.rgb:r for r in range(3,11)};counts=Counter();order=[]
    for r0 in range(15,145,5):
        for c0 in range(2,22,2):
            for c in range(c0,c0+2):
                for r in range(r0,r0+5):
                    cell=calc.cell(r,c);filled=cell.fill.patternType=='solid';order.append(filled)
                    if filled:counts[palette[cell.fill.fgColor.rgb]]+=1
    assert order==[True]*1321+[False]*(len(order)-1321)
    for row in range(3,11):
        assert counts[row]==d['10월계산기'][f'C{row}'].value==d['10월계산기'][f'AA{row}'].value
    expected={'C37':0,'C40':9283,'C41':509283,'H37':470022,'H38':1299740,'I38':1321000,'H44':1818762,'H45':-1309479,'E74':762073,'E80':103600,'E83':0,'E84':0,'E85':0,'D14':63000,'D15':62400,'E14':20000,'E15':18900}
    for addr,value in expected.items():assert d['10월'][addr].value==value,(addr,d['10월'][addr].value,value)
    assert d['잔액']['G16'].value==6624466
    assert d['잔액']['G7'].value==4474466
    assert d['잔액']['G17'].value==2150000
    assert d['10월계산기']['W12'].value==1321000
    assert d['10월계산기']['W13'].value==1321000
    errors=[];false=[]
    for sheet in d:
        for row in sheet:
            for cell in row:
                if cell.data_type=='e':errors.append((sheet.title,cell.coordinate,cell.value))
                if sheet.title=='검증' and cell.value is False:false.append(cell.coordinate)
    assert not errors,errors[:30]
    assert not false,false
    for address in ['U28','U29','U30','U31']:assert d['10월'][address].value is True
    for address in ['O149','O150']:assert d['10월계산기'][address].value is True
    audit=json.loads(Path(previous_audit).read_text())
    audit['as_of']='2026-10-15'
    audit['workbook_sha256']=sha256(Path(candidate).read_bytes()).hexdigest()
    audit['raw']['count']=2464;audit['raw']['end']=max(rv('A',r) for r in range(2323,2466)).isoformat()
    audit['raw']['by_month']['2026-10']=143;audit['raw']['october_scope']=dict(scope)
    audit['october_ledger'].update(rows=13,raw_link_missing=[],raw_link_duplicates={},unexpected_raw_links=[],invalid_rows=[],amount_mismatches=[])
    audit['wrap'].update(total=6624466,inflows=635602,outflows=2549739)
    audit['wrap']['balances']['생활비']=4474466
    for result in audit['calculators']:
        if result['month']==10:result.update(color_total=1321000,counts=dict(counts),display=1321000,salary_burden=1321000)
    audit['october_block_layout'].update(colored_cells=1321,full_blocks=132,partial_block_cells=1)
    audit['october_lifestyle'].update(records=75,actual=actual,tetris=tetris,calendar_actual=d['10월']['I33'].value,calendar_tetris=d['10월']['M28'].value)
    audit['october_summary']={a:d['10월'][a].value for a in audit['october_summary'] if a!='fixed_independent'}
    audit['october_summary']['fixed_independent']=sum(d['10월'][f'Q{r}'].value or 0 for r in range(4,25) if d['10월'][f'U{r}'].value in ['확정','결제'])
    audit['vba'].update(binary_identical_to_previous_version=vba_same,executed=False)
    audit['delta_20261015']={k:v for k,v in s.items() if k not in ['source_mapping']}
    audit['update_checks']={'original_131_transactions_preserved':True,'past_month_values_unchanged':True,'merged_cells_and_existing_borders_preserved':True,'original_zip_parts_preserved':len(preserved),'vba_and_buttons_byte_preserved':True,'independent_colors_and_accounting_match':True,'formula_errors':errors,'false_validation_cells':false,'recalculation':'LibreOffice macro-free temporary copy; original XLSM package and formulas preserved; values imported only','excel_vba_executed':False}
    Path(output).write_text(json.dumps(audit,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'validated':True,'october_rows':143,'personal_actual':actual,'tetris':tetris,'wrap':6624466,'advance_outstanding':103600,'formula_errors':errors,'false_checks':false,'preserved_parts':len(preserved)},ensure_ascii=False))


if __name__=='__main__':
    cli=argparse.ArgumentParser()
    for name in ['baseline','candidate','source','summary','previous-audit','output']:cli.add_argument('--'+name,required=True)
    args=cli.parse_args()
    verify(args.baseline,args.candidate,args.source,args.summary,args.previous_audit,args.output)
