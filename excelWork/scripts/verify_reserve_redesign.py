"""Recalculate disposable fixtures and independently check reserve accounting."""
import argparse
from copy import copy
from collections import Counter
from datetime import datetime
import gc
from hashlib import sha256
import json
import math
import os
from pathlib import Path
import subprocess
from zipfile import ZipFile

import openpyxl
from lxml import etree as ET

from redesign_reserve import ExtendedPackage, FUNDS, BACK, U0, U1, S0, S1
from update_20261015 import macro_free, Q


def read(path):
    return openpyxl.load_workbook(path,data_only=True)


def recalculate(candidate, folder, edit=None):
    folder=Path(folder);folder.mkdir()
    (folder/'out').mkdir();(folder/'cache').mkdir()
    p=ExtendedPackage(candidate)
    if edit:edit(p)
    p.save(folder/'input.xlsm')
    macro_free(folder/'input.xlsm',folder/'input.xlsx')
    result=subprocess.run(['soffice','-env:UserInstallation='+(folder/'profile').as_uri(),
                           '--headless','--convert-to','xlsx:Calc MS Excel 2007 XML',
                           '--outdir',str(folder/'out'),str(folder/'input.xlsx')],
                          env={**os.environ,'XDG_CACHE_HOME':str(folder/'cache')},
                          capture_output=True,text=True,timeout=120)
    out=folder/'out/input.xlsx'
    assert result.returncode==0 and out.is_file(),(result.returncode,result.stdout,result.stderr)
    return out


def assert_no_errors(w):
    errors=[(s.title,c.coordinate,c.value) for s in w for row in s for c in row if c.data_type=='e']
    assert not errors,errors[:20]


def end_date(month):
    return datetime(2027,1,1) if month==12 else datetime(2026,month+1,1)


def verify_month(w, month):
    s=w['예비비'];ledger=w['잔액'];start=datetime(2026,month,1);end=end_date(month)
    cash={fund:0 for fund in FUNDS};due={fund:0 for fund in FUNDS};used={fund:0 for fund in FUNDS}
    for r in range(64,501):
        date=ledger[f'A{r}'].value
        if not isinstance(date,datetime) or date>=end:continue
        amount=ledger[f'Y{r}'].value or 0
        sender,recipient=ledger[f'C{r}'].value,ledger[f'D{r}'].value
        if recipient in cash:cash[recipient]+=amount
        if sender in cash:cash[sender]-=amount
    for r in range(U0,U1+1):
        if s[f'R{r}'].value!='OK':continue
        uid,category,date,amount=[s[f'{c}{r}'].value for c in ['A','C','D','E']]
        paid=sum(s[f'D{x}'].value for x in range(S0,S1+1)
                 if s[f'B{x}'].value==uid and s[f'N{x}'].value=='OK' and s[f'E{x}'].value<end)
        if date<end and month>=3:due[category]+=max(0,amount-paid)
        if start<=date<end:used[category]+=amount
    for r,fund in enumerate(FUNDS,8):
        if month>=3:
            assert s[f'C{r}'].value==cash[fund],(month,fund,'cash',s[f'C{r}'].value,cash[fund])
            assert s[f'F{r}'].value==due[fund],(month,fund,'due')
            assert s[f'I{r}'].value==cash[fund]-due[fund],(month,fund,'available')
            assert s[f'N{r}'].value=='일치',(month,fund,'balance tab')
            assert s[f'C{r}'].value==ledger[f'G{r}'].value
        else:
            assert s[f'C{r}'].value=='원장 없음'
            assert s[f'F{r}'].value=='원장 없음'
            assert s[f'I{r}'].value=='원장 없음'
        assert s[f'L{r}'].value==used[fund],(month,fund,'monthly use')
        assert w[f'{month}월'][f'N{r+56}'].value==used[fund],(month,fund,'monthly link')
    for r in range(193,197):assert w['검증'][f'M{r}'].value==1,(month,r,w['검증'][f'I{r}'].value)
    assert s['L13'].value==sum(used.values())
    return {'month':month,'ledger_total':sum(cash.values()) if month>=3 else None,
            'outstanding':sum(due.values()) if month>=3 else None,'actual_use':sum(used.values())}


def cache_original(candidate, recalculated, output, baseline):
    p=ExtendedPackage(candidate);values=read(recalculated);before=read(baseline)
    bools={(s.title,c.coordinate) for s in before for row in s for c in row if isinstance(c.value,bool)}
    bools.update((BACK,c.coordinate) for row in before['예비비'] for c in row if isinstance(c.value,bool))
    bools.update(('검증',f'M{r}') for r in range(193,197))
    count=0
    for sheet in p.sheets:
        for cell in p.sheet(sheet).iter(Q('c')):
            if cell.find(Q('f')) is None:continue
            addr=cell.get('r');value=values[sheet][addr].value
            if (sheet,addr) in bools and value in [0,1]:value=bool(value)
            assert values[sheet][addr].data_type!='e',(sheet,addr,value)
            for child in list(cell):
                if child.tag in [Q('v'),Q('is')]:cell.remove(child)
            cell.attrib.pop('t',None)
            if value is not None:
                from openpyxl.utils.datetime import to_excel
                if isinstance(value,datetime):value=to_excel(value)
                if isinstance(value,bool):cell.set('t','b');text='1' if value else '0'
                elif isinstance(value,(int,float)):text=str(value)
                else:cell.set('t','str');text=str(value)
                ET.SubElement(cell,Q('v')).text=text
            count+=1
    p.save(output)
    return count


def verify_artifact(baseline, final):
    before=read(baseline);after=read(final)
    b=openpyxl.load_workbook(baseline);a=openpyxl.load_workbook(final)
    assert after.sheetnames==before.sheetnames+[BACK]
    assert [s.title for s in after if s.sheet_state=='visible']==[s.title for s in before if s.sheet_state=='visible']
    assert after[BACK].sheet_state=='hidden'
    for name in ['RAW','잔액']:
        for row in before[name]:
            for c in row:
                actual=after[name][c.coordinate].value
                if c.value in [None,''] and actual in [None,'']:continue
                if isinstance(c.value,(int,float)) and isinstance(actual,(int,float)):
                    assert math.isclose(c.value,actual,abs_tol=1e-6),(name,c.coordinate,c.value,actual)
                else:assert c.value==actual,(name,c.coordinate,c.value,actual)
        assert str(a[name].data_validations)==str(b[name].data_validations)
        assert set(map(str,a[name].merged_cells.ranges))==set(map(str,b[name].merged_cells.ranges))
    for m in range(1,13):
        name=f'{m}월'
        for addr in ['C37','C38','C39','C40','C41','H37','I37','H38','I38','H40','I40','H41','I41','H44','I44','H45','I45','N53']:
            x,y=before[name][addr].value,after[name][addr].value
            assert x==y,(name,addr,x,y)
        for addr in ['W12','W13','W14','O149','O150']:
            assert before[f'{m}월계산기'][addr].value==after[f'{m}월계산기'][addr].value,(m,addr)
    for sheet in b:
        if sheet.title=='예비비':continue
        for row in sheet:
            for c in row:
                target=a[sheet.title][c.coordinate]
                if c._style!=target._style:
                    # Appending styles can cause openpyxl's merged-cell border
                    # cache to choose an equivalent border at a different ID.
                    for attr in ['font','fill','border','alignment','number_format','protection']:
                        assert copy(getattr(c,attr))==copy(getattr(target,attr)),(sheet.title,c.coordinate,attr)
        assert set(map(str,sheet.merged_cells.ranges))==set(map(str,a[sheet.title].merged_cells.ranges))
    for c in [c for row in after['검증'] for c in row if isinstance(c.value,bool)]:assert c.value is True,c.coordinate
    assert_no_errors(after)
    source_map={after['예비비'][f'F{r}'].value:r for r in range(U0,U1+1) if after['예비비'][f'F{r}'].value}
    for rawrow in [2379,2380,2433,2459]:assert after['예비비'][f'C{source_map[rawrow]}'].value=='생활비'
    assert after['예비비']['C13'].value==2150000
    assert after['예비비']['F13'].value==0
    assert after['예비비']['I13'].value==2150000
    assert after['예비비']['B36'].value=='부모 생신 본인 부담'
    assert after['예비비']['A52'].value=='2026-09'
    # Every 2026 plan cell is preserved and connected to its legacy contract.
    for r,oldr in enumerate(range(5,30,3),20):
        for m in range(1,13):
            assert after['예비비'].cell(r,m+3).value==(before['예비비'].cell(oldr,m+7).value or 0)
    with ZipFile(baseline) as z0,ZipFile(final) as z1:
        assert z0.read('xl/vbaProject.bin')==z1.read('xl/vbaProject.bin')
        for n in z0.namelist():
            if n.startswith(('xl/drawings/','xl/media/','xl/ctrlProps/','xl/activeX/')):
                assert z0.read(n)==z1.read(n),n
            if n.endswith('.rels') and n!='xl/_rels/workbook.xml.rels':assert z0.read(n)==z1.read(n),n
        assert z1.testzip() is None
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in z1.read('[Content_Types].xml')
        assert 'xl/calcChain.xml' not in z1.namelist()
    return {'raw_and_balance_values_unchanged':True,'monthly_living_and_salary_totals_unchanged':True,
            'plans_preserved':True,'original_styles_outside_reserve_preserved':True,
            'vba_buttons_drawings_preserved':True,'visible_sheet_order_unchanged':True}


def run(baseline,candidate,output,workdir,report):
    root=Path(workdir);root.mkdir(parents=True,exist_ok=True)
    f=openpyxl.load_workbook(candidate,data_only=False)
    rows={f['예비비'][f'F{r}'].value:r for r in range(U0,U1+1) if isinstance(f['예비비'][f'F{r}'].value,int)}
    del f;gc.collect()
    def query(p,month):
        p.write('예비비','B3',datetime(2026,month,1));p.write('잔액','B3',datetime(2026,month,1))
    months=[];baseline_recalc=None
    for month in range(1,13):
        path=recalculate(candidate,root/f'month_{month:02d}',lambda p,m=month:query(p,m))
        w=read(path);assert_no_errors(w)
        months.append(verify_month(w,month))
        if month==10:baseline_recalc=path
        print(json.dumps({'month_passed':month,**months[-1]},ensure_ascii=False),flush=True)
        del w;gc.collect()
    def choose(p,raws):
        for raw in raws:p.write('예비비',f'C{rows[raw]}','여행비')
    def transfer(p,amount,date):
        vals={'A246':date,'B246':'이체','C246':'여행비','D246':'외부','E246':amount,
              'F246':'검증용 정산 이체','G246':'임시 테스트에만 존재','J246':'미래에셋_WRAP형CMA',
              'K246':'검증용 카드계좌','L246':'WRAP계좌이체'}
        for addr,v in vals.items():p.write('잔액',addr,v)
    def link(p,row,raw,amount):
        for c,v in [('B','R'+str(raw)),('C',246),('D',amount)]:p.write('예비비',f'{c}{row}',v)
    cases=[]
    def check_case(name,edit,check,expected_warning=False):
        path=recalculate(candidate,root/name,edit);w=read(path);assert_no_errors(w);check(w)
        if not expected_warning:
            for r in range(193,197):assert w['검증'][f'M{r}'].value==1,(name,r,w['검증'][f'I{r}'].value)
        cases.append({'name':name,'passed':True,'expected_warning':expected_warning})
        print(json.dumps(cases[-1]),flush=True);del w;gc.collect()
    def amount(w,addr,value):assert w['예비비'][addr].value==value,(addr,w['예비비'][addr].value,value)
    def purchase(p):query(p,10);choose(p,[2379])
    def purchase_check(w):
        amount(w,'C9',1400000);amount(w,'F9',39000);amount(w,'I9',1361000);amount(w,'L9',39000)
        assert w['10월']['N65'].value==39000
        assert w['10월']['H38'].value==1299740-39000
        assert w['10월']['I38'].value==1321000-39000
        assert w['10월계산기']['W13'].value==1321000-39000
    check_case('card_use_before_transfer',purchase,purchase_check)
    def payment_first(p):
        purchase(p)
        # A card payment from another bank is not a reserve funding transfer,
        # even if its amount coincides with the use being reserved.
        p.write('RAW','G2455',39000);p.write('RAW','K2455',39000)
    check_case('card_paid_elsewhere_still_requires_reserve_funding',payment_first,purchase_check)
    def settle(p,amount_=39000):
        query(p,11);choose(p,[2379]);transfer(p,amount_,datetime(2026,11,5));link(p,440,2379,amount_)
    def full(w):
        amount(w,'C9',1361000);amount(w,'F9',0);amount(w,'I9',1361000);amount(w,'L9',0)
        assert w['10월']['N65'].value==39000 and w['11월']['N65'].value==0
    check_case('next_month_settlement',settle,full)
    def future_funding(p):settle(p);query(p,10)
    def future_funding_check(w):
        purchase_check(w)
        assert not any(w['예비비'][f'A{r}'].value=='2026-11-05' for r in range(52,64))
    check_case('future_funding_not_visible_in_prior_month',future_funding,future_funding_check)
    def partial(w):amount(w,'C9',1380000);amount(w,'F9',19000);amount(w,'I9',1361000)
    check_case('partial_settlement',lambda p:settle(p,20000),partial)
    def prepay(p):
        query(p,9);choose(p,[2379]);transfer(p,39000,datetime(2026,9,25));link(p,440,2379,39000)
    def advance(w):
        amount(w,'C9',1361000);amount(w,'F9',0);amount(w,'I9',1361000);amount(w,'L9',0)
        assert any(w['예비비'][f'I{r}'].value=='선이체' for r in range(36,48))
    check_case('prepayment_before_use',prepay,advance)
    def prepay_used(p):prepay(p);query(p,10)
    def advanced_used(w):amount(w,'C9',1361000);amount(w,'F9',0);amount(w,'I9',1361000);amount(w,'L9',39000)
    check_case('prepayment_then_use',prepay_used,advanced_used)
    def bulk(p):
        query(p,11);choose(p,[2379,2380]);transfer(p,74000,datetime(2026,11,5));link(p,440,2379,39000);link(p,441,2380,35000)
    def bulk_check(w):amount(w,'C9',1326000);amount(w,'F9',0);amount(w,'I9',1326000)
    check_case('one_transfer_multiple_uses',bulk,bulk_check)
    def duplicate(p):settle(p);link(p,441,2379,39000)
    def duplicate_check(w):
        amount(w,'F9',39000);assert w['예비비']['N440'].value.startswith('확인');assert w['검증']['I193'].value>=2
    check_case('duplicate_settlement_blocked',duplicate,duplicate_check,True)
    def duplicate_use(p):
        purchase(p);p.write('예비비','F300',2379);p.write('예비비','C300','여행비')
    def duplicate_use_check(w):
        assert w['예비비'][f'R{rows[2379]}'].value.startswith('확인')
        assert w['예비비']['R300'].value.startswith('확인')
    check_case('duplicate_raw_use_blocked',duplicate_use,duplicate_use_check,True)
    def monthend(p):purchase(p);p.write('RAW','A2379',datetime(2026,10,31,23,59,59))
    def last_day_check(w):
        # This fixture moves a RAW date, not the manual calendar/paint history.
        # Check reserve period boundaries independently of that separate import.
        amount(w,'C9',1400000);amount(w,'F9',39000);amount(w,'I9',1361000);amount(w,'L9',39000)
        assert w['10월']['N65'].value==39000
        assert not any(w['예비비'][f'I{r}'].value=='선이체' for r in range(36,48))
    check_case('month_end_timestamp_included',monthend,last_day_check)
    def plan(p):query(p,11);p.write('예비비','N25',150000)
    def plan_check(w):
        amount(w,'C9',1400000);assert w['11월']['M49'].value==150000;assert w['11월']['M53'].value==250000
    check_case('plan_edit_links_monthly_without_creating_cash',plan,plan_check)
    def pagination(p):
        query(p,10)
        for raw,row in rows.items():
            if raw in [2354,2367,2379,2380,2399,2414,2418,2433,2459,2312,2317,2318]:
                p.write('예비비',f'C{row}','여행비')
        p.write('예비비','B33','전체');p.write('예비비','B34',2)
    def page_check(w):
        assert w['예비비']['AB36'].value>0
        assert '페이지 2' in w['예비비']['D34'].value
    check_case('pagination_has_no_silent_truncation',pagination,page_check)
    def stale_page(p):query(p,10);p.write('예비비','B34',99);p.write('예비비','B50',99)
    def stale_page_check(w):
        assert w['예비비']['AC34'].value==1 and w['예비비']['AB36'].value>0
        assert w['예비비']['AC50'].value==1 and w['예비비']['AB52'].value>0
    check_case('stale_page_is_clamped_after_month_change',stale_page,stale_page_check)
    def insufficient(p):purchase(p);p.write('예비비',f'C{rows[2379]}','명절비')
    def insufficient_check(w):
        amount(w,'C8',0);amount(w,'F8',39000);amount(w,'I8',-39000)
        assert '부족' in w['예비비']['A15'].value
    check_case('insufficient_fund_not_hidden_or_auto_reclassified',insufficient,insufficient_check)
    count=cache_original(candidate,baseline_recalc,output,baseline)
    artifact=verify_artifact(baseline,output)
    reportdata={'as_of':'2026-10-15','default_month':'2026-10','month_tests':months,
                'scenario_tests':cases,'checks':artifact,'cached_formulas':count,
                'formula_errors':0,'visible_sheets':29,'total_sheets':46,
                'wrap_purpose_cash':2150000,'outstanding_reserve_settlement':0,
                'confirmed_living_travel_amount':181800,
                'workbook_sha256':sha256(Path(output).read_bytes()).hexdigest(),
                'excel_vba_executed':False,
                'recalculation':'Disposable macro-free XLSX copies in LibreOffice; original XLSM formulas/package retained.'}
    Path(report).write_text(json.dumps(reportdata,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'verified':True,'months':len(months),'scenarios':len(cases),'hash':reportdata['workbook_sha256']},ensure_ascii=False),flush=True)


def targeted_recheck(baseline,candidate,previous_candidate,previous_report,output,workdir,report):
    """Reuse the full suite only after proving the sole delta is display ranking.

    Financial formulas, inputs, styles, controls and all other XML parts must be
    byte-identical. Recalculate the new default and the changed temporal case.
    """
    old=ExtendedPackage(previous_candidate);new=ExtendedPackage(candidate)
    assert set(old.parts)==set(new.parts)
    allowed=new.sheets[BACK]
    changed=[]
    for part in old.parts:
        if part!=allowed:assert old.parts[part]==new.parts[part],part
    oldtree=old.sheet(BACK);newtree=new.sheet(BACK)
    oldcells={c.get('r'):c for c in oldtree.iter(Q('c'))}
    for c in newtree.iter(Q('c')):
        addr=c.get('r');prior=oldcells[addr]
        if ET.tostring(c)!=ET.tostring(prior):
            assert addr.startswith('BH') and int(addr[2:]) in range(120,557),addr
            changed.append(addr)
            c.find(Q('f')).text=prior.find(Q('f')).text
    assert changed and ET.tostring(oldtree)==ET.tostring(newtree)
    del old,new;gc.collect()
    root=Path(workdir);root.mkdir(parents=True,exist_ok=True)
    original=openpyxl.load_workbook(candidate)
    row=next(r for r in range(U0,U1+1) if original['예비비'][f'F{r}'].value==2379)
    del original;gc.collect()
    default=recalculate(candidate,root/'default')
    w=read(default);assert_no_errors(w);verify_month(w,10);del w;gc.collect()
    def fixture(p):
        p.write('예비비',f'C{row}','여행비')
        vals={'A246':datetime(2026,11,5),'B246':'이체','C246':'여행비','D246':'외부','E246':39000,
              'F246':'검증용 미래 정산','J246':'미래에셋_WRAP형CMA','K246':'검증용 카드계좌','L246':'WRAP계좌이체'}
        for a,v in vals.items():p.write('잔액',a,v)
        for c,v in [('B','R2379'),('C',246),('D',39000)]:p.write('예비비',f'{c}440',v)
    path=recalculate(candidate,root/'future_funding_visibility',fixture)
    w=read(path);assert_no_errors(w)
    assert w['예비비']['C9'].value==1400000 and w['예비비']['F9'].value==39000
    assert w['예비비']['I9'].value==1361000 and w['10월']['N65'].value==39000
    assert w['10월']['H38'].value==1260740 and w['10월계산기']['W13'].value==1282000
    assert not any(w['예비비'][f'A{r}'].value=='2026-11-05' for r in range(52,64))
    assert w['예비비']['A52'].value=='2026-09' # Existing prepayment still visible.
    for r in range(193,197):assert w['검증'][f'M{r}'].value==1
    del w;gc.collect()
    count=cache_original(candidate,default,output,baseline)
    checks=verify_artifact(baseline,output)
    result=json.loads(Path(previous_report).read_text())
    result['scenario_tests'].append({'name':'future_funding_not_visible_in_prior_month','passed':True,'expected_warning':False})
    result.update(checks=checks,cached_formulas=count,workbook_sha256=sha256(Path(output).read_bytes()).hexdigest())
    result['targeted_final_recheck']={
        'only_changed_cells':changed,'change':'Hide future funding from past-month transfer detail',
        'all_financial_formulas_inputs_and_other_parts_byte_identical':True,
        'previous_candidate_sha256':sha256(Path(previous_candidate).read_bytes()).hexdigest(),
        'default_and_temporal_case_recalculated':True,
    }
    Path(report).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'verified':True,'months':len(result['month_tests']),'scenarios':len(result['scenario_tests']),'hash':result['workbook_sha256']},ensure_ascii=False),flush=True)


if __name__=='__main__':
    p=argparse.ArgumentParser()
    for name in ['baseline','candidate','output','workdir','report']:p.add_argument('--'+name,required=True)
    p.add_argument('--previous-candidate');p.add_argument('--previous-report')
    a=p.parse_args()
    if a.previous_candidate:
        assert a.previous_report
        targeted_recheck(a.baseline,a.candidate,a.previous_candidate,a.previous_report,a.output,a.workdir,a.report)
    else:run(a.baseline,a.candidate,a.output,a.workdir,a.report)
