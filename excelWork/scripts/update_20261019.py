"""Import seven new source transactions without reverting confirmed decisions.

Keep the first 164 October records, workbook row identities, KakaoPay cash-out
classification, macros and original uploads. Paint six card purchases; show
the housing guarantee premium as an additional cash cost, not recurring cover.
"""
import argparse
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime
import gc
from hashlib import sha256
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter as col

from balance_selected_ledger import check_view
from redesign_reserve import ExtendedPackage
from update_20261015 import Q, CATEGORIES
from update_cash_details_20261019 import import_changed_caches, paint_order
from validate_excel_structure import validate_package
from verify_reserve_cash_alignment import independent_balances, verify_month
from verify_reserve_redesign import recalculate, assert_no_errors

ASOF = datetime(2026, 10, 19)
START, OLD_END, END = 2323, 2486, 2493
ACTUAL, TETRIS = 1582350, 1607000


def digest(path): return sha256(Path(path).read_bytes()).hexdigest()
def normal(row): return tuple(None if v == '' else v for v in row)
def asset(value): return '모임통장' if value == '20260606' else value
def identity(row): return row[0], asset(row[1]), row[5], row[6]


def prepare(baseline, source, previous, output):
    w = openpyxl.load_workbook(baseline)
    d = openpyxl.load_workbook(baseline, data_only=True)
    incoming = list(openpyxl.load_workbook(source, data_only=True).active.values)
    old_source = list(openpyxl.load_workbook(previous, data_only=True).active.values)
    assert incoming[0] == old_source[0] == ('날짜','자산','분류','소분류','내용','금액(원)','수입/지출','메모','금액','화폐')
    prior, latest = Counter(map(normal, old_source[1:])), Counter(map(normal, incoming[1:]))
    assert not prior-latest and sum(prior.values()) == 164 and sum(latest.values()) == 171
    raw = d['RAW']; assert raw.max_row == OLD_END
    known = {(raw[f'A{r}'].value,raw[f'C{r}'].value,raw[f'G{r}'].value,raw[f'H{r}'].value):r for r in range(START,OLD_END+1)}
    assert len(known) == 164
    additions, mapping, matched = [], [], set()
    for sr, row in enumerate(incoming[1:], 2):
        row = normal(row); key = identity(row)
        assert row[9] == 'KRW' and row[5] == row[8] and row[0].year == 2026 and row[0].month == 10
        if key in known:
            rr = known[key]; assert rr not in matched
            matched.add(rr); mapping.append({'source_row':sr,'workbook_row':rr,'status':'preserved'})
        else: additions.append((sr,row))
    assert len(matched) == 164 and len(additions) == 7
    additions.sort(key=lambda v:(v[1][0],v[0]))
    p = ExtendedPackage(baseline)
    # Grow only the known October range; all earlier record identities remain.
    for sheet in w:
        if sheet.title.startswith(('기존_', '보관_')): continue
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f':
                    formula = re.sub(r'(RAW!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?)2486\b',lambda m:m[1]+str(END),cell.value)
                    if formula != cell.value: p.write(sheet.title,cell.coordinate,formula,True)
    personal, colors, metadata = [], Counter(), []
    for rr,(sr,row) in enumerate(additions,OLD_END+1):
        date,ac,category,sub,title,amount,kind,note,_,_ = row
        assert kind == '지출' and isinstance(amount,int) and amount > 0
        cash = ac == '미래에셋_WRAP형CMA'
        if cash:
            assert rr == 2493 and title == '보증보험료' and amount == 110075
            classification, scope, direction = '이사부대','WRAP','출금'
            memo = '10/19 첨부 RAW 2행. 9월 동일 보증보험료의 이사부대 분류 유지. 연간 주거부대 현금 지출로 추가 실지출 표에 1회 반영; 월 정기보험비·예비비와 분리.'
        else:
            assert '카드' in ac or ac == 'KB쿠팡와우'
            classification, scope, direction = '변동비','카드사용',None
            palette = {'식비':3,'기타':4,'패션/미용':6,'건강/보험/의료':7,'잡화/소품/가구/생필품':9}[category]
            units = math.ceil(amount/1000); colors[palette] += units
            personal.append({'row':rr,'amount':amount,'units':units,'calculator_row':palette})
            memo = f'10/19 첨부 RAW {sr}행. 본인 카드 변동지출·달력·계산기 반영. WRAP에서 재차 차감하지 않음.'
        p.clone_row('RAW',2485,rr)
        values=[date,ASOF.replace(day=1),ac,category,sub,title,amount,kind,note,classification,amount,'본인',memo,None,None,None,None,scope,direction]
        for c,value in enumerate(values,1):p.write('RAW',f'{col(c)}{rr}',value)
        link='(COUNTIF(잔액!$W$64:$W$500,ROW())+COUNTIF(잔액!$X$64:$X$500,ROW()))'
        p.write('RAW',f'N{rr}',f'=IF(R{rr}="WRAP",IF({link}=1,"연결",IF({link}=0,"누락","중복")),IF(R{rr}="카드사용","월간·계산기","타계좌·잔액제외"))',True)
        mapping.append({'source_row':sr,'workbook_row':rr,'status':'added'})
        metadata.append({'row':rr,'source_row':sr,'date':date.isoformat(),'asset':ac,'title':title,'amount':amount,'classification':classification})
    assert sum(x['amount'] for x in personal)==105450 and sum(colors.values())==106
    for name in ['dimension','autoFilter']:p.sheet('RAW').find(Q(name)).set('ref',f'A1:S{END}')
    for cf in p.sheet('RAW').findall(Q('conditionalFormatting')):
        cf.set('sqref',cf.get('sqref').replace('N2323:N2486',f'N2323:N{END}'))
    p.write('RAW','Q1','10월 RAW 갱신 · 2026-10-19 기준')
    p.write('RAW','Q2','10/1~19 RAW 171건: 기존 164건 유지 + 신규 7건. WRAP 16건·타계좌 48건·카드사용 107건.')
    # Existing blank ledger row already owns shared normalization formulas.
    assert d['잔액']['A249'].value is None
    inputs={'A249':additions[-1][1][0],'B249':'출금','C249':'생활비','D249':'지출','E249':110075,
            'F249':'보증보험료','G249':'RAW 2493행. 계약기간 2026.09.10~2027.09.09의 주거부대 비용. WRAP 출금 110,075원 1회, 10월 추가 실지출 표에 1회 반영. 달력·월 정기보험비·예비비와 중복 제외.',
            'J249':'미래에셋_WRAP형CMA','K249':'외부','L249':'WRAP입출금','W249':2493}
    for addr,value in inputs.items():p.write('잔액',addr,value)
    # Additional cash cost: the shared F/H formulas remain untouched.
    assert d['10월']['C54'].value is None
    for addr,value in {'B54':'현금','C54':'보증보험료','I54':'연간 주거부대비. 달력 제외·추가 실지출에 1회 반영.'}.items():p.write('10월',addr,value)
    covered='+'.join(f'(RAW!J2493="{x}")' for x in sorted(CATEGORIES))
    for addr,formula in {'A54':'=INT(RAW!A2493)','D54':'=RAW!K2493',
                         'G54':f'=IF(AND(RAW!H2493="지출",RAW!L2493="본인",({covered})>0),RAW!K2493,0)'}.items():
        p.write('10월',addr,formula,True)
    for c in ['C','I']:p.cell('10월',f'{c}54').set('s',p.cell('10월',f'{c}52').get('s'))
    p.row_maps['10월'][54].set('ht','44');p.row_maps['10월'][54].set('customHeight','1')
    p.write('10월','U1',ASOF)
    for dr in [3,8,13,18,23,28]:
        for c in 'BCDEFGH':
            date=d['10월'][f'{c}{dr}'].value
            if not isinstance(date,datetime) or date.month!=10 or date.day not in [18,19]:continue
            p.write('10월',f'{c}{dr+2}',Translator(w['10월']['E15'].value,origin='E15').translate_formula(f'{c}{dr+2}'),True)
            rr=f'RAW!$A${START}:$A${END}';cl=f'RAW!$J${START}:$J${END}'
            terms='+'.join(f'({cl}="{x}")' for x in sorted(CATEGORIES))
            p.write('10월',f'{c}{dr+1}',f'=SUMPRODUCT(({rr}>={c}{dr})*({rr}<{c}{dr}+1)*(RAW!$H${START}:$H${END}="지출")*(RAW!$L${START}:$L${END}="본인")*({terms})*ROUNDUP(RAW!$K${START}:$K${END}/1000,0)*1000)',True)
    for addr,formula in {'E83':f'=COUNT(RAW!G{START}:G{END})-171','E84':f'=I33-{ACTUAL}',
                         'E85':f'=SUM(I4,I9,I14,I19,I24,I29)-{TETRIS}'}.items():p.write('10월',addr,formula,True)
    for addr,value in {'A70':'10/1~19 RAW 171건. 기존 164건 보존 + 카드 6건 105,450원·보증보험료 현금 110,075원 추가.',
                       'A72':'10/19 실제 / 예정 / 공동자금 · 달력 합계와 연결',
                       'J83':'0이면 최신 10/19 첨부 RAW 171건 모두 유지. 확정 분류·행번호 보존.',
                       'J86':'급여 미입금: 10/19 중간 집계이며 월말 결과가 아님'}.items():p.write('10월',addr,value)
    p.write('10월','K32',w['10월']['K32'].value.replace('10/13까지','10/19까지').replace('10/18까지','10/19까지'))
    for r in range(4,25):
        text=w['10월'][f'V{r}'].value
        if text:p.write('10월',f'V{r}',text.replace('10/18 기준','10/19 기준').replace('10/18 RAW 집계','10/19 RAW 집계'))
    calc=w['10월계산기'];order=paint_order()
    assert [calc[a].fill.patternType=='solid' for a in order]==[True]*1501+[False]*(len(order)-1501)
    xfs=p.tree('xl/styles.xml').find(Q('cellXfs'));cache={ET.tostring(x):i for i,x in enumerate(xfs)};idx=1501
    for item in personal:
        for _ in range(item['units']):
            cell=p.cell('10월계산기',order[idx]);style=deepcopy(xfs[int(cell.get('s','0'))])
            style.set('fillId',xfs[int(p.cell('10월계산기',f'B{item["calculator_row"]}').get('s'))].get('fillId'))
            key=ET.tostring(style)
            if key not in cache:cache[key]=len(xfs);xfs.append(style)
            cell.set('s',str(cache[key]));idx+=1
    xfs.set('count',str(len(xfs)))
    for r in range(3,11):p.write('10월계산기',f'AA{r}',d['10월계산기'][f'AA{r}'].value+colors[r])
    p.write('10월계산기','A154','10/19까지 카드 변동지출 6건·106칸 추가. 보증보험료 현금 110,075원은 월간 추가 실지출에서 별도 집계하며 색상 칸에 중복 합산하지 않음.')
    p.write('잔액','D3','2026-10-19 기준 · WRAP RAW 16건 연결. 카카오페이 충전·보증보험료는 출금에 1회 반영. 카드 사용은 WRAP 중복 제외.')
    p.write('사용안내','A22','10월은 10/19까지 실제 자료입니다. 미수령 급여·미결제 고정비·공동자금·대납은 실제 소비와 구분합니다.')
    p.write('사용안내','A23','10월 RAW 171건 중 WRAP 16건을 잔액 원장과 대조합니다. 타계좌 48건·카드사용 107건 보존. 보증보험료는 추가 실지출에 1회 반영합니다.')
    p.trees.pop('xl/workbook.xml');p.save(output)
    w.close();d.close();gc.collect()
    return {'new_rows':metadata,'new_personal':personal,'source_mapping':sorted(mapping,key=lambda v:v['source_row'])}


def query_check(w,month):
    s=w['잔액'];end=datetime(2027,1,1) if month==12 else datetime(2026,month+1,1)
    balances,total,_=independent_balances(w,end)
    for r in range(7,16):assert s[f'G{r}'].value==balances[s[f'A{r}'].value]
    assert s['G16'].value==total
    return {'month':month,'wrap_total':total,'view':check_view(w,month)}


def verify(baseline,final,source,metadata):
    b=openpyxl.load_workbook(baseline,data_only=True);w=openpyxl.load_workbook(final,data_only=True)
    raw,month,calc,ledger=[w[n] for n in ['RAW','10월','10월계산기','잔액']]
    assert raw.max_row==END and w.sheetnames==b.sheetnames and '검증' not in w.sheetnames
    incoming=list(openpyxl.load_workbook(source,data_only=True).active.values)[1:]
    assert Counter(identity(row) for row in incoming)==Counter((raw[f'A{r}'].value,raw[f'C{r}'].value,raw[f'G{r}'].value,raw[f'H{r}'].value) for r in range(START,END+1))
    for item in metadata['source_mapping']:
        if item['status']!='added':continue
        source_values=normal(incoming[item['source_row']-2]);rr=item['workbook_row']
        expected={'A':source_values[0],'C':asset(source_values[1]),'D':source_values[2],
                  'E':source_values[3],'F':source_values[4],'G':source_values[5],
                  'H':source_values[6],'I':source_values[7]}
        for c,value in expected.items():assert raw[f'{c}{rr}'].value==value,('source values',rr,c)
    assert raw['J2493'].value==raw['J2157'].value=='이사부대'
    variable=[r for r in range(START,END+1) if raw[f'H{r}'].value=='지출' and raw[f'L{r}'].value=='본인' and raw[f'J{r}'].value in CATEGORIES]
    assert len(variable)==98
    assert sum(raw[f'K{r}'].value for r in variable)==ACTUAL==month['I33'].value
    assert sum(math.ceil(raw[f'K{r}'].value/1000)*1000 for r in variable)==TETRIS==calc['W12'].value
    for dr in [3,8,13,18,23,28]:
        for c in range(2,9):
            date=month.cell(dr,c).value
            if not isinstance(date,datetime) or date.month!=10:continue
            rows=[r for r in variable if raw[f'A{r}'].value.date()==date.date()]
            assert (month.cell(dr+2,c).value or 0)==sum(raw[f'K{r}'].value for r in rows),(date,'actual')
            assert (month.cell(dr+1,c).value or 0)==sum(math.ceil(raw[f'K{r}'].value/1000)*1000 for r in rows),(date,'tetris')
    assert month['C54'].value=='보증보험료'
    assert [month[f'{c}54'].value for c in 'DFGH']==[110075,110075,0,110075]
    assert month['H39'].value==110075 and month['I39'].value==111000
    assert month['H44'].value==2211447 and month['I44'].value==2241000
    assert month['H53'].value==0 and month['H52'].value==0
    assert all(month[a].value==0 for a in ['E83','E84','E85'])
    assert all(bool(month[f'U{r}'].value) for r in [28,29,30,31])
    assert month['E80'].value==103600 and month['N69'].value==239000
    assert month['H37'].value==470022 and month['E74'].value==762073
    assert Counter(raw[f'R{r}'].value for r in range(START,END+1))==Counter({'WRAP':16,'타계좌':48,'카드사용':107})
    links=Counter(ledger[f'{c}{r}'].value for r in range(64,501) for c in 'WX' if ledger[f'{c}{r}'].value is not None)
    wrap_rows={r for r in range(START,END+1) if raw[f'R{r}'].value=='WRAP'}
    assert set(links)==wrap_rows and set(links.values())=={1}
    assert all(raw[f'N{r}'].value=='연결' for r in wrap_rows)
    assert ledger['B248'].value=='출금' and ledger['P248'].value=='내역'
    assert ledger['B249'].value=='출금' and ledger['W249'].value==2493 and ledger['E249'].value==110075
    assert ledger['G7'].value==4409691 and ledger['G16'].value==6559691
    assert ledger['D7'].value==137675 and ledger['E7'].value==2537839
    left=[int(ledger[f'M{r}'].value)+63 for r in range(38,60) if ledger[f'M{r}'].value]
    right=[int(ledger[f'N{r}'].value)+63 for r in range(38,60) if ledger[f'N{r}'].value]
    assert left==[249,248,245,238,242] and len(right)==12 and 248 not in right and 249 not in right
    assert ledger['F38'].value==-110075 and ledger['F39'].value==-15700
    accounts={ledger[f'A{r}'].value for r in range(7,16)}
    october=[r for r in range(64,501) if isinstance(ledger[f'A{r}'].value,datetime) and datetime(2026,10,1)<=ledger[f'A{r}'].value<datetime(2026,11,1)]
    cash_in=sum(ledger[f'Y{r}'].value for r in october if ledger[f'C{r}'].value not in accounts and ledger[f'D{r}'].value in accounts)
    cash_out=sum(ledger[f'Y{r}'].value for r in october if ledger[f'C{r}'].value in accounts and ledger[f'D{r}'].value not in accounts)
    assert cash_in==696602 and cash_out==2675514 and 8538603+cash_in-cash_out==6559691
    palette={calc[f'B{r}'].fill.fgColor.rgb:r for r in range(3,11)};counts=Counter();painted=[]
    for a in paint_order():
        c=calc[a];solid=c.fill.patternType=='solid';painted.append(solid)
        if solid:counts[palette[c.fill.fgColor.rgb]]+=1
    assert painted==[True]*1607+[False]*(2600-1607)
    for r in range(3,11):assert counts[r]==calc[f'C{r}'].value==calc[f'AA{r}'].value
    for a in paint_order()[:1501]:assert calc[a]._style==b['10월계산기'][a]._style
    for m in range(1,10):
        for name in [f'{m}월',f'{m}월계산기']:
            for row in b[name]:
                for c in row:assert c.value==w[name][c.coordinate].value,(name,c.coordinate)
    for r in range(20,29):
        for c in range(1,16):assert b['예비비'].cell(r,c).value==w['예비비'].cell(r,c).value
    reserve=verify_month(w,10);view=query_check(w,10);assert_no_errors(w)
    before,after=ExtendedPackage(baseline),ExtendedPackage(final)
    for c in before.sheet('RAW').iter(Q('c')):
        if c.get('r') in ['Q1','Q2']:continue
        assert ET.tostring(c)==ET.tostring(after.cell('RAW',c.get('r'))),('old RAW',c.get('r'))
    for row in range(64,249):
        for c in 'ABCDEFGHIJKLM':
            addr=f'{c}{row}'
            assert ET.tostring(before.cell('잔액',addr))==ET.tostring(after.cell('잔액',addr)),('old ledger input',addr)
    with ZipFile(baseline) as a,ZipFile(final) as z:
        assert a.namelist()==z.namelist() and z.testzip() is None
        for name in a.namelist():
            if name.startswith(('xl/drawings/','xl/media/','xl/ctrlProps/','xl/activeX/')) or name=='xl/vbaProject.bin' or name.endswith('.rels'):
                assert a.read(name)==z.read(name),name
        os,ns=ET.fromstring(a.read('xl/styles.xml')),ET.fromstring(z.read('xl/styles.xml'))
        for component in os:
            target=ns.find(component.tag)
            if component.tag==Q('cellXfs'):
                assert [ET.tostring(x) for x in component]==[ET.tostring(x) for x in list(target)[:len(component)]]
            else:assert ET.tostring(component)==ET.tostring(target),('style component',component.tag)
    structure=validate_package(final);assert not structure['errors'],structure
    b.close();w.close();gc.collect()
    return {'date':'2026-10-19','as_of':'2026-10-19','latest_transaction':'2026-10-19T20:48:00.227000',
            'baseline_sha256':digest(baseline),'source_sha256':digest(source),'workbook_sha256':digest(final),
            'old_october_records_preserved':164,'new_transactions':7,'october_records':171,'total_raw_records':2492,
            'new_card_transactions':6,'new_card_actual':105450,'new_card_tetris':106000,
            'housing_cash_premium':110075,'housing_classification':'이사부대','additional_actual':110075,'additional_budget':111000,
            'calendar_actual':ACTUAL,'calendar_tetris':TETRIS,'monthly_total_actual':2211447,'monthly_total_budget':2241000,
            'living_balance':4409691,'wrap_total':6559691,'scope_counts':{'WRAP':16,'타계좌':48,'카드사용':107},
            'wrap_inflows':cash_in,'wrap_outflows':cash_out,'october_ledger_rows':len(october),
            'food_and_category_block_counts':dict(counts),'colored_cells':1607,'daily_calendar_checks':31,
            'default_query':view,'reserve_alignment':reserve,'cash_history_count':5,'transfer_count':12,
            'kakaopay_outflow_preserved':True,'prior_raw_and_ledger_inputs_preserved':True,'prior_month_values_preserved':True,
            'vba_drawings_controls_relationships_byte_identical':True,'formula_errors':0,'structure_validation':structure,
            'native_excel_opened':False,'sharepoint_file_updated':False,**metadata}


def update_audit(previous,final,report,output):
    """Regenerate current audit summaries from verified results; retain history."""
    a=json.loads(Path(previous).read_text());w=openpyxl.load_workbook(final,data_only=True)
    a['raw']['count']=report['total_raw_records'];a['raw']['end']=report['latest_transaction']
    a['raw']['by_month']['2026-10']=report['october_records'];a['raw']['october_scope']=report['scope_counts']
    a['october_ledger'].update(rows=17,raw_linked_rows=16,internal_classification_rows=1)
    a['wrap']['balances']['생활비']=report['living_balance'];a['wrap']['total']=report['wrap_total']
    a['wrap']['inflows']=report['wrap_inflows'];a['wrap']['outflows']=report['wrap_outflows'];a['wrap']['cached_matches']=True
    for item in a['calculators']:
        if item['month']==10:item.update(color_total=TETRIS,display=TETRIS,salary_burden=TETRIS,
                                         counts=report['food_and_category_block_counts'])
    a['october_block_layout'].update(colored_cells=1607,full_blocks=160,partial_block_cells=7)
    a['october_lifestyle'].update(records=98,actual=ACTUAL,tetris=TETRIS,calendar_actual=ACTUAL,calendar_tetris=TETRIS)
    for addr in set(a['october_summary'])|{'H39','I39','I38','I44','I45'}:
        if re.fullmatch(r'[A-Z]+[0-9]+',addr):a['october_summary'][addr]=w['10월'][addr].value
    a['as_of']='2026-10-19';a['ui_updated_at']='2026-10-19'
    a['current_verification']='raw_update_20261019_validation.json';a['workbook_sha256']=report['workbook_sha256']
    a['kakaopay_wrap_outflow_20261019']['scope']='카카오페이 출금 재분류 당시 기록. 해당 분류는 유지하며 추가 보증보험료·최신 잔액·해시는 raw_update_20261019_validation.json을 따릅니다.'
    a['delta_20261019']={k:report[k] for k in ['date','latest_transaction','old_october_records_preserved',
        'new_transactions','october_records','total_raw_records','new_card_transactions','new_card_actual',
        'new_card_tetris','housing_cash_premium','housing_classification','additional_actual','additional_budget',
        'calendar_actual','calendar_tetris','monthly_total_actual','monthly_total_budget','living_balance','wrap_total',
        'scope_counts','colored_cells','daily_calendar_checks','cash_history_count','transfer_count',
        'kakaopay_outflow_preserved','prior_raw_and_ledger_inputs_preserved','prior_month_values_preserved']}
    a['delta_20261019'].update(details='raw_update_20261019_validation.json',guide='2026-10-19_RAW반영_검증.md',
        previous_version='history/2026-10-19/04',source='raw/2026-10-19/2026-10-01_2026-10-31.xlsx',
        unit_tests_passed=30,native_excel_opened=False,sharepoint_file_updated=False)
    w.close();gc.collect()
    Path(output).write_text(json.dumps(a,ensure_ascii=False,indent=2)+'\n')


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for name in ['baseline','source','previous','workdir','output','report']:parser.add_argument('--'+name,required=True)
    a=parser.parse_args();folder=Path(a.workdir);candidate=folder/'candidate.xlsm'
    metadata=prepare(a.baseline,a.source,a.previous,candidate)
    print('Seven new rows prepared; recalculating disposable copy.',flush=True)
    calculated=recalculate(candidate,folder/'default')
    caches=import_changed_caches(candidate,calculated,a.output)
    report=verify(a.baseline,a.output,a.source,metadata);report['changed_formula_caches']=caches;report['month_regressions']=[]
    for m in [9,11]:
        def edit(p,month=m):p.write('잔액','B3',datetime(2026,month,1))
        path=recalculate(candidate,folder/f'month_{m}',edit)
        w=openpyxl.load_workbook(path,data_only=True)
        report['month_regressions'].append(query_check(w,m));w.close();gc.collect()
    Path(a.report).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k not in ['changed_formula_caches','source_mapping']},ensure_ascii=False),flush=True)
