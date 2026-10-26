"""Incremental RAW-to-month/calendar engine preserving the recorded history.

The delivery XLSM is patched as an OPC package. Spreadsheet libraries author
only new RAW/helper sheets; the existing VBA, cash ledger and reserve survive.
Historical daily balances are anchors, not silently reconstructed estimates.
"""
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime, timedelta
from hashlib import sha256
from io import BytesIO
import argparse
import json
import math
from pathlib import Path
import re
from zipfile import ZipFile

from lxml import etree as ET
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side, Protection
from openpyxl.worksheet.table import Table, TableColumn, TableFormula, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.utils import get_column_letter as col

from redesign_reserve import ExtendedPackage, append_styles
from update_20261015 import Q, RNS, PNS
from update_cash_details_20261019 import paint_order, import_changed_caches
from validate_excel_structure import validate_package
from verify_reserve_redesign import recalculate, assert_no_errors

BACK = '_RAW자동연동'
TABLE = 'RawData2026'
PREPARED_END = 4001
CAPACITY = 8000
YEAR = 2026
VARIABLE = {'변동비','담배','일회성전자','경조·선물','여행','약속·만남','이사부대','이사정착품','공연·휴가지원'}
FLOWS = {'자산형성-희망두배':'희망두배','자산형성-청약':'청약','자산형성-증권':'증권',
         '자산형성-IRP':'IRP','자산형성-ISA':'ISA','자산형성-연금':'연금','대출상환':'원금','의무상환':'원금',
         '공동-모임통장':'모임','공동-가족통장':'가족','공동-여행통장':'여행'}
FLOW_ROWS = {37:'희망두배',38:'청약',39:'증권',40:'ISA',41:'연금',42:'IRP',43:'원금',57:'모임',58:'가족',59:'여행'}
SAFE_FIXED = ['월세','관리비','서울도시가스','예스코','도시가스','롯데손해보험','SKT','유모비','LG헬로비전','티브로드',
              '후불교통','청년 버팀목 전세대출이자','토스뱅크 이자','카카오뱅크 이자','우리은행 이자',
              '구글플레이','T우주 CU구독','T우주+Youtube Premium','ChatGPT 구독','로켓와우','카카오 정기결제',
              'Apple Appstore','네이버플러스 멤버쉽','티빙+디즈니플러스','우버택시 멤버쉽']
HEADERS = ['일시','월','자산','분류','소분류','내용','금액(KRW)','수입/지출','추가입력','재분류','본인부담','결제구분','비고','입력확인',
           '유효일시','거래키','동일거래순번','잔액 관리범위','WRAP 입출금 방향','적용분류','적용결제','실금액','적용부담','계산기분류',
           '달력금액','달력칸수','고정항목','고정금액','급여','용돈','상여','기타수입','대납원금','적립분류','적립금액','기본분류','기본결제','기본부담','카드할인액']
H = {col(i): title for i,title in enumerate(HEADERS,1)}
RR = lambda c: f'{TABLE}[{H[c]}]'
MONEY = '#,##0;[Red]-#,##0;–'


def digest(path): return sha256(Path(path).read_bytes()).hexdigest()
def ceil_units(amount): return math.ceil((amount or 0)/1000)
def fill_key(cell):
    x=cell.fill.fgColor
    return x.type, x.rgb if x.type=='rgb' else x.indexed if x.type=='indexed' else x.theme, x.tint


def fixed_slot(title, sheet):
    text=title.replace(' ','').lower()
    labels={r:str(sheet[f'L{r}'].value or '') for r in range(4,25)}
    def find(word, default=26): return next((r for r,v in labels.items() if word in v),default)
    if title=='월세':return 4
    if title=='관리비':return 5
    if any(t in text for t in ['예스코','도시가스']):return 6
    if title=='롯데손해보험':return 7
    if title=='SKT':return 8
    if title=='유모비':return 9
    if any(t in text for t in ['티브로드','헬로비전']):return 10
    if any(t in text for t in ['후불교통','교통카드후불','티머니','이즐충전','기후동행']):return 11
    if title=='청년 버팀목 전세대출이자':return 13
    if any(t in text for t in ['토스뱅크이자','카카오뱅크이자','우리은행이자']):
        return find(title.split()[0])
    if title=='구글플레이':return find('유튜브')
    if title=='T우주 CU구독':return find('CU구독',find('T우주'))
    if 't우주' in text:return find('T우주')
    if title=='우버택시 멤버쉽':return find('Uber')
    if title=='ChatGPT 구독':return find('ChatGPT')
    if title=='로켓와우':return 16
    if title=='카카오 정기결제':return 17
    if title=='Apple Appstore':return 19
    if '네이버플러스멤버' in text:return 20
    if title=='티빙+디즈니플러스':return 21
    if text.startswith('11번가'):return 22
    if '롯데카드띵샵' in text:return find('에어팟',23)
    if '애플케어' in text:return find('애플케어',24)
    if '네이버플러스스토어(' in text:return find('가습기')
    return 26


def palette_for(classification, category, sub):
    if classification in {'담배','여행'}:return 5
    if classification=='약속·만남':return 8
    if classification=='경조·선물':return 10
    if classification=='일회성전자':return 9
    if category=='식비':return 3
    if category=='패션/미용':return 6
    if category=='건강/보험/의료':return 7
    if category=='잡화/소품/가구/생필품':return 9
    if category=='교통/차량':return 5
    if category=='여행/약속/공사':return 8 if sub=='만남' else 5
    return 4


def sum_formula(column, month_ref, *criteria):
    parts=[RR(column),RR('B'),month_ref]
    for c,value in criteria:parts += [RR(c),str(value)]
    return '=SUMIFS('+','.join(parts)+')'


def count_formula(month_ref, slot):
    return f'COUNTIFS({RR("B")},{month_ref},{RR("AA")},{slot},{RR("T")},"고정비",{RR("H")},"지출",{RR("Q")},1)'


def default_class(r, safe_end):
    # D is the original export category; J remains an optional explicit override.
    expense=f'IF(OR(D{r}="정기구독",D{r}="고정비",COUNTIF(\'{BACK}\'!$AN$2:$AN${safe_end},F{r})>0,LEFT(SUBSTITUTE(F{r}," ",""),4)="11번가",ISNUMBER(SEARCH("띵샵",F{r})),ISNUMBER(SEARCH("애플케어",F{r}))),"고정비",IF(OR(ISNUMBER(SEARCH("대리결제",I{r})),ISNUMBER(SEARCH("대납",I{r}))),"대납",IF(F{r}="보증보험료","이사부대",IF(OR(C{r}="가족통장",C{r}="모임통장",C{r}="20260606"),"공금","변동비"))))'
    expense=f'IF(D{r}="투자","투자",IF(OR(D{r}="보증금",F{r}="계약금",F{r}="잔금"),"보증금",IF(AND(V{r}<0,ISNUMBER(SEARCH("카드",F{r})),ISNUMBER(SEARCH("할인",F{r}))),"카드할인",{expense})))'
    transfer=f'IF(ISNUMBER(SEARCH("희망두배",D{r})),"자산형성-희망두배",IF(ISNUMBER(SEARCH("주택청약",D{r})),"자산형성-청약",IF(OR(D{r}="모임통장",D{r}="20260606"),"공동-모임통장",IF(D{r}="가족통장","공동-가족통장",IF(D{r}="여행통장","공동-여행통장","내부이체")))))'
    transfer=f'IF(ISNUMBER(SEARCH("IRP",D{r})),"자산형성-IRP",IF(ISNUMBER(SEARCH("ISA",D{r})),"자산형성-ISA",IF(ISNUMBER(SEARCH("연금저축",D{r})),"자산형성-연금",IF(ISNUMBER(SEARCH("증권",D{r})),"자산형성-증권",{transfer}))))'
    income=f'IF(OR(D{r}="급여",D{r}="월급",ISNUMBER(SEARCH("급여",F{r})),F{r}="월급"),"월급",IF(OR(D{r}="용돈",F{r}="용돈"),"용돈",IF(OR(D{r}="상여",ISNUMBER(SEARCH("상여",F{r}))),"상여","이자·환급")))'
    return f'=IF(A{r}="","",IF(LEFT(H{r},2)="이체",{transfer},IF(F{r}="카드정산 결제","내부이체",IF(H{r}="수입",{income},IF(H{r}="지출",{expense},"확인필요")))))'


def row_formulas(r, safe_end, alias_end, exclusions, last):
    exclusion='OR('+','.join(f'ROW()={i}' for i in sorted(exclusions))+')' if exclusions else 'FALSE'
    eligible='+'.join(f'(T{r}="{tag}")' for tag in sorted(VARIABLE))
    old_special=f'AND(ROW()<={last},IFERROR(MONTH(O{r}),0)<=8,OR(T{r}="이사부대",T{r}="이사정착품",T{r}="공연·휴가지원"))'
    palette=f'IF(OR(T{r}="담배",T{r}="여행"),5,IF(T{r}="약속·만남",8,IF(T{r}="경조·선물",10,IF(T{r}="일회성전자",9,IF(D{r}="식비",3,IF(D{r}="패션/미용",6,IF(D{r}="건강/보험/의료",7,IF(D{r}="잡화/소품/가구/생필품",9,IF(D{r}="교통/차량",5,IF(D{r}="여행/약속/공사",IF(E{r}="만남",8,5),4))))))))))'
    flow='""'
    for tag,value in reversed(list(FLOWS.items())):flow=f'IF(T{r}="{tag}","{value}",{flow})'
    key_parts=[f'TEXT(O{r},"0.0000000000")']+[f'LEN({c}{r})&":"&{c}{r}' for c in ['C','D','E','F']]+[f'V{r}']+[f'LEN({c}{r})&":"&{c}{r}' for c in ['H','I']]
    key='&"|"&'.join(key_parts)
    criterion=f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(P{r},"~","~~"),"*","~*"),"?","~?")'
    return {
        'O':f'=IF(A{r}="","",IF(ISNUMBER(A{r}),A{r},IFERROR(VALUE(A{r}),"")))',
        'B':f'=IF(O{r}="","",DATE(YEAR(O{r}),MONTH(O{r}),1))',
        'P':f'=IF(O{r}="","",{key})',
        'Q':f'=IF(P{r}="","",IF(LEN({criterion})<=250,COUNTIF($P$2:P{r},{criterion}),SUMPRODUCT(($O$2:O{r}=O{r})*($C$2:C{r}=C{r})*($D$2:D{r}=D{r})*($E$2:E{r}=E{r})*($F$2:F{r}=F{r})*($V$2:V{r}=V{r})*($H$2:H{r}=H{r})*($I$2:I{r}=I{r}))))',
        'AJ':default_class(r,safe_end), 'J':f'=AJ{r}',
        'T':f'=IF(A{r}="","",IF(J{r}="",AJ{r},J{r}))',
        'AK':f'=IF(A{r}="","",IF(T{r}="대납","정산대기",IF(OR(C{r}="가족통장",C{r}="모임통장",C{r}="20260606",T{r}="공금"),"공동","본인")))',
        'L':f'=AK{r}', 'U':f'=IF(A{r}="","",IF(L{r}="",AK{r},L{r}))',
        'V':f'=IF(A{r}="","",IFERROR(VALUE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(G{r},",",""),"원",""),"₩","")),0))',
        'AL':f'=IF(A{r}="","",IF(OR(T{r}="대납",U{r}<>"본인"),0,V{r}))',
        'K':f'=AL{r}', 'W':f'=IF(A{r}="","",IF(K{r}="",AL{r},K{r}))',
        'X':f'=IF(A{r}="","",{palette})',
        'Y':f'=IF(A{r}="",0,IF(AND(ISNUMBER(O{r}),H{r}="지출",U{r}="본인",Q{r}=1,({eligible})>0,NOT({exclusion}),NOT({old_special})),W{r},0))',
        'Z':f'=IF(Y{r}=0,0,CEILING(Y{r}/1000,1))',
        'AA':f'=IF(T{r}<>"고정비",0,IFERROR(VLOOKUP(MONTH(O{r})&"|"&F{r},\'{BACK}\'!$AK$2:$AL${alias_end},2,FALSE),IF(LEFT(SUBSTITUTE(F{r}," ",""),4)="11번가",22,IF(ISNUMBER(SEARCH("띵샵",F{r})),23,IF(ISNUMBER(SEARCH("애플케어",F{r})),24,26)))))',
        'AB':f'=IF(AND(T{r}="고정비",H{r}="지출",U{r}="본인",Q{r}=1),W{r},0)',
        'AC':f'=IF(AND(T{r}="월급",H{r}="수입",U{r}="본인",Q{r}=1),W{r},0)',
        'AD':f'=IF(AND(T{r}="용돈",H{r}="수입",U{r}="본인",Q{r}=1),W{r},0)',
        'AE':f'=IF(AND(T{r}="상여",H{r}="수입",U{r}="본인",Q{r}=1),W{r},0)',
        'AF':f'=IF(AND(OR(T{r}="이자·환급",T{r}="부수입"),F{r}<>"토스뱅크카드 추가환급",H{r}="수입",U{r}="본인",Q{r}=1),W{r},0)',
        'AG':f'=IF(AND(T{r}="대납",H{r}="지출",Q{r}=1),V{r},0)',
        'AH':f'=IF(A{r}="","",{flow})',
        'AI':f'=IF(AND(AH{r}<>"",OR(H{r}="지출",H{r}="이체지출",H{r}="이체출금"),U{r}="본인",Q{r}=1),V{r},0)',
        'AM':f'=IF(AND(Q{r}=1,U{r}="본인"),IF(AND(T{r}="카드할인",H{r}="지출"),-V{r},IF(AND(H{r}="수입",OR(T{r}="카드할인",F{r}="토스뱅크카드 추가환급")),V{r},0)),0)',
        'R':f'=IF(A{r}="","",IF(OR(C{r}="미래에셋_WRAP형CMA",D{r}="미래에셋_WRAP형CMA"),"WRAP",IF(OR(ISNUMBER(SEARCH("카드",C{r})),C{r}="KB쿠팡와우"),"카드사용","타계좌")))',
        'S':f'=IF(R{r}<>"WRAP","",IF(OR(H{r}="수입",H{r}="이체수입",H{r}="이체입금"),"입금","출금"))',
        'N':f'=IF(A{r}="","",IF(O{r}="","날짜 확인",IF(Q{r}>1,IF(ROW()<={last},"기존 중복 확인","중복 제외"),IF(T{r}="확인필요","수입/지출 확인",IF(OR(G{r}="",AND(G{r}<>0,V{r}=0)),"금액 확인",IF(YEAR(O{r})<>2026,"2026년 외 거래",IF(AND(ROW()>{last},R{r}="WRAP"),"월간 반영 · 잔액원장 별도",IF(ROW()<={last},"기존자료","자동 반영"))))))))',
    }


def base_model(w, exclusions):
    actual=defaultdict(float); units=defaultdict(int); colors=defaultdict(int)
    fixed=defaultdict(float); counts=Counter(); income=defaultdict(float); flows=defaultdict(float);discounts=defaultdict(float); seen=set()
    for r,row in enumerate(list(w['RAW'].values)[1:],2):
        dt,_,asset,category,sub,title,amount,kind,note,tag,personal,payer= row[:12]
        if not isinstance(dt,datetime):continue
        key=(dt,asset,category,sub,title,amount,kind,note)
        if key in seen:continue
        seen.add(key); m=dt.month
        if kind=='지출' and payer=='본인' and tag in VARIABLE and r not in exclusions and not(m<=8 and tag in {'이사부대','이사정착품','공연·휴가지원'}):
            actual[dt.date()]+=personal or 0; units[dt.date()]+=ceil_units(personal)
            colors[m,palette_for(tag,category,sub)]+=ceil_units(personal)
        if tag=='고정비' and kind=='지출' and payer=='본인':
            slot=fixed_slot(title,w[f'{m}월']); fixed[m,slot]+=personal or 0;counts[m,slot]+=1
        if kind=='수입' and payer=='본인' and title!='토스뱅크카드 추가환급' and tag in ['월급','용돈','상여','이자·환급','부수입']:
            income[m,tag if tag in ['월급','용돈','상여'] else '기타']+=personal or 0
        if tag in FLOWS and kind in ['지출','이체지출','이체출금'] and payer=='본인':flows[m,FLOWS[tag]]+=amount or 0
        if payer=='본인':
            if tag=='카드할인' and kind=='지출':discounts[m]-=amount or 0
            elif kind=='수입' and (tag=='카드할인' or title=='토스뱅크카드 추가환급'):discounts[m]+=amount or 0
    return actual,units,colors,fixed,counts,income,flows,discounts


def grid(s,m):
    for r in [3,8,13,18,23,28]:
        for c in range(2,9):
            dt=s.cell(r,c).value
            if isinstance(dt,datetime) and dt.month==m:yield dt,r,c


def author(source, output):
    original=openpyxl.load_workbook(source); data=openpyxl.load_workbook(source,data_only=True)
    if BACK in original.sheetnames:
        raise ValueError('이미 자동 연동된 파일입니다. 새 기준을 만들지 말고 기존 RAW 표의 입력 행만 갱신하세요.')
    last=data['RAW'].max_row
    exclusions=set()
    for m in range(1,9):
        for r in range(52,62):
            match=re.fullmatch(r'=RAW![GK](\d+)',str(original[f'{m}월'][f'D{r}'].value))
            if match:exclusions.add(int(match[1]))
    actual,units,colors,fixed,counts,income,flows,discounts=base_model(data,exclusions)
    wb=openpyxl.Workbook();raw=wb.active;raw.title='RAW';back=wb.create_sheet(BACK)
    raw.append(HEADERS)
    titles=sorted({row[5] for row in list(data['RAW'].values)[1:] if row[9]=='고정비'}|set(SAFE_FIXED))
    alias_row=2
    for m in range(1,13):
        for title in titles:
            back.cell(alias_row,37,f'{m}|{title}');back.cell(alias_row,38,fixed_slot(title,data[f'{m}월']));alias_row+=1
    for r,title in enumerate(SAFE_FIXED,2):back.cell(r,40,title)
    forms={}
    for r in range(2,PREPARED_END+1):
        fs=row_formulas(r,len(SAFE_FIXED)+1,alias_row-1,exclusions,last)
        if r==2:forms=fs
        if r<=last:
            for c in range(1,14):
                value=original['RAW'].cell(r,c).value
                if value is not None:raw.cell(r,c,value)
        for c,f in fs.items():
            existing=raw[f'{c}{r}'].value
            if c in ['J','L'] and r<=last and existing is not None:continue
            if c=='K' and r<=last and existing is not None:
                old=data['RAW'];same=old[f'K{r}'].value==old[f'G{r}'].value and old[f'L{r}'].value=='본인'
                if not same:continue
            raw[f'{c}{r}']=f
        if r<=last:
            for c in ['R','S']:
                v=original['RAW'][f'{c}{r}'].value
                if v is not None:raw[f'{c}{r}']=v
    tab=Table(displayName=TABLE,ref=f'A1:AM{PREPARED_END}')
    tab.autoFilter=AutoFilter(ref=tab.ref)
    tab.tableColumns=[TableColumn(id=i,name=t) for i,t in enumerate(HEADERS,1)]
    for i,c in enumerate(tab.tableColumns,1):
        if col(i) in forms:c.calculatedColumnFormula=TableFormula(attr_text=forms[col(i)].lstrip('='))
    tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True,showColumnStripes=False)
    raw.add_table(tab);raw.freeze_panes='C2';raw.sheet_view.zoomScale=80
    # Legacy cash/reserve links use stable physical RAW rows. Filtering is safe;
    # sorting/deleting rows is not. No password: the owner can expand deliberately.
    raw.protection.sheet=True;raw.protection.autoFilter=False;raw.protection.sort=True
    raw.protection.insertRows=True;raw.protection.deleteRows=True
    raw.protection.formatRows=False;raw.protection.formatColumns=False
    widths={'A':21,'B':10,'C':22,'D':21,'E':15,'F':27,'G':15,'H':12,'I':28,'J':19,'K':15,'L':13,'M':34,'N':29}
    for c,width in widths.items():raw.column_dimensions[c].width=width
    raw.column_dimensions.group('O','AP',hidden=True)
    for row in raw.iter_rows(min_row=1,max_row=PREPARED_END,max_col=len(HEADERS)):
        for cell in row:
            cell.font=Font(name='맑은 고딕',size=10,color='243B53')
            cell.alignment=Alignment(vertical='center',wrap_text=cell.column in [6,9,13,14])
            cell.protection=Protection(locked=cell.row==1 or cell.column not in [1,3,4,5,6,7,8,9,10,11,12,13])
            if cell.row==1:
                cell.fill=PatternFill('solid',fgColor='17324D');cell.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
            elif cell.column<=14:
                cell.fill=PatternFill('solid',fgColor='EDF7F2' if cell.column in [2,10,11,12,14] else ('FFFFFF' if cell.row%2 else 'F5F8FC'))
            if cell.column in [1,15]:cell.number_format='yyyy-mm-dd hh:mm'
            if cell.column==2:cell.number_format='yyyy-mm'
            if cell.column in [7,11,22,23,25,28,29,30,31,32,33,35,38]:cell.number_format=MONEY
    raw.row_dimensions[1].height=30
    raw.sheet_properties.pageSetUpPr.fitToPage=True
    raw.print_options.horizontalCentered=False
    raw.print_area=f'A1:N{last}';raw.print_title_rows='1:1'
    categories=list(dict.fromkeys([row[9] for row in list(data['RAW'].values)[1:] if row[9]]+list(FLOWS)))
    for r,value in enumerate(categories,50):raw.cell(r,43,value)
    raw.column_dimensions['AQ'].width=26;raw.column_dimensions['AR'].width=20
    for c in ['AS','AT','AU','AV']:raw.column_dimensions[c].width=12
    guide=[('RAW 입력 · 월별 필터',''),('입력 칸','A 일시 / C~I 원본 거래'),('월 필터','B열 ▼에서 월 선택'),('초록 칸','월·분류·부담 자동 계산'),('다음 입력행',f'=IFERROR(MATCH(TRUE,INDEX(A2:A{PREPARED_END}="",0),0)+1,{PREPARED_END+1})'),
           ('거래 건수',f'=COUNT({RR("O")})'),('중복 확인·추가분 제외',f'=COUNTIF({RR("Q")},">1")'),('분류 수정','J열에서 직접 선택 가능'),('이체·대납','소비와 구분; 필요 시 J 수정'),('WRAP 잔액','기존 잔액 원장 별도 관리'),('계산기','RAW에서 숫자·색상 자동'),('새 입력 전','필터를 전체로 해제하세요'),('붙여넣기','A 날짜 / C~I 거래를 입력')]
    for r,(label,value) in enumerate(guide,1):
        raw.cell(r,43,label);raw.cell(r,44,value)
        raw.cell(r,43).font=Font(name='맑은 고딕',size=10,bold=True,color='17324D')
        raw.merge_cells(start_row=r,start_column=44,end_row=r,end_column=48)
        raw.cell(r,44).alignment=Alignment(wrap_text=True,vertical='center')
        if r<=13:raw.row_dimensions[r].height=28
    raw.cell(15,43,'다음 입력 위치');raw.cell(15,44,'=HYPERLINK("#\'RAW\'!A"&AR5,"입력할 빈 행으로 이동 →")')
    raw.cell(14,43,'입력 공간');raw.cell(14,44,'4,000건 준비 · 확장 시 시트 보호 해제')
    raw.cell(16,43,'행 연결 보호');raw.cell(16,44,'필터 허용 · 정렬/행삭제는 보호')
    dv=DataValidation(type='list',formula1=f'$AQ$50:$AQ${49+len(categories)}',allow_blank=True);dv.errorTitle='재분류 목록';dv.showErrorMessage=True
    dv.add(f'J2:J{PREPARED_END}');raw.add_data_validation(dv)
    kinds=DataValidation(type='list',formula1='"지출,수입,이체지출,이체수입,이체출금,이체입금"',allow_blank=True);kinds.add(f'H2:H{PREPARED_END}');raw.add_data_validation(kinds)
    for i in range(365):
        dt=datetime(YEAR,1,1)+timedelta(days=i);s=data[f'{dt.month}월']
        match=next((r,c) for day,r,c in grid(s,dt.month) if day.date()==dt.date());r,c=match
        vals=[dt,s.cell(r+2,c).value or 0,(s.cell(r+1,c).value or 0)/1000,actual[dt.date()],units[dt.date()]]
        for j,v in enumerate(vals,1):back.cell(i+2,j,v)
        back.cell(i+2,6,s.cell(r+2,c).value is None)
    meta={'last_raw':last,'exclusions':sorted(exclusions),'months':{},'source_sha256':digest(source),'alias_end':alias_row-1}
    for m in range(1,13):
        s=data[f'{m}월'];calc=data[f'{m}월계산기'];legend={fill_key(calc[f'B{r}']):r for r in range(3,11)}
        count=Counter();last_painted=0;block=500+(m-1)*2600
        for i,a in enumerate(paint_order()):
            palette=legend.get(fill_key(calc[a]),0) if calc[a].fill.patternType=='solid' else 0
            if palette:count[palette]+=1;last_painted=i+1
            back.cell(block+i,16,palette);back.cell(block+i,17,count[palette] if palette else 0)
        for palette in range(3,11):
            mr=2+(m-1)*8+palette-3
            for c,v in enumerate([datetime(YEAR,m,1),palette,colors[m,palette],count[palette]],7):back.cell(mr,c,v)
        for slot in range(4,27):
            mr=2+(m-1)*23+slot-4
            for c,v in enumerate([m,slot,fixed[m,slot],counts[m,slot],s[f'O{slot}'].value if slot<=24 else 0,s[f'U{slot}'].value if slot<=24 else None],28):back.cell(mr,c,v)
        for idx,(key,amt) in enumerate([('월급',income[m,'월급']),('용돈',income[m,'용돈']),('상여',income[m,'상여']),('기타',income[m,'기타'])]):back.cell(m+1,20+idx,amt)
        back.cell(m+1,24,discounts[m])
        for idx,flow in enumerate(FLOW_ROWS.values()):
            mr=2+(m-1)*10+idx;back.cell(mr,42,m);back.cell(mr,43,flow);back.cell(mr,44,flows[m,flow])
        meta['months'][str(m)]={'base_colors':dict(count),'last_painted':last_painted,'base_block':block}
    back.sheet_state='hidden'
    wb.save(output);original.close();data.close();wb.close()
    return meta


def inject_authored(p, authored):
    """Import two authored sheets and the RAW table, preserving existing style IDs."""
    with ZipFile(authored) as z:
        stylemap=append_styles(p,z)
        raw=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        back=ET.fromstring(z.read('xl/worksheets/sheet2.xml'))
        for tree in [raw,back]:
            for cell in tree.iter(Q('c')):
                if cell.get('s') is not None:cell.set('s',str(stylemap[int(cell.get('s'))]))
        old_pr=p.sheet('RAW').find(Q('sheetPr'));new_pr=raw.find(Q('sheetPr'))
        if old_pr is not None and old_pr.get('codeName'):
            new_pr.set('codeName',old_pr.get('codeName'))
        p.parts[p.sheets['RAW']]=ET.tostring(raw,encoding='UTF-8',xml_declaration=True)
        p.trees.pop(p.sheets['RAW'],None);p.cell_maps.pop('RAW',None);p.row_maps.pop('RAW',None)
        book=p.tree('xl/workbook.xml');sheets=book.find(Q('sheets'))
        sid=max(int(s.get('sheetId')) for s in sheets)+1
        part=f'xl/worksheets/sheet{sid}.xml'
        while part in p.parts:sid+=1;part=f'xl/worksheets/sheet{sid}.xml'
        p.parts[part]=ET.tostring(back,encoding='UTF-8',xml_declaration=True);p.sheets[BACK]=part
        rels=p.tree('xl/_rels/workbook.xml.rels')
        ET.SubElement(rels,'{'+PNS+'}Relationship',Id='rIdRawAutoModel',Type=RNS+'/worksheet',Target=f'worksheets/sheet{sid}.xml')
        ET.SubElement(sheets,Q('sheet'),name=BACK,sheetId=str(sid),state='hidden',attrib={'{'+RNS+'}id':'rIdRawAutoModel'})
        ctypes=p.tree('[Content_Types].xml');ns=ET.QName(ctypes).namespace
        ET.SubElement(ctypes,'{'+ns+'}Override',PartName='/'+part,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml')
        table_id=max([0]+[int(re.search(r'table(\d+)\.xml$',n)[1]) for n in p.parts if re.search(r'xl/tables/table\d+\.xml$',n)])+1
        table=ET.fromstring(z.read('xl/tables/table1.xml'));table.set('id',str(table_id))
        tablepart=f'xl/tables/table{table_id}.xml';p.parts[tablepart]=ET.tostring(table,encoding='UTF-8',xml_declaration=True)
        ET.SubElement(ctypes,'{'+ns+'}Override',PartName='/'+tablepart,ContentType='application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml')
        rawpart=Path(p.sheets['RAW']);rawrels=str(rawpart.parent/'_rels'/(rawpart.name+'.rels'))
        rtree=ET.fromstring(p.parts[rawrels]) if rawrels in p.parts else ET.Element('{'+PNS+'}Relationships',nsmap={None:PNS})
        ET.SubElement(rtree,'{'+PNS+'}Relationship',Id='rIdRawAutoTable',Type=RNS+'/table',Target=f'../tables/table{table_id}.xml')
        p.parts[rawrels]=ET.tostring(rtree,encoding='UTF-8',xml_declaration=True)
        p.sheet('RAW').find(Q('tableParts'))[0].set('{'+RNS+'}id','rIdRawAutoTable')
        app=p.tree('docProps/app.xml');ans=ET.QName(app).namespace;vns='http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes'
        pair=app.find(f'{{{ans}}}HeadingPairs/{{{vns}}}vector/{{{vns}}}variant[2]/{{{vns}}}i4')
        if pair is not None:pair.text=str(int(pair.text)+1)
        titles=app.find(f'{{{ans}}}TitlesOfParts/{{{vns}}}vector')
        if titles is not None:ET.SubElement(titles,'{'+vns+'}lpstr').text=BACK;titles.set('size',str(len(titles)))
        calc=book.find(Q('calcPr'))
        if calc is None:calc=ET.SubElement(book,Q('calcPr'))
        calc.set('calcMode','auto');calc.set('fullCalcOnLoad','1');calc.set('forceFullCalc','1')


def insert_before_conditional(s,node):
    from validate_excel_structure import WORKSHEET_RANK
    rank=WORKSHEET_RANK[ET.QName(node).localname]
    position=next((i for i,x in enumerate(s) if WORKSHEET_RANK.get(ET.QName(x).localname,-1)>rank),len(s))
    s.insert(position,node)


def add_dxf(p, fill):
    styles=p.tree('xl/styles.xml');dxfs=styles.find(Q('dxfs'))
    if dxfs is None:
        dxfs=ET.Element(Q('dxfs'));styles.insert(list(styles).index(styles.find(Q('cellStyles')))+1,dxfs)
    idx=len(dxfs);node=ET.SubElement(dxfs,Q('dxf'));node.append(deepcopy(fill));dxfs.set('count',str(len(dxfs)))
    return idx


def conditional(p,name,sqref,formula,dxf,priority):
    node=ET.Element(Q('conditionalFormatting'),sqref=sqref)
    rule=ET.SubElement(node,Q('cfRule'),type='expression',dxfId=str(dxf),priority=str(priority),stopIfTrue='1')
    ET.SubElement(rule,Q('formula')).text=formula.lstrip('=')
    insert_before_conditional(p.sheet(name),node)


def slot_formula(start=15,offset=0):
    return f'({offset}+(INT((ROW()-{start})/5)*10+INT((COLUMN()-2)/2))*10+MOD(COLUMN()-2,2)*5+MOD(ROW()-{start},5)+1)'


def install_calculator(p,month,source,meta):
    name=f'{month}월계산기';s=p.sheet(name);base=meta['months'][str(month)]
    book=p.tree('xl/workbook.xml');names=book.find(Q('definedNames'))
    if names is None:names=ET.SubElement(book,Q('definedNames'))
    for suffix,c in [('Palette','P'),('Rank','Q')]:
        ET.SubElement(names,Q('definedName'),name=f'_RawBase{suffix}{month:02d}').text=f"'{BACK}'!${c}${base['base_block']}:${c}${base['base_block']+2599}"
    for r in range(3,11):
        mr=2+(month-1)*8+r-3
        extra=f'+SUMIF(\'{month}월\'!$AC$52:$AC$61,{r},\'{month}월\'!$AB$52:$AB$61)' if month>=9 else ''
        current=sum_formula('Z',f"'{month}월'!$T$1",('X',r)).lstrip('=')
        p.write(name,f'C{r}',f"='{BACK}'!J{mr}+{current}-'{BACK}'!I{mr}{extra}",True)
        p.write(name,f'AB{r}',f'=MAX(0,C{r}-AA{r})',True)
        p.write(name,f'AC{r}','=0' if r==3 else f'=SUM($AB$3:AB{r-1})',True)
        p.write(name,f'AD{r}',f'=MAX(0,MIN(C{r},AA{r}))',True)
    p.write(name,'AF1',base['last_painted']);p.write(name,'AF2',CAPACITY)
    p.write(name,'AF3','=AF1+SUM(AB3:AB10)',True)
    white=ET.Element(Q('fill'));pat=ET.SubElement(white,Q('patternFill'),patternType='solid');ET.SubElement(pat,Q('fgColor'),rgb='FFFFFFFF')
    white_id=add_dxf(p,white);idx=slot_formula()
    hidden=f'IFERROR(AND(INDEX(_RawBasePalette{month:02d},{idx})>0,INDEX(_RawBaseRank{month:02d},{idx})>INDEX($AD$3:$AD$10,INDEX(_RawBasePalette{month:02d},{idx})-2)),FALSE)'
    conditional(p,name,'B15:U144',hidden,white_id,1)
    idx2=f'IF(ROW()<145,{idx},{slot_formula(165,2600)})'
    style=p.tree('xl/styles.xml');xfs=style.find(Q('cellXfs'));fills=style.find(Q('fills'))
    for r in range(3,11):
        legend=p.cell(name,f'B{r}');xf=xfs[int(legend.get('s'))];fill=fills[int(xf.get('fillId'))]
        dxf=add_dxf(p,fill)
        conditional(p,name,'B15:U144 B165:U434',f'AND({idx2}>$AF$1+$AC${r},{idx2}<=$AF$1+$AC${r}+$AB${r})',dxf,r-1)
    columns=s.find(Q('cols'))
    for c in list(columns):
        if int(c.get('min'))>=27:columns.remove(c)
    ET.SubElement(columns,Q('col'),min='27',max='32',hidden='1',width='12',customWidth='1')
    p.write(name,'A154','RAW 자동 계산 · 색상 블록도 조건부 서식으로 갱신됩니다. 수동 색상 집계 버튼은 사용하지 않습니다.')
    p.write(name,'A156','=IF(AF3>AF2,"색상 표시 범위 초과 · 금액 합계는 정상. 블록 영역 확장이 필요합니다.",IF(AF3>2600,"추가 블록은 아래 영역에 이어서 표시됩니다.","RAW 추가·수정 → 숫자와 색상 자동 반영"))',True)
    p.write(name,'A163','=IF(AF3>2600,"추가 블록 · 자동 표시","")',True)
    p.cell(name,'U434');s.find(Q('dimension')).set('ref','A1:AF434')
    # Disconnect the old macro which overwrites C3:C10 with physical-fill counts.
    drawings=s.find(Q('drawing'))
    if drawings is not None:
        relfile=str(Path(p.sheets[name]).parent/'_rels'/(Path(p.sheets[name]).name+'.rels'))
        rels=ET.fromstring(p.parts[relfile]);rid=drawings.get('{'+RNS+'}id')
        target=next(x.get('Target') for x in rels if x.get('Id')==rid)
        import posixpath
        part=target.lstrip('/') if target.startswith('/') else posixpath.normpath(str(Path(p.sheets[name]).parent/target))
        tree=p.tree(part)
        for shape in tree.iter():
            if shape.get('macro'):
                shape.attrib.pop('macro',None)
                texts=[x for x in shape.iter() if ET.QName(x).localname=='t']
                for i,text in enumerate(texts):text.text='RAW 자동' if i==0 else ''


def install_month(p,m,data,meta):
    name=f'{m}월';s=data[name]
    # Known expense details are views of the same RAW records, not new spending.
    if m>=9:
        for r in range(52,62):
            source=None
            value=str(s[f'D{r}'].value or '')
            if data[name][f'C{r}'].value=='보증보험료':source=2157 if m==9 else 2493 if m==10 else None
            if (m,r) in {(9,53),(9,54),(10,52),(10,53)}:
                source={(9,53):2080,(9,54):2286,(10,52):2334,(10,53):2474}[m,r]
            if source:
                p.write(name,f'G{r}',f'=INDEX({RR("Y")},{source-1})',True)
                p.write(name,f'AC{r}',f'=INDEX({RR("X")},{source-1})',True)
            else:p.write(name,f'AC{r}',4)
            p.write(name,f'AB{r}',f'=IF(D{r}="","",CEILING(F{r}/1000,1)-IF(G{r}="",0,CEILING(G{r}/1000,1)))',True)
        # AB is now in units (not won), making per-record rounding explicit.
        cols=p.sheet(name).find(Q('cols'))
        for node in list(cols):
            if int(node.get('min'))>=27:cols.remove(node)
        ET.SubElement(cols,Q('col'),min='27',max='29',hidden='1',width='12',customWidth='1')
    for dt,r,c in grid(s,m):
        base=(dt.date()-datetime(YEAR,1,1).date()).days+2;date=f'{col(c)}{r}'
        for target,rawcol,baseline,modelcol,scale in [(f'{col(c)}{r+2}','Y','B','D',1),(f'{col(c)}{r+1}','Z','C','E',1000)]:
            current=f'SUMIFS({RR(rawcol)},{RR("O")},">="&{date},{RR("O")},"<"&{date}+1)'
            delta=f'({current}-\'{BACK}\'!{modelcol}{base})'
            if m>=9:
                extra='H' if rawcol=='Y' else 'AB'
                delta+=f'+SUMIFS(${extra}$52:${extra}$61,$A$52:$A$61,">="&{date},$A$52:$A$61,"<"&{date}+1)'
            expression=f"('{BACK}'!{baseline}{base}+({delta}))*{scale}"
            p.write(name,target,f'=IF(AND(\'{BACK}\'!F{base},({delta})=0),"",{expression})',True)
        p.write(name,f'{col(c)}{r+3}',f'=IF(COUNT({col(c)}{r+1},{col(c)}{r+2})=0,"",SUM({col(c)}{r+1})-SUM({col(c)}{r+2}))',True)
    # Income follows all recorded dates in its month, not a stale as-of literal.
    for addr,rawcol,basecol in [('C37','AC','T'),('C38','AD','U'),('C39','AE','V'),('C40','AF','W')]:
        legacy=s[addr].value or 0
        if m>10:legacy=0
        p.write(name,addr,f'={legacy}+{sum_formula(rawcol,"$T$1").lstrip("=")}-\'{BACK}\'!{basecol}{m+1}',True)
    p.write(name,'H43',f'={s["H43"].value or 0}+{sum_formula("AM","$T$1").lstrip("=")}-\'{BACK}\'!X{m+1}',True)
    for r in range(4,25):
        mr=2+(m-1)*23+r-4;base=s[f'O{r}'].value
        bstatus=s[f'U{r}'].value
        cur=sum_formula('AB','$T$1',('AA',r)).lstrip('=');cnt=count_formula('$T$1',r)
        forecast=bstatus in ['예상','R'] or m>10
        if forecast:
            formula=f'=IF({cnt}>0,{cur},'+('""' if base is None else str(base))+')'
        else:formula=f'={base or 0}+{cur}-\'{BACK}\'!AD{mr}'
        p.write(name,f'O{r}',formula,True)
        p.write(name,f'Q{r}',f'=IF(COUNT(O{r}:P{r})=0,"",SUM(O{r})-SUM(P{r}))',True)
        fallback='예상' if forecast and base is not None else '기존 입력' if base is not None else ''
        p.write(name,f'U{r}',f'=IF({cnt}>0,"확정","{fallback}")',True)
    mr=2+(m-1)*23+26-4
    p.write(name,'K26','기타 고정비 · RAW 추가/변경')
    p.write(name,'O26',f'={sum_formula("AB","$T$1",("AA",26)).lstrip("=")}-\'{BACK}\'!AD{mr}',True)
    p.write(name,'Q26','=O26',True);p.write(name,'R26','=CEILING(Q26,1000)',True)
    p.write(name,'U26','=IF(O26=0,"","확정")',True)
    p.write(name,'Q25','=SUM(Q4:Q24,Q26)',True);p.write(name,'R25','=SUM(R4:R24,R26)',True)
    actual='+'.join(f'SUMIF(U4:U24,"{status}",Q4:Q24)' for status in ['확정','결제','RAW 확인','기존 입력'])+'+Q26'
    p.write(name,'H37','='+actual,True)
    rounded_actual='+'.join(f'IF(OR(U{r}="확정",U{r}="결제",U{r}="RAW 확인",U{r}="기존 입력"),R{r},0)' for r in range(4,25))+'+R26'
    p.write(name,'I37','='+rounded_actual,True)
    for i,(r,flow) in enumerate(FLOW_ROWS.items()):
        mr=2+(m-1)*10+i;base=s[f'N{r}'].value or 0
        change=sum_formula('AI','$T$1',('AH',f'"{flow}"')).lstrip('=')+f'-\'{BACK}\'!AR{mr}'
        expr=f'{base}+{change}'
        p.write(name,f'N{r}',f'=IF(({expr})=0,"",{expr})' if s[f'N{r}'].value is None else '='+expr,True)
    p.write(name,'I40','=SUM('+','.join(f'ROUNDUP(SUM(N{r}),-3)' for r in range(37,44))+')',True)
    p.write(name,'I42','=SUM('+','.join(f'ROUNDUP(SUM(N{r}),-3)' for r in range(57,60))+')',True)
    p.write(name,'AA1',f'=OR(COUNTIFS({RR("B")},$T$1)>0,COUNT(B5:H5,B10:H10,B15:H15,B20:H20,B25:H25,B30:H30)>0,SUM(D52:D61)<>0)',True)
    p.write(name,'U1',f'=IF(COUNTIFS({RR("B")},$T$1)=0,"",INT(MAX(IFERROR({RR("O")}*({RR("B")}=$T$1),0))))',True)
    p.cell(name,'U1').find(Q('f')).set('t','array');p.cell(name,'U1').find(Q('f')).set('ref','U1')
    p.write(name,'H45','=IF($AA$1,C41-H44,"")',True);p.write(name,'I45','=IF($AA$1,C41-I44,"")',True)
    p.write(name,'K32','RAW 자동 연동 · 새 거래는 RAW 표 안 빈 행에 입력합니다. 이체·대납은 소비와 분리되며 계산기 색상도 자동 갱신됩니다.')
    if m>=9:
        p.write(name,'E74','='+actual,True);p.write(name,'E75','=Q25',True)
        p.write(name,'E78','=E74+E76+'+sum_formula('AG','$T$1').lstrip('='),True)
        p.write(name,'E79','=E78',True)
    p.write(name,'A33','달력 지출 합계 · RAW 자동')


def prepare(source,workdir,output):
    workdir=Path(workdir);authored=workdir/'authored.xlsx'
    meta=author(source,authored)
    p=ExtendedPackage(source);inject_authored(p,authored)
    for priority,color,formula in [
        (1,'FFFDE7E9','OR($N2="날짜 확인",$N2="금액 확인",$N2="수입/지출 확인")'),
        (2,'FFFFEBC6','OR($N2="중복 제외",$N2="기존 중복 확인")'),
        (3,'FFFFF6D8','$N2="월간 반영 · 잔액원장 별도"'),
    ]:
        fill=ET.Element(Q('fill'));pattern=ET.SubElement(fill,Q('patternFill'),patternType='solid');ET.SubElement(pattern,Q('fgColor'),rgb=color)
        conditional(p,'RAW',f'A2:N{PREPARED_END}',formula,add_dxf(p,fill),priority)
    data=openpyxl.load_workbook(source,data_only=True)
    for m in range(1,13):install_month(p,m,data,meta);install_calculator(p,m,source,meta)
    for addr,text in {
        'A18':'RAW 입력 · 월별 필터',
        'A19':'RAW 표의 빈 행에 A 일시·C~I 거래를 추가하세요. B 월과 분류·본인부담은 자동 계산됩니다. B열 ▼에서 월별로 필터링합니다.',
        'A20':'1~12월 달력·월 합계·계산기 숫자와 색상이 RAW에서 자동 갱신됩니다. 이전 Calculate 버튼은 자동 표시로 바꾸어 매크로 허용 없이 사용합니다.',
        'A22':'기존 월별 이력과 목적자금 정산은 보존합니다. 중복 거래는 제외하고 RAW 입력확인에 표시합니다. 특수 용도는 J 재분류로 수정하세요.',
        'A23':'WRAP 잔액은 기존 잔액 원장 기준입니다. 새 RAW 입출금의 목적 계정 배분은 자동 추정하지 않으므로 잔액 원장 등록은 별도입니다.',
    }.items():p.write('사용안내',addr,text)
    p.save(output);data.close()
    return meta


if __name__=='__main__':
    parser=argparse.ArgumentParser(description='One-time RAW automation migration. Future inputs use the Excel table directly.')
    for arg in ['source','workdir','output','report']:parser.add_argument('--'+arg,required=True)
    args=parser.parse_args();folder=Path(args.workdir);folder.mkdir(parents=True,exist_ok=True)
    metadata=prepare(args.source,folder,folder/'candidate.xlsm')
    calculated=recalculate(folder/'candidate.xlsm',folder/'calculation')
    import_changed_caches(folder/'candidate.xlsm',calculated,args.output)
    from verify_raw_auto_link_20261026 import retain_unchanged_reserve,package_checks,check_default,check_source_rows
    preserved=retain_unchanged_reserve(args.source,args.output)
    report=package_checks(args.source,args.output)
    report['baseline_checks']=check_default(args.source,args.output)
    report['source_rows']=check_source_rows(args.source,args.output)
    report['blank_reference_caches_retained']=preserved
    report['metadata']=metadata
    Path(args.report).write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:v for k,v in report.items() if k!='metadata'},ensure_ascii=False,indent=2))
