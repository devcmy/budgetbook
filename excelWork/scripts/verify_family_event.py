"""Verify the requested label/239,000-won change and evidenced cash settlement."""
import argparse
from collections import Counter
from datetime import datetime
import gc
from hashlib import sha256
import json
from pathlib import Path
from zipfile import ZipFile

import openpyxl

from redesign_reserve import ExtendedPackage
from verify_reserve_redesign import (
    read, recalculate, assert_no_errors, verify_month, cache_original,
)


def verify_artifact(baseline, final, incoming):
    before, after = read(baseline), read(final)
    assert before.sheetnames == after.sheetnames
    assert [(s.title, s.sheet_state) for s in before] == [(s.title, s.sheet_state) for s in after]
    assert after['예비비']['L7'].value == '선택월 예비비 사용액'
    for a in ['L10', 'L13', 'E107', 'T107']:
        assert after['예비비'][a].value == 239000, a
    for a in ['F13', 'U107']:
        assert after['예비비'][a].value == 0, a
    assert after['예비비']['C13'].value == after['예비비']['I13'].value == 2150000
    assert after['예비비']['N437'].value == after['예비비']['N438'].value == 'OK'
    assert after['예비비']['G36'].value == 239000
    assert after['예비비']['K36'].value == '2건'
    assert after['잔액']['G16'].value == 6685466
    assert after['잔액']['G15'].value == 61000
    assert after['잔액']['G7'].value == before['잔액']['G7'].value == 4474466
    assert after['10월']['N66'].value == after['10월']['E79'].value == 239000
    assert after['10월']['E78'].value == 800000
    assert after['RAW']['N2466'].value == '연결'
    assert after['잔액']['W246'].value == 2466
    assert sum(after['잔액'][f'W{r}'].value == 2466 for r in range(64,501)) == 1
    # The new upload must be covered exactly, including the edited old refund.
    source = openpyxl.load_workbook(incoming, data_only=True).active
    wanted = Counter((v[0], v[1], v[5], v[6]) for v in source.values if isinstance(v[0],datetime))
    actual = Counter()
    for row in after['RAW'].iter_rows(min_row=2323,max_row=2466):
        v = [c.value for c in row]
        account = '20260606' if v[2] == '모임통장' else v[2]
        actual[(v[0],account,v[6],v[7])] += 1
    assert wanted == actual, (wanted-actual, actual-wanted)
    changed_raw = {f'{c}2422' for c in ['A','D','E','M']} | {'M2401','Q2','Q5','Q6'}
    for row in before['RAW']:
        for c in row:
            if c.coordinate not in changed_raw:
                assert c.value == after['RAW'][c.coordinate].value, ('RAW',c.coordinate)
    # No second salary/living charge, and no paint/calendar edits in any month.
    for m in range(1,13):
        for a in ['C37','C38','C39','C40','C41','H37','I37','H38','I38','H40','I40','H41','I41','H44','I44','H45','I45','N53']:
            assert before[f'{m}월'][a].value == after[f'{m}월'][a].value, (m,a)
        for a in ['W12','W13','W14','O149','O150']:
            assert before[f'{m}월계산기'][a].value == after[f'{m}월계산기'][a].value, (m,a)
    assert_no_errors(after)
    bools=[('검증',c.coordinate) for row in after['검증'] for c in row if isinstance(c.value,bool)]
    for sn,a in bools:
        assert after[sn][a].value is True, (sn,a)
    for s in before:
        if s.title == '검증':
            continue
        for row in s:
            for c in row:
                if isinstance(c.value,bool):
                    assert after[s.title][c.coordinate].value == c.value, (s.title,c.coordinate)
    with ZipFile(baseline) as old, ZipFile(final) as new:
        assert old.read('xl/styles.xml') == new.read('xl/styles.xml')
        for n in old.namelist():
            if n == 'xl/vbaProject.bin' or n.startswith(('xl/drawings/','xl/media/','xl/ctrlProps/','xl/activeX/')) or n.endswith('.rels'):
                assert old.read(n) == new.read(n), n
        assert b'application/vnd.ms-excel.sheet.macroEnabled.main+xml' in new.read('[Content_Types].xml')
        assert new.testzip() is None
    return {'label_and_usage_linked':True,'new_raw_exactly_covered':True,
            'principal_return_linked_once':True,'other_raw_values_preserved':True,
            'monthly_living_salary_and_tetris_totals_unchanged':True,
            'original_style_definitions_preserved':True,'vba_controls_and_drawings_preserved':True,
            'sheet_order_and_visibility_preserved':True,'all_boolean_checks_true':len(bools)}


def run(a):
    root=Path(a.workdir);root.mkdir(parents=True,exist_ok=True)
    months=[];default=None
    for month in range(1,13):
        def query(p,m=month):
            p.write('예비비','B3',datetime(2026,m,1));p.write('잔액','B3',datetime(2026,m,1))
        if month == 10 and a.default_recalc:
            path=Path(a.default_recalc)
        else:
            path=recalculate(a.candidate,root/f'month_{month:02d}',query)
        w=read(path);assert_no_errors(w);months.append(verify_month(w,month))
        if month==10:default=path
        print(json.dumps({'month_passed':month,**months[-1]},ensure_ascii=False),flush=True)
        del w;gc.collect()
    cases=[]
    def case(name,edit,check):
        path=recalculate(a.candidate,root/name,edit)
        w=read(path);assert_no_errors(w);check(w)
        cases.append({'name':name,'passed':True,'expected_guard_or_outstanding':True})
        print(json.dumps(cases[-1]),flush=True);del w;gc.collect()
    def remove(p):
        for c in ['B','C','D']:p.write('예비비',f'{c}438',None)
    def remove_check(w):
        assert w['예비비']['F10'].value==39000
        assert w['예비비']['L10'].value==239000
    case('missing_cash_link_remains_outstanding',remove,remove_check)
    def over(p):
        p.write('예비비','E107',240000);p.write('예비비','D438',40000)
    def over_check(w):
        assert w['예비비']['M438'].value==0
        assert w['예비비']['N438'].value.startswith('확인')
        assert w['예비비']['F10'].value==40000
    case('cannot_spend_returned_61000_again',over,over_check)
    def wrong(p):p.write('예비비','B438','U007')
    def wrong_check(w):
        assert w['예비비']['M438'].value==0
        assert w['예비비']['N438'].value.startswith('확인')
    case('cash_exception_does_not_allow_other_uses',wrong,wrong_check)
    def duplicate(p):
        p.write('예비비','E107',278000)
        for c,v in [('B','U008'),('C',208),('D',39000)]:p.write('예비비',f'{c}439',v)
    def duplicate_check(w):
        assert w['예비비']['M438'].value==w['예비비']['M439'].value==0
        assert w['예비비']['F10'].value==78000
    case('cash_source_total_cannot_exceed_net_39000',duplicate,duplicate_check)
    count=cache_original(a.candidate,default,a.output,a.baseline)
    checks=verify_artifact(a.baseline,a.output,a.incoming)
    result={'as_of':'2026-10-15','month_tests':months,'scenario_tests':cases,
            'checks':checks,'formula_errors':0,'cached_formulas':count,
            'selected_month_usage':239000,'funding':{'명절비':200000,'임시보관':39000},
            'principal_return':61000,'family_interest_remaining':25,
            'wrap_total':6685466,'wrap_purpose_cash':2150000,'outstanding_reserve_settlement':0,
            'raw_count':2465,'october_raw_count':144,
            'workbook_sha256':sha256(Path(a.output).read_bytes()).hexdigest(),
            'recalculation':'Disposable macro-free XLSX in LibreOffice; only caches copied to original XLSM.',
            'excel_ui_and_vba_executed':False}
    Path(a.report).write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'verified':True,'months':12,'targeted_scenarios':len(cases),'hash':result['workbook_sha256']}),flush=True)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    for n in ['baseline','candidate','incoming','output','workdir','report']:
        parser.add_argument('--'+n,required=True)
    parser.add_argument('--default-recalc')
    run(parser.parse_args())
