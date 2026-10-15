"""Apply the confirmed birthday settlement without reserializing Excel controls.

The 2026-10-15 replacement RAW adds a 61,000-won principal return and edits
metadata of the existing 100,000-won sibling refund. Neither is new consumption.
"""
import argparse
from datetime import datetime
from pathlib import Path
import re

from lxml import etree as ET
import openpyxl

from redesign_reserve import (
    ExtendedPackage, BACK, S0, S1, settlement_permission_formula,
)
from update_20261015 import Q


def apply(source, incoming, output):
    w = openpyxl.load_workbook(source, data_only=False)
    d = openpyxl.load_workbook(source, data_only=True)
    raw = openpyxl.load_workbook(incoming, data_only=True).active
    rows = [tuple(c.value for c in row) for row in raw.iter_rows(min_row=2)]
    assert len(rows) == 144 and d['RAW'].max_row == 2465
    returns = [v for v in rows if v[1:3] == ('가족통장', '미래에셋_WRAP형CMA')]
    assert len(returns) == 1
    v = returns[0]
    assert v[0].date() == datetime(2026, 10, 15).date()
    assert v[5] == v[8] == 61000 and v[6] == '이체지출'
    refunds = [v for v in rows if v[1] == '가족통장' and v[5] == 100000 and v[6] == '지출']
    assert len(refunds) == 1
    assert d['예비비']['A107'].value == 'U008'
    assert d['예비비']['E107'].value == 200000
    assert d['잔액']['A246'].value is None
    assert d['예비비']['B438'].value is None
    p = ExtendedPackage(source)

    # Extend only RAW range endpoints, never stable RAW row identities.
    for s in w:
        if s.title.startswith(('기존_', '보관_')):
            continue
        for row in s:
            for c in row:
                if c.data_type == 'f':
                    formula = re.sub(r'(RAW!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?)2465\b',
                                     r'\g<1>2466', c.value)
                    if formula != c.value:
                        p.write(s.title, c.coordinate, formula, True)

    p.clone_row('RAW', 2454, 2466)
    note = ('10/15 재첨부 RAW 2행. 생신 본인 납입 300,000 중 미사용 원금 61,000 회수. '
            '수입·환불 소비 아님. WRAP 임시보관에 1회 복귀; 가족통장 이자 25원은 제외.')
    values = [v[0], datetime(2026, 10, 1), v[1], v[2], None, v[4],
              61000, v[6], None, '내부이체', 61000, '본인', note,
              None, None, None, None, 'WRAP', '이체입금']
    for i, value in enumerate(values, 1):
        p.write('RAW', f'{openpyxl.utils.get_column_letter(i)}2466', value)
    link_count = '(COUNTIF(잔액!$W$64:$W$500,ROW())+COUNTIF(잔액!$X$64:$X$500,ROW()))'
    p.write('RAW', 'N2466', f'=IF(R2466="WRAP",IF({link_count}=1,"연결",'
            f'IF({link_count}=0,"누락","중복")),"타계좌·잔액제외")', True)
    for tag in ['dimension', 'autoFilter']:
        p.sheet('RAW').find(Q(tag)).set('ref', 'A1:S2466')
    for cf in p.sheet('RAW').findall(Q('conditionalFormatting')):
        cf.set('sqref', cf.get('sqref').replace('N2323:N2465', 'N2323:N2466'))
    # This is an edit to the existing refund, not a second 100,000-won outflow.
    changed = refunds[0]
    for col, value in [('A', changed[0]), ('D', changed[2]), ('E', changed[3])]:
        p.write('RAW', f'{col}2422', value)
    p.write('RAW', 'M2422', '10/15 재첨부 RAW 43행으로 시각·분류 수정. 기존 형제자매 잔금 반환 100,000원. 신규 소비·본인 반환 아님.')
    p.write('RAW', 'M2401', '부모 생신 용돈 500,000원. 본인 행사 부담 239,000원과 형제자매 부담 400,000원으로 식사까지 충당.')
    p.write('RAW', 'Q2', '10/1~15 RAW 144건: WRAP 14건 연결. 타계좌 현금 45건·카드 사용 85건 보존. 가족통장 원금 61,000원 회수 추가.')
    p.write('RAW', 'Q5', '생신 본인 사용 239,000원 = 명절비 200,000 + 임시보관 39,000. 기존 현금 납입으로 충당; 월급 부담에 다시 더하지 않음.')
    p.write('RAW', 'Q6', '생신 실비 639,000 / 형제자매 반환 100,000 / 본인 원금 회수 61,000. 철도 대납 미회수 103,600원. 이체 중복 소비 제외.')
    ledger = {'A246': v[0], 'B246': '이체', 'C246': '외부', 'D246': '임시보관',
              'E246': 61000, 'F246': '부모 생신 미사용 원금 회수',
              'G246': 'RAW 2466행. 8월 가족통장 납입 100,000 중 행사 39,000 사용, 원금 61,000 회수. 신규 수입·예비비 적립 아님.',
              'J246': '가족통장', 'K246': '미래에셋_WRAP형CMA',
              'L246': 'WRAP계좌이체', 'W246': 2466}
    for addr, value in ledger.items():
        p.write('잔액', addr, value)
    p.write('잔액', 'G208', '원본 8월 임시보관의 가족통장 납입 100,000원. 10월 생신 사용 39,000원에 배분하고 잔금 61,000원은 원장 246행에서 회수. 날짜 미기록: 월 기준일.')

    p.write('예비비', 'L7', '선택월 예비비 사용액')
    row7 = p.sheet('예비비').find(Q('sheetData')).find(f'{Q("row")}[@r="7"]')
    row7.set('ht', '42'); row7.set('customHeight', '1')
    p.write('예비비', 'E107', 239000)
    p.write('예비비', 'H107', '본인 납입 300,000 − 잔금 회수 61,000 = 239,000. 명절비 200,000 + 임시보관 39,000. 가족통장 이자 25원 별도.')
    p.write('예비비', 'A66', '생신 사용 239,000원 = 명절비 200,000 + 임시보관 39,000. 잔금 61,000원 회수는 신규 수입·사용이 아닙니다.')
    for col, value in [('B', 'U008'), ('C', 208), ('D', 39000), ('H', '원장'),
                       ('I', '임시보관 납입 100,000 중 행사 부담 39,000. 나머지 61,000은 10/15 원장 246행 회수')]:
        p.write('예비비', f'{col}438', value)
    for r in range(S0, S1 + 1):
        p.write('예비비', f'M{r}', settlement_permission_formula(r), True)
    # Show both prepayments and the principal return without changing any cash math.
    for r in range(120, 557):
        existing = w[BACK][f'BH{r}'].value
        relation = (f'OR(COUNTIFS(예비비!$C$430:$C$729,BA{r},예비비!$N$430:$N$729,"OK")>0,'
                    f'AND(INDEX(잔액!$J:$J,BA{r})="가족통장",'
                    f'INDEX(잔액!$K:$K,BA{r})="미래에셋_WRAP형CMA"),'
                    f'COUNTIF(예비비!$A$8:$A$12,INDEX(잔액!$C:$C,BA{r}))=1,')
        old_start = f'OR(COUNTIF(예비비!$A$8:$A$12,INDEX(잔액!$C:$C,BA{r}))=1,'
        assert old_start in existing
        p.write(BACK, f'BH{r}', existing.replace(old_start, relation), True)
    p.write(BACK, 'Q7', '=예비비!E107', True)
    p.write(BACK, 'F7', '10월 실제 사용 239,000원. 명절비 선이체 200,000 + 임시보관 39,000. WRAP 추가 차감 없음.')
    p.write(BACK, 'A65', '10월 생신 실제 사용 239,000원은 명절비 선이체 200,000원과 임시보관 39,000원으로 충당. 잔금 61,000원은 WRAP 임시보관으로 회수; 소비·적립 아님.')
    p.write(BACK, 'N74', '원본 적립 차이 −100,000 보존. 생신 사용 239,000은 명절비 200,000 + 임시보관 39,000으로 충당하여 경조사비 잔액에서 다시 빼지 않음.')
    p.write('10월', 'S66', '생신 본인 사용 239,000 = 명절비 200,000 + 임시보관 39,000. 기존 현금 납입과 연결; WRAP·월급 재차감 없음.')
    p.write('10월', 'J76', '용돈 500,000 + 식사 139,000 = 639,000. 본인 부담 239,000 + 형제자매 400,000.')
    p.write('10월', 'E78', '=SUMIFS(RAW!$G$2323:$G$2466,RAW!$C$2323:$C$2466,"가족통장",RAW!$H$2323:$H$2466,"지출")+'
            'SUMIFS(RAW!$G$2323:$G$2466,RAW!$C$2323:$C$2466,"가족통장",RAW!$H$2323:$H$2466,"이체지출")', True)
    p.write('10월', 'J78', '실제 출금 800,000 = 행사 639,000 + 형제자매 반환 100,000 + 본인 원금 회수 61,000. 식사 정산 중복 제외.')
    p.write('10월', 'J79', '본인 납입 300,000 − 회수 61,000 = 사용 239,000. 가족통장 이자 25원 잔존. WRAP·당월 월급 중복 차감 없음.')
    p.write('10월', 'A70', '10/1~15 RAW 144건. 생신 본인 사용 239,000원과 원금 61,000원 회수 연결. 카드대금 5건은 중복 소비·WRAP 차감 제외.')
    p.write('10월', 'E83', '=COUNT(RAW!G2323:G2466)-144', True)
    p.write('10월', 'J83', '0이면 최신 10/15 첨부 RAW 144건 모두 유지. 기존 반환 건은 수정이며 중복 추가하지 않음.')
    for row, expected in [(35, 144), (42, 239000), (63, 144), (64, 14), (66, 14),
                          (90, 239000), (91, 239000), (170, 687510), (174, 6685466)]:
        p.write('검증', f'M{row}', f'=I{row}={expected}', True)
    p.write('검증', 'I44', '=COUNTIF(잔액!$I$231:$I$246,"입력 확인")+COUNTIF(잔액!$I$231:$I$246,"WRAP 범위 확인")', True)
    calc = p.tree('xl/workbook.xml').find(Q('calcPr'))
    calc.set('calcMode', 'auto'); calc.set('fullCalcOnLoad', '1'); calc.set('forceFullCalc', '1')
    p.save(output)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['source', 'incoming', 'output']:
        parser.add_argument('--' + name, required=True)
    args = parser.parse_args()
    apply(args.source, args.incoming, args.output)
