"""Apply the 2026-10-15 RAW delta without reserializing the XLSM package.

VBA, drawings, controls, relationships and untouched sheets remain original ZIP
parts. A macro-free temporary copy may be recalculated; only its value caches
are imported into the original package. No uploaded instructions are executed.
"""
from collections import Counter, defaultdict
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from zipfile import ZipFile
import argparse
import json
import posixpath
import re

from lxml import etree as ET
import openpyxl
from openpyxl.formula.translate import Translator
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.utils.datetime import to_excel

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
PNS = 'http://schemas.openxmlformats.org/package/2006/relationships'
Q = lambda tag: '{'+NS+'}'+tag
ASOF = datetime(2026, 10, 15)
MONTH = datetime(2026, 10, 1)
START = 2323
CATEGORIES = {'변동비','담배','일회성전자','경조·선물','여행','약속·만남'}


class Package:
    def __init__(self, path):
        self.path = Path(path)
        with ZipFile(path) as z:
            self.infos = z.infolist()
            self.parts = {n:z.read(n) for n in z.namelist()}
        self.trees = {}
        rels = ET.fromstring(self.parts['xl/_rels/workbook.xml.rels'])
        targets = {r.get('Id'):r.get('Target') for r in rels}
        workbook = self.tree('xl/workbook.xml')
        self.sheets = {}
        for sheet in workbook.find(Q('sheets')):
            target = targets[sheet.get('{'+RNS+'}id')]
            self.sheets[sheet.get('name')] = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/'+target)
        self.cell_maps = {}
        self.row_maps = {}

    def tree(self, part):
        if part not in self.trees:
            self.trees[part] = ET.fromstring(self.parts[part], ET.XMLParser(resolve_entities=False))
        return self.trees[part]

    def sheet(self, name):
        return self.tree(self.sheets[name])

    def cell(self, sheet, addr):
        data = self.sheet(sheet).find(Q('sheetData'))
        if sheet not in self.cell_maps:
            self.row_maps[sheet] = {int(r.get('r')):r for r in data}
            self.cell_maps[sheet] = {c.get('r'):c for r in data for c in r}
        cells = self.cell_maps[sheet]
        if addr not in cells:
            rownum = int(re.search(r'\d+',addr)[0])
            rows = self.row_maps[sheet]
            if rownum not in rows:
                row = ET.Element(Q('row'),r=str(rownum))
                pos = next((i for i,r in enumerate(data) if int(r.get('r'))>rownum),len(data))
                data.insert(pos,row);rows[rownum]=row
            row = rows[rownum]
            cell = ET.Element(Q('c'),r=addr)
            col = column_index_from_string(re.match('[A-Z]+',addr)[0])
            pos=next((i for i,c in enumerate(row) if column_index_from_string(re.match('[A-Z]+',c.get('r'))[0])>col),len(row))
            row.insert(pos,cell);cells[addr]=cell
        return cells[addr]

    def write(self, sheet, addr, value, formula=False):
        cell=self.cell(sheet,addr)
        for child in list(cell):
            if child.tag in [Q('f'),Q('v'),Q('is')]:cell.remove(child)
        cell.attrib.pop('t',None)
        if formula:
            ET.SubElement(cell,Q('f')).text=value.lstrip('=')
        elif value is not None:
            if isinstance(value,datetime):value=to_excel(value)
            if isinstance(value,bool):
                cell.set('t','b');ET.SubElement(cell,Q('v')).text='1' if value else '0'
            elif isinstance(value,(int,float)):
                ET.SubElement(cell,Q('v')).text=str(value)
            else:
                cell.set('t','inlineStr');text=ET.SubElement(ET.SubElement(cell,Q('is')),Q('t'))
                text.set('{http://www.w3.org/XML/1998/namespace}space','preserve');text.text=str(value)

    def clone_row(self, sheet, source, destination):
        self.cell(sheet,'A'+str(source))
        data=self.sheet(sheet).find(Q('sheetData'))
        row=deepcopy(self.row_maps[sheet][source]);row.set('r',str(destination))
        for cell in row:cell.set('r',re.sub(r'\d+$',str(destination),cell.get('r')))
        data.append(row)
        self.row_maps[sheet][destination]=row
        self.cell_maps[sheet].update({c.get('r'):c for c in row})

    def save(self,path):
        for part,tree in self.trees.items():
            self.parts[part]=ET.tostring(tree,encoding='UTF-8',xml_declaration=True)
        with ZipFile(path,'w') as z:
            for info in self.infos:z.writestr(info,self.parts[info.filename])


def prepare(baseline, source, output, summary_path):
    w=openpyxl.load_workbook(baseline,data_only=False)
    d=openpyxl.load_workbook(baseline,data_only=True)
    raw=d['RAW'];incoming=openpyxl.load_workbook(source,data_only=True).active
    assert list(next(incoming.values))==['날짜','자산','분류','소분류','내용','금액(원)','수입/지출','메모','금액','화폐']
    old=defaultdict(list)
    for r in range(START,raw.max_row+1):old[(raw[f'A{r}'].value,raw[f'G{r}'].value,raw[f'H{r}'].value)].append(r)
    matched=set(); additions=[];source_map=[]
    for r in range(2,incoming.max_row+1):
        v=[incoming.cell(r,c).value for c in range(1,11)]
        assert isinstance(v[0],datetime) and v[0].year==2026 and v[0].month==10
        assert v[9]=='KRW' and v[5]==v[8] and isinstance(v[5],int)
        match=old.get((v[0],v[5],v[6]),[])
        if match:
            assert len(match)==1
            canonical=match[0];assert canonical not in matched;matched.add(canonical)
            assert (raw[f'C{canonical}'].value==v[1] or (raw[f'C{canonical}'].value=='모임통장' and v[1]=='20260606'))
            source_map.append({'source_row':r,'workbook_row':canonical,'status':'preserved'})
        else:additions.append((r,v))
    assert len(matched)==131 and len(additions)==12
    additions.sort(key=lambda x:x[1][0])
    p=Package(baseline);end=raw.max_row+len(additions)
    new_rows=[];new_personal=[];new_counts=Counter()
    payment_total=0
    for offset,(sr,v) in enumerate(additions,1):
        r=raw.max_row+offset
        date,asset,cat,sub,title,amount,kind,note,_,_=v
        if kind=='이체지출' and title=='카드정산 결제':
            classification='내부이체';payment='본인';burden=amount;scope='타계좌';direction=None
            candidates=[i for i in range(START,raw.max_row+1) if raw[f'C{i}'].value=='미래에셋_WRAP형CMA' and raw[f'D{i}'].value==asset and raw[f'G{i}'].value==amount and raw[f'H{i}'].value=='이체지출']
            assert len(candidates)==1
            memo=f'10/15 RAW {sr}행 · 카드대금 납부. WRAP 준비송금 RAW {candidates[0]}행과 연결. 소비·WRAP에서 다시 차감하지 않음.'
            payment_total+=amount
        elif kind=='수입' and asset=='미래에셋_WRAP형CMA':
            classification='이자·환급';payment='본인';burden=amount;scope='WRAP';direction='입금'
            memo=f'10/15 RAW {sr}행 · WRAP 실제 이자 입금. 원장에 1회 가산.'
        else:
            assert kind=='지출' and ('카드' in asset or asset=='KB쿠팡와우')
            scope='카드사용';direction=None;payment='본인';burden=amount
            classification='여행' if cat=='교통/차량' else '변동비'
            memo=f'10/15 RAW {sr}행 · 신규 카드 사용. 월간·계산기 반영, WRAP 원장 제외.'
            if title=='코레일' and amount==51800:
                classification='대납';burden=0;payment='정산대기';note='대리결제·미회수'
                memo=f'10/15 RAW {sr}행 · 사용자 확인: 10/14의 별도 신규 철도 대리결제. 기존 10/12 대납과 각각 유지. 회수 입금 없음, 본인 소비 제외.'
            else:
                units=(amount+999)//1000
                category_row=5 if classification=='여행' else 4 if cat=='기타' else 3
                assert classification=='여행' or cat in ['식비','기타']
                new_personal.append({'row':r,'date':date.isoformat(),'amount':amount,'units':units,'calculator_row':category_row})
                new_counts[category_row]+=units
        p.clone_row('RAW',2453,r)
        values=[date,MONTH,asset,cat,sub or None,title,amount,kind,note or None,classification,burden,payment,memo,None,None,None,None,scope,direction]
        for c,value in enumerate(values,1):p.write('RAW',get_column_letter(c)+str(r),value)
        formula=f'=IF(R{r}="WRAP",IF((COUNTIF(잔액!$W$64:$W$500,ROW())+COUNTIF(잔액!$X$64:$X$500,ROW()))=1,"연결",IF((COUNTIF(잔액!$W$64:$W$500,ROW())+COUNTIF(잔액!$X$64:$X$500,ROW()))=0,"누락","중복")),IF(R{r}="카드사용","월간·계산기","타계좌·잔액제외"))'
        p.write('RAW',f'N{r}',formula,True)
        new_rows.append({'source_row':sr,'row':r,'date':date.isoformat(),'amount':amount,'classification':classification,'scope':scope})
        source_map.append({'source_row':sr,'workbook_row':r,'status':'added'})
    assert payment_total==2396039
    rawtree=p.sheet('RAW');rawtree.find(Q('dimension')).set('ref',f'A1:S{end}')
    rawtree.find(Q('autoFilter')).set('ref',f'A1:S{end}')
    for c in rawtree.findall(Q('conditionalFormatting')):
        c.set('sqref',c.get('sqref').replace('N2323:N2453',f'N2323:N{end}'))
    p.write('RAW','Q1','10월 RAW 갱신 · 2026-10-15 기준')
    p.write('RAW','Q2','10/1~15 RAW 143건 보존: WRAP 13건만 잔액 원장에 연결. 타계좌 현금 45건·카드 사용 85건은 월간 자료로 보존.')
    interest=next(r for r in new_rows if r['scope']=='WRAP')
    assert interest['amount']==535
    ledger_values={'A245':datetime.fromisoformat(interest['date']),'B245':'입금','C245':'외부','D245':'생활비','E245':535,'F245':'미래에셋 WRAP 이자','G245':f'RAW {interest["row"]}행 · 10/13 추가 이자 535원. 10/15 RAW 재대조, WRAP에 1회 가산.','J245':'외부','K245':'미래에셋_WRAP형CMA','L245':'WRAP입출금','W245':interest['row'],'X245':None}
    for addr,value in ledger_values.items():p.write('잔액',addr,value)
    p.write('잔액','D3','2026-10-15 기준 · WRAP 입출금 RAW 13건. 타계좌 카드대금 납부 5건은 준비송금 후속 결제로 중복 반영하지 않습니다.')
    p.write('10월','U1',ASOF)
    for col in ['D','E']:
        date=d['10월'][col+'13'].value
        assert date.day in [14,15]
        p.write('10월',col+'15',Translator(w['10월']['B15'].value,origin='B15').translate_formula(col+'15'),True)
        rr=f'RAW!$A${START}:$A${end}';kind=f'RAW!$H${START}:$H${end}';pay=f'RAW!$L${START}:$L${end}';cl=f'RAW!$J${START}:$J${end}';burden=f'RAW!$K${START}:$K${end}'
        conditions='+'.join(f'({cl}="{c}")' for c in sorted(CATEGORIES))
        f=f'=SUMPRODUCT(({rr}>={col}13)*({rr}<{col}13+1)*({kind}="지출")*({pay}="본인")*({conditions})*ROUNDUP({burden}/1000,0)*1000)'
        p.write('10월',col+'14',f,True)
    for r in range(4,25):
        note=w['10월'][f'V{r}'].value
        if note:
            # Keep the fixed row heights: update the audit date in-place rather
            # than appending paragraphs that would be clipped in the memo cell.
            note=note.replace('10/13 기준','10/15 기준').replace('10/13 RAW 집계','10/15 RAW 집계')
            p.write('10월',f'V{r}',note)
    p.write('10월','E80',f'=SUMIFS(RAW!$G${START}:$G${end},RAW!$J${START}:$J${end},"대납",RAW!$L${START}:$L${end},"정산대기",RAW!$H${START}:$H${end},"지출")',True)
    p.write('10월','J80','사용자 확인: 10/12·10/14 각 51,800원 철도 대납. 회수 입금 없음. 미회수 합계 103,600원, 본인 소비 제외.')
    p.write('10월','A70','10/1~15 RAW 143건 반영. 기존 131건 유지 + 신규 12건. 타계좌 카드대금 5건 2,396,039원은 중복 소비·WRAP 차감 제외. 신규 철도 대납 51,800원 별도 추적.')
    p.write('10월','A72','10/15 실제 / 예정 / 공동자금 · 달력 합계와 연결')
    p.write('10월','J82','각 본인 소비 거래를 천원 단위 올림해 합산. 대납 2건·목적자금·공동자금·카드대금 납부 제외.')
    p.write('10월','E83',f'=COUNT(RAW!G{START}:G{end})-143',True)
    p.write('10월','J83','0이면 10/15 첨부 RAW 143건 모두 유지')
    actual=1218440+sum(r['amount'] for r in new_personal)
    tetris=1238000+sum(r['units']*1000 for r in new_personal)
    assert actual==1299740 and tetris==1321000
    p.write('10월','E84',f'=I33-{actual}',True)
    p.write('10월','E85',f'=SUM(I4,I9,I14,I19,I24,I29)-{tetris}',True)
    p.write('10월','J86','급여 미입금: 10/15 중간 집계이며 월말 결과가 아님')
    # Append to the original block sequence; never repaint prior work.
    calc=w['10월계산기'];order=[f'{get_column_letter(c)}{r}' for r0 in range(15,145,5) for c0 in range(2,22,2) for c in range(c0,c0+2) for r in range(r0,r0+5)]
    filled=[calc[a].fill.patternType=='solid' for a in order]
    assert filled==[True]*1238+[False]*(len(order)-1238)
    styles=p.tree('xl/styles.xml');xfs=styles.find(Q('cellXfs'))
    cache={ET.tostring(x):i for i,x in enumerate(xfs)}
    def paint(addr,palette_row):
        cell=p.cell('10월계산기',addr);xf=deepcopy(xfs[int(cell.get('s','0'))])
        legend=p.cell('10월계산기','B'+str(palette_row));fillid=xfs[int(legend.get('s','0'))].get('fillId')
        xf.set('fillId',fillid)
        key=ET.tostring(xf)
        if key not in cache:cache[key]=len(xfs);xfs.append(xf)
        cell.set('s',str(cache[key]))
    index=1238
    for row in new_personal:
        for _ in range(row['units']):paint(order[index],row['calculator_row']);index+=1
    xfs.set('count',str(len(xfs)))
    for r in range(3,11):p.write('10월계산기',f'AA{r}',d['10월계산기'][f'AA{r}'].value+new_counts[r])
    p.write('10월계산기','A154','10/15까지 본인 생활비 반영. 기존 색상·블록 유지 + 83칸 추가. 미회수 대납 2건 103,600원·가족 공동자금·예비비 200,000원은 월간 별도 표시.')
    # Current coverage checks grow; the dated 10/13 preparation batch stays fixed.
    for row in w['검증']:
        for cell in row:
            if cell.data_type=='f' and ('2323' in cell.value):
                p.write('검증',cell.coordinate,cell.value.replace('$2453',f'${end}').replace(':G2453',f':G{end}').replace(':R2453',f':R{end}').replace(':N2453',f':N{end}').replace(':S2453',f':S{end}'),True)
    for addr,n in {'M35':143,'M63':143,'M64':13,'M65':45,'M66':13,'M70':13,'M169':9092,'M174':6624466}.items():
        sourceaddr='I'+addr[1:];p.write('검증',addr,f'={sourceaddr}={n}',True)
    p.write('검증','I44','=COUNTIF(잔액!$I$231:$I$245,"입력 확인")+COUNTIF(잔액!$I$231:$I$245,"WRAP 범위 확인")',True)
    p.write('검증','I70','=COUNT(잔액!W231:W245)',True)
    p.write('검증','Q14','10/15까지 · 대납 2건 103,600원·가족 공동자금 분리')
    p.write('검증','A24','9월 가족통장 선이체 20만원은 10월 사용입니다. 행사 실비 639,000원·반환 100,000원·본인 예비비 200,000원 구분을 유지합니다.')
    p.write('검증','A25','식사 정산이체 139,000원 중복 제외. 철도 대납은 10/12·10/14 각 51,800원, 미회수 103,600원입니다.')
    p.write('검증','A49','10/13 결제 준비 이체 최초 반영분 검증')
    p.write('검증','A50','10/13 최초 반영 RAW 거래')
    historical=w['검증']['I53'].value
    historical=historical.replace('DATE(2026,10,13))','DATE(2026,10,13),잔액!$A$64:$A$500,"<"&DATE(2026,10,14))')
    p.write('검증','I53',historical,True)
    p.write('검증','A59','결제 준비 이체는 소비·자산형성에서 제외합니다. 후속 카드대금 납부 5건은 타계좌 RAW에 보존하고 WRAP 원장에서 재차 차감하지 않습니다. 실제 대출이자는 월간 비용에 1회 반영합니다.')
    p.write('검증','A72','WRAP 13건만 연결합니다. 타계좌 현금 45건·카드사용 85건은 RAW와 월간에 보존합니다. 가족통장 사용 20만원·철도 미회수 대납 103,600원은 월간 별도 표시합니다.')
    p.write('사용안내','A16','가족통장·모임통장의 실제 사용은 월간에서 확인합니다. WRAP 밖 보관금은 합계에 포함하지 않습니다. 가족행사 예비비 20만원과 철도 미회수 대납 2건 103,600원은 별도 표시합니다.')
    p.write('사용안내','A22','10월은 10/15까지 실제 자료입니다. 미수령 급여·미결제 고정비·공동자금·대납은 월간 확인표에 별도로 표시합니다.')
    p.write('사용안내','A23','10월 RAW 143건 중 WRAP 13건만 잔액 원장 연결을 검증합니다. 타계좌 45건·카드사용 85건은 보존하며 월간·계산기의 검증 자료로 사용합니다.')
    calculation=p.tree('xl/workbook.xml').find(Q('calcPr'))
    calculation.set('fullCalcOnLoad','1');calculation.set('forceFullCalc','1')
    p.save(output)
    summary={'as_of':ASOF.date().isoformat(),'source_sha256':sha256(Path(source).read_bytes()).hexdigest(),'baseline_sha256':sha256(Path(baseline).read_bytes()).hexdigest(),'preserved_october_rows':131,'new_rows':new_rows,'source_mapping':sorted(source_map,key=lambda x:x['source_row']),'october_count':143,'raw_total':2464,'last_raw_row':end,'personal_actual':actual,'personal_tetris':tetris,'new_personal':new_personal,'new_card_payments':{'count':5,'amount':payment_total,'wrap_effect':0,'consumption_effect':0},'new_wrap_interest':535,'wrap_total':6624466,'outstanding_advances':103600,'user_confirmed':'10/12 and 10/14 rail payments are distinct; each 51,800 is an unreimbursed advance.'}
    Path(summary_path).write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({k:summary[k] for k in ['as_of','preserved_october_rows','october_count','raw_total','personal_actual','personal_tetris','wrap_total','outstanding_advances']},ensure_ascii=False))


def macro_free(source,destination):
    with ZipFile(source) as z:
        parts={n:z.read(n) for n in z.namelist()}
    types=ET.fromstring(parts['[Content_Types].xml'])
    for x in list(types):
        if 'vbaProject' in x.get('ContentType',''):types.remove(x)
        elif x.get('PartName')=='/xl/workbook.xml':x.set('ContentType','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml')
    parts['[Content_Types].xml']=ET.tostring(types,encoding='UTF-8',xml_declaration=True)
    rels=ET.fromstring(parts['xl/_rels/workbook.xml.rels'])
    for r in list(rels):
        if 'vbaProject' in r.get('Type',''):rels.remove(r)
    parts['xl/_rels/workbook.xml.rels']=ET.tostring(rels,encoding='UTF-8',xml_declaration=True)
    # Calc otherwise trusts Excel's pre-existing caches, including dependents of
    # changed inputs. Force a real recalculation in this disposable copy.
    for name,body in list(parts.items()):
        if name.startswith('xl/worksheets/sheet') and name.endswith('.xml'):
            sheet=ET.fromstring(body)
            for cell in sheet.iter(Q('c')):
                if cell.find(Q('f')) is not None:
                    cell.attrib.pop('t',None)
                    for child in list(cell):
                        if child.tag in [Q('v'),Q('is')]:cell.remove(child)
            parts[name]=ET.tostring(sheet,encoding='UTF-8',xml_declaration=True)
    with ZipFile(destination,'w') as z:
        for name,body in parts.items():
            if not name.startswith('xl/vbaProject'):z.writestr(name,body)


def import_caches(original,recalculated,output,boolean_reference):
    p=Package(original)
    result=openpyxl.load_workbook(recalculated,data_only=True)
    reference=openpyxl.load_workbook(boolean_reference,data_only=True)
    boolean_cells={(s.title,c.coordinate) for s in reference for row in s for c in row if isinstance(c.value,bool)}
    changed=Counter();errors=[]
    for sheet in p.sheets:
        for row in p.sheet(sheet).find(Q('sheetData')):
            for c in row:
                if c.find(Q('f')) is None:continue
                addr=c.get('r');rc=result[sheet][addr]
                if rc.data_type=='e':errors.append((sheet,addr,rc.value))
                value=rc.value
                if (sheet,addr) in boolean_cells and value in [0,1]:value=bool(value)
                for child in list(c):
                    if child.tag in [Q('v'),Q('is')]:c.remove(child)
                c.attrib.pop('t',None)
                if value is not None:
                    if isinstance(value,datetime):value=to_excel(value)
                    if isinstance(value,bool):c.set('t','b');text='1' if value else '0'
                    elif isinstance(value,(int,float)):text=str(value)
                    else:c.set('t','str');text=str(value)
                    ET.SubElement(c,Q('v')).text=text
                changed[sheet]+=1
    assert not errors,errors[:20]
    p.save(output)
    print(json.dumps({'cached_formulas':sum(changed.values()),'errors':errors,'sheets':dict(changed)},ensure_ascii=False))


if __name__=='__main__':
    cli=argparse.ArgumentParser();sub=cli.add_subparsers(dest='mode',required=True)
    prep=sub.add_parser('prepare')
    for name in ['baseline','source','output','summary']:prep.add_argument('--'+name,required=True)
    temp=sub.add_parser('macro-free');temp.add_argument('source');temp.add_argument('destination')
    cache=sub.add_parser('cache');cache.add_argument('original');cache.add_argument('recalculated');cache.add_argument('output');cache.add_argument('--boolean-reference',required=True)
    args=cli.parse_args()
    if args.mode=='prepare':prepare(args.baseline,args.source,args.output,args.summary)
    elif args.mode=='macro-free':macro_free(args.source,args.destination)
    else:import_caches(args.original,args.recalculated,args.output,args.boolean_reference)
