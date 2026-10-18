"""Align spendable reserve envelopes with balance accounts, not bank cashflow.

Actual ledger/RAW identities are immutable. A month-specific allocation moves
already committed money into the existing payment envelope with zero net cash
effect. Both legacy purpose payouts and new payment-envelope payouts work.
"""
import argparse
from datetime import datetime
import json
from pathlib import Path

from lxml import etree as ET
import openpyxl
from openpyxl.utils import get_column_letter as col

from redesign_reserve import ExtendedPackage, BACK, FUNDS, U0, U1, S0, S1
from remove_validation_sheet import remove_validation_sheet
from reserve_status_ui import apply_ui
from update_20261015 import Q
from validate_excel_structure import validate_package


def raw_balance(account, cutoff):
    return (f'SUMIFS(잔액!$Y$64:$Y$500,잔액!$D$64:$D$500,{account},잔액!$A$64:$A$500,"<"&{cutoff})'
            f'-SUMIFS(잔액!$Y$64:$Y$500,잔액!$C$64:$C$500,{account},잔액!$A$64:$A$500,"<"&{cutoff})')


def allocation(account, month, previous=False):
    index = f'MONTH({month})' + ('' if previous else '+1')
    return (f'IFERROR(INDEX(\'{BACK}\'!$BI$81:$BU$86,'
            f'MATCH({account},\'{BACK}\'!$BH$81:$BH$86,0),{index}),0)')


def available_formula(row):
    account = f'$A{row}'
    return (f'=IF($B$3<DATE(2026,3,1),"원장 없음",'
            f'{raw_balance(account,"EDATE($B$3,1)")}-{allocation(account,"$B$3")})')


def payment_formula():
    account = '"결제대금비"'
    return (f'=IF($B$3<DATE(2026,3,1),"원장 없음",'
            f'{raw_balance(account,"EDATE($B$3,1)")}+{allocation(account,"$B$3")})')


def settlement_permission(r):
    """Same validity checks and historic exceptions; add payment-envelope source."""
    ledger = f'MAX(64,C{r})'
    allocated = f'SUMIF($C${S0}:$C${S1},C{r},$D${S0}:$D${S1})'
    holiday = f'AND(B{r}="U008",C{r}=223,J{r}="경조사비",F{r}="명절비")'
    family = (f'AND(B{r}="U008",C{r}=208,J{r}="경조사비",'
              f'F{r}="임시보관",{allocated}<=39000)')
    opening = (f'AND(H{r}="기초반영",OR(B{r}="U001",B{r}="U002",B{r}="U003"),'
               f'C{r}="",E{r}=DATE(2026,3,1))')
    linked = (f'AND(H{r}="원장",C{r}>=64,C{r}<=500,'
              f'INDEX(잔액!$I:$I,{ledger})="OK",'
              f'E{r}=INDEX(잔액!$A:$A,{ledger}),F{r}=INDEX(잔액!$C:$C,{ledger}),'
              f'OR(F{r}=J{r},F{r}="결제대금비",{holiday},{family}),'
              f'OR(COUNTIF($A$8:$A$12,F{r})=1,F{r}="결제대금비",{family}),'
              f'COUNTIF($A$8:$A$12,INDEX(잔액!$D:$D,{ledger}))=0,'
              f'{allocated}<=INDEX(잔액!$Y:$Y,{ledger}))')
    return f'=IF(B{r}="","",IFERROR(IF(OR({opening},{linked}),1,0),0))'


def prepare(source, output):
    w = openpyxl.load_workbook(source)
    assert '검증' in w.sheetnames
    assert w['예비비']['C7'].value == '원장잔액', 'Use the pre-migration snapshot.'
    assert w['잔액']['A13'].value == '결제대금비'
    p = ExtendedPackage(source)
    for r in range(S0, S1 + 1):
        p.write('예비비', f'M{r}', settlement_permission(r), True)
    # The backend uses only immutable use/settlement records, never query-month
    # UI values. Each of thirteen columns is a separate month-end stock.
    for r, fund in enumerate(FUNDS + ['결제대금비'], 81):
        p.write(BACK, f'BH{r}', fund)
    for c in range(61, 74):
        letter = col(c)
        date = datetime(2025, 12, 1) if c == 61 else datetime(2026, c - 61, 1)
        for r in [80, 99]:
            p.write(BACK, f'{letter}{r}', date)
        cutoff = f'EDATE({letter}$99,1)'
        for r in range(U0, U1 + 1):
            uid, use_date, amount = [f'예비비!${x}{r}' for x in ['A', 'D', 'E']]
            paid = (f'SUMIFS(예비비!$O${S0}:$O${S1},예비비!$B${S0}:$B${S1},{uid},'
                    f'예비비!$E${S0}:$E${S1},"<"&{cutoff})')
            payment = paid[:-1] + f',예비비!$F${S0}:$F${S1},"결제대금비")'
            recognized = f'IF({use_date}<{cutoff},{amount},0)'
            formula = f'=IF(예비비!$R{r}="OK",{payment}+MAX(0,{recognized}-{paid}),0)'
            p.write(BACK, f'{letter}{r}', formula, True)
        for r in range(81, 86):
            p.write(BACK, f'{letter}{r}',
                    f'=SUMIF(예비비!$C${U0}:$C${U1},$BH{r},{letter}${U0}:{letter}${U1})', True)
        p.write(BACK, f'{letter}86', f'=SUM({letter}81:{letter}85)', True)
    p.sheet(BACK).find(Q('dimension')).set('ref', 'A1:BU556')
    backend_cols = p.sheet(BACK).find(Q('cols'))
    if backend_cols is None:
        backend_cols = ET.Element(Q('cols'))
        tree = p.sheet(BACK)
        tree.insert(list(tree).index(tree.find(Q('sheetData'))), backend_cols)
    # Existing columns end at 60; preserve original style/width definitions.
    assert all(int(x.get('max')) < 61 for x in backend_cols)
    ET.SubElement(backend_cols, Q('col'), min='61', max='73', hidden='1', width='12', customWidth='1')

    # Preserve the source ledger; expose the zero-total automatic allocation in
    # balances and transfer columns. Every row still satisfies B+C-D-E+F=G.
    for r in range(8, 14):
        account = f'$A{r}'
        current = allocation(account, '$B$3')
        previous = allocation(account, '$B$3', previous=True)
        old = {c: w['잔액'][f'{c}{r}'].value.lstrip('=') for c in ['B', 'E', 'F']}
        if r < 13:
            outflow = f'MAX({current}-{previous},0)'
            inflow = f'MAX({previous}-{current},0)'
            prior = f'({old["B"]})-{previous}'
        else:
            changes = [(allocation(f'"{fund}"', '$B$3'), allocation(f'"{fund}"', '$B$3', True)) for fund in FUNDS]
            outflow = '+'.join(f'MAX({before}-{after},0)' for after, before in changes)
            inflow = '+'.join(f'MAX({after}-{before},0)' for after, before in changes)
            prior = f'({old["B"]})+{previous}'
        for c, formula in {'B': prior, 'E': f'({old["E"]})+{outflow}', 'F': f'({old["F"]})+{inflow}'}.items():
            p.write('잔액', f'{c}{r}', '=' + formula, True)
    for c in range(3, 13):
        letter = col(c)
        month = f'{letter}$21'
        cutoff = f'EDATE({month},1)'
        for r in range(22, 31):
            account = f'$A{r}'
            base = raw_balance(account, cutoff).replace('잔액!', '')
            if 23 <= r <= 27:
                base += '-' + allocation(account, month)
            elif r == 28:
                base += '+' + allocation(account, month)
            p.write('잔액', f'{letter}{r}',
                    f'=IF({month}>MAX($A$64:$A$500),"",{base})', True)
    p.write('잔액', 'E6', '이체·분류 보냄')
    p.write('잔액', 'F6', '이체·분류 받음')
    p.write('잔액', 'A18', '계정 잔액에는 예비비 사용분의 결제대금비 자동 확보가 포함됩니다. 은행 거래가 아니며, 아래 원장 조회는 실제 입력 기록입니다.')
    p.write('잔액', 'A33', '목적계정 사용가능액 + 결제대금비 등 모든 계정 = WRAP 현금. 결제대금 자동 확보는 총현금을 바꾸지 않습니다.')
    p.write('잔액', 'A525', '선택월 실제 입력 원장입니다. 자동 결제대금 분류는 상단 계정 잔액에 별도로 반영되며 새 은행 거래가 아닙니다. 원장 편집은 왼쪽 +를 펼칩니다.')
    ui_report = apply_ui(p, available_formula, payment_formula())
    removal = remove_validation_sheet(p)
    book = p.tree('xl/workbook.xml')
    reserve_index = [node.get('name') for node in book.find(Q('sheets'))].index('예비비')
    for name in book.find(Q('definedNames')):
        if name.get('name') == '_xlnm.Print_Area' and name.get('localSheetId') == str(reserve_index):
            name.text = "'예비비'!$A$1:$O$79"
    calc = book.find(Q('calcPr'))
    for k, v in [('calcMode', 'auto'), ('fullCalcOnLoad', '1'), ('forceFullCalc', '1')]:
        calc.set(k, v)
    p.write('사용안내', 'A11', '예비비는 선택월 사용계획·실제 사용·사용가능액을 봅니다. 사용계획은 잔액을 차감하지 않고, 실제 사용분은 결제대금비로 자동 확보합니다.')
    p.write('사용안내', 'A27', '예비비 사용은 등록할 때 한 번 반영합니다. 실제 이체를 사용ID·원장행에 연결하면 자동 확보와 중복되지 않습니다. 적립계획과 사용계획은 별도 입력입니다.')
    p.save(output)
    validation = validate_package(output)
    assert not validation['errors'], validation
    report = {'ui': ui_report, 'removed_validation_sheet': removal, 'structure': validation,
              'reservation_formula': 'paid_from_payment_envelope + MAX(0, recognized_use - all_valid_settlements)',
              'raw_and_source_ledger_preserved': True, 'no_bank_transaction_created': True,
              'new_usage_plans_initially_blank': True, 'savings_plans_preserved': True}
    print(json.dumps(report, ensure_ascii=False), flush=True)
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    for name in ['source', 'output']:
        parser.add_argument('--' + name, required=True)
    parser.add_argument('--manifest')
    args = parser.parse_args()
    result = prepare(args.source, args.output)
    if args.manifest:
        Path(args.manifest).write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n')
